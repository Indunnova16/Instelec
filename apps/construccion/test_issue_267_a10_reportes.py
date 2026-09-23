"""Instelec#267 — A10: 4 Reportes descargables del Presupuesto Planeado de
Construcción (issue Fase 5): PDF ejecutivo, Excel 4 hojas, CSV contable, PPT
7 diapositivas — todos gateados por FORMATO server-side (A7,
``permissions_fin.py::formatos_reporte_permitidos``).

Fixture compartida (mismo criterio de magnitud que
``test_issue_267_a5_kpi_cards.py``, con las 3 bandas de semáforo A4
representadas a propósito):

    Ingresos Operacionales (Ingresos) = -100  -> pct 200% -> rojo
    Gastos de Personal    (Fijo)      =   40  -> pct  80% -> amarillo
    Materiales            (Variable)  =   10  -> pct  20% -> verde
    total_general = -50

Cobertura (tests_requeridos de F2):
- PDF: el HTML fuente (``_construir_html_pdf_presupuesto`` — ver docstring
  de ``reportes_presupuesto.py`` sobre por qué se testea el HTML y no bytes
  del PDF binario) contiene los datos del período; smoke test de que el PDF
  final es un archivo válido con tamaño realista.
- Excel: 4 hojas nombradas (Resumen/Matriz/Rubros/Histórico) con datos
  correctos, incluida ≥1 carga de Histórico.
- CSV: encabezado EXACTO (reuso literal de PlanoFinancieroCsvView) + fila
  agregada correcta.
- PPT: exactamente 7 diapositivas.
- Gate de rol (A7): contador -> solo Excel (200), resto 403; supervisor ->
  los 4 en 403; gerente_financiero -> los 4 en 200; director/
  admin_construccion -> PDF/Excel 200, PPT/CSV 403.
- Botones del template: gateados por ``formatos_reporte_permitidos`` — NO es
  solo un gate server-side, la UI tampoco debe ofrecer lo que el rol no
  puede descargar.
- Edge case dato legacy/ausente: sin presupuesto cargado para el año pedido,
  los 4 reportes responden 200 con el mismo empty-state que la pestaña
  Tabla (nunca 404 — la descarga la pidió el usuario a propósito).
"""
import csv
import io
from decimal import Decimal

from django.test import Client, TestCase
from django.urls import reverse

from apps.construccion.models_fin import (
    HistorialCargaPresupuestoConstruccion,
    PresupuestoDetalladoConstruccion,
)
from apps.construccion.reportes_presupuesto import (
    _construir_html_pdf_presupuesto,
    construir_contexto_reporte,
    escribir_csv_presupuesto,
    generar_excel_presupuesto,
    generar_pdf_presupuesto,
    generar_ppt_presupuesto,
)
from apps.construccion.test_issue_267_a7_matriz_roles import (
    _crear_proyecto,
    _crear_usuario,
    _invalidar_cache_roles_a7,
    _seed_matriz_roles_a7,
)


def _datos_finv2_bd_caso_a10():
    return {
        'finv2_bd': {
            'rubros': {
                'Ingresos Operacionales': {'total': -100.0, 'meses': {'septiembre': -100.0}, 'cuentas': []},
                'Gastos de Personal': {'total': 40.0, 'meses': {'septiembre': 40.0}, 'cuentas': []},
                'Materiales': {'total': 10.0, 'meses': {'septiembre': 10.0}, 'cuentas': []},
            },
            'total': -50.0,
            'filas_detalle': [
                {'rubro': 'Ingresos Operacionales', 'clasificacion': 'Ingresos',
                 'valor': -100.0, 'mes': 9, 'anio': 2026, 'ciudad': 'Barranquilla',
                 'codigo_contable': '4135'},
                {'rubro': 'Gastos de Personal', 'clasificacion': 'Fijo',
                 'valor': 40.0, 'mes': 9, 'anio': 2026, 'ciudad': 'Barranquilla',
                 'codigo_contable': '5105'},
                {'rubro': 'Materiales', 'clasificacion': 'Variable',
                 'valor': 10.0, 'mes': 9, 'anio': 2026, 'ciudad': 'Cartagena',
                 'codigo_contable': '5195'},
            ],
        },
    }


def _crear_presupuesto(proyecto, anio=2026):
    return PresupuestoDetalladoConstruccion.objects.create(
        proyecto=proyecto, anio=anio,
        tipo=PresupuestoDetalladoConstruccion.Tipo.PLANEADO,
        datos=_datos_finv2_bd_caso_a10(),
    )


# ===========================================================================
# Unit — construir_contexto_reporte (sin filtro de mes, con filtro de mes)
# ===========================================================================
class ConstruirContextoReporteTests(TestCase):
    def setUp(self):
        self.proyecto = _crear_proyecto('Proyecto A10 contexto')
        _crear_presupuesto(self.proyecto)

    def test_sin_mes_trae_los_3_rubros(self):
        ctx = construir_contexto_reporte(self.proyecto, 2026)
        self.assertEqual(len(ctx['rubro_rows']), 3)
        self.assertTrue(ctx['tiene_datos'])
        semaforos = {r['rubro']: r['semaforo'] for r in ctx['rubro_rows']}
        self.assertEqual(semaforos['Ingresos Operacionales'], 'rojo')
        self.assertEqual(semaforos['Gastos de Personal'], 'amarillo')
        self.assertEqual(semaforos['Materiales'], 'verde')

    def test_con_mes_filtra_filas_detalle_y_reagrega(self):
        ctx = construir_contexto_reporte(self.proyecto, 2026, mes=9)
        self.assertEqual(len(ctx['filas_detalle']), 3)
        self.assertTrue(ctx['tiene_datos'])

    def test_mes_sin_datos_da_contexto_vacio_sin_lanzar(self):
        ctx = construir_contexto_reporte(self.proyecto, 2026, mes=1)
        self.assertFalse(ctx['tiene_datos'])
        self.assertEqual(ctx['rubro_rows'], [])

    def test_presupuesto_inexistente_da_contexto_vacio_sin_lanzar(self):
        otro_proyecto = _crear_proyecto('Proyecto A10 sin presupuesto')
        ctx = construir_contexto_reporte(otro_proyecto, 2099)
        self.assertFalse(ctx['tiene_datos'])
        self.assertIsNone(ctx['presupuesto'])


# ===========================================================================
# 1. PDF — HTML fuente (ver docstring del módulo: no se extrae del binario)
# ===========================================================================
class PdfPresupuestoTests(TestCase):
    def setUp(self):
        self.proyecto = _crear_proyecto('Proyecto A10 PDF')
        _crear_presupuesto(self.proyecto)

    def test_html_contiene_datos_del_periodo(self):
        ctx = construir_contexto_reporte(self.proyecto, 2026, mes=9)
        html = _construir_html_pdf_presupuesto(self.proyecto, 2026, 9, ctx)
        self.assertIn(self.proyecto.nombre, html)
        self.assertIn('09/2026', html)
        self.assertIn('Ingresos Operacionales', html)
        self.assertIn('Gastos de Personal', html)
        self.assertIn('Materiales', html)
        # Alertas: solo el rubro en rojo (>100%) aparece en la lista de alertas.
        self.assertIn('Alertas', html)
        self.assertIn('Ingresos Operacionales: 200.0% sobre', html)

    def test_html_sin_datos_no_lanza_y_declara_vacio(self):
        otro = _crear_proyecto('Proyecto A10 PDF sin datos')
        ctx = construir_contexto_reporte(otro, 2026)
        html = _construir_html_pdf_presupuesto(otro, 2026, None, ctx)
        self.assertIn('Sin rubros cargados para el período', html)
        self.assertIn('Ningún rubro supera el 100%', html)

    def test_generar_pdf_devuelve_binario_valido(self):
        """Smoke test del binario final -- NO se le extrae texto (ver
        docstring del módulo), solo se valida que sea un PDF real de tamaño
        no trivial (WeasyPrint corrió sin lanzar)."""
        contenido = generar_pdf_presupuesto(self.proyecto, 2026, 9)
        self.assertTrue(contenido.startswith(b'%PDF-'))
        self.assertGreater(len(contenido), 500)


# ===========================================================================
# 2. Excel — 4 hojas nombradas
# ===========================================================================
class ExcelPresupuestoTests(TestCase):
    def setUp(self):
        self.proyecto = _crear_proyecto('Proyecto A10 Excel')
        _crear_presupuesto(self.proyecto)
        HistorialCargaPresupuestoConstruccion.objects.create(
            proyecto=self.proyecto, anio=2026, mes=9,
            filas_procesadas=3, valor_total=Decimal('-50.0'),
            estado=HistorialCargaPresupuestoConstruccion.Estado.PROCESADA,
            archivo_nombre='Presupuesto_267_mes09.xlsx',
        )

    def test_4_hojas_nombradas_en_orden(self):
        from openpyxl import load_workbook

        contenido = generar_excel_presupuesto(self.proyecto, 2026)
        libro = load_workbook(io.BytesIO(contenido))
        self.assertEqual(libro.sheetnames, ['Resumen', 'Matriz', 'Rubros', 'Histórico'])

    def test_hoja_resumen_trae_kpi_correctos(self):
        from openpyxl import load_workbook

        contenido = generar_excel_presupuesto(self.proyecto, 2026)
        libro = load_workbook(io.BytesIO(contenido))
        hoja = libro['Resumen']
        valores = {fila[0].value: fila[1].value for fila in hoja.iter_rows(min_row=5) if fila[0].value}
        self.assertEqual(valores['Ingreso'], -100.0)
        self.assertEqual(valores['Costos Fijos'], 40.0)
        self.assertEqual(valores['Costos Variables'], 10.0)
        self.assertEqual(valores['Resultado'], -150.0)

    def test_hoja_matriz_tiene_encabezados_de_mes_y_semaforo(self):
        from openpyxl import load_workbook

        contenido = generar_excel_presupuesto(self.proyecto, 2026)
        libro = load_workbook(io.BytesIO(contenido))
        hoja = libro['Matriz']
        encabezados = [c.value for c in next(hoja.iter_rows(min_row=1, max_row=1))]
        self.assertEqual(encabezados[0], 'Rubro')
        self.assertIn('Julio', encabezados)
        self.assertIn('Junio', encabezados)
        self.assertEqual(encabezados[-1], 'Semáforo')

    def test_hoja_rubros_incluye_los_3_rubros_con_semaforo(self):
        from openpyxl import load_workbook

        contenido = generar_excel_presupuesto(self.proyecto, 2026)
        libro = load_workbook(io.BytesIO(contenido))
        hoja = libro['Rubros']
        filas = list(hoja.iter_rows(min_row=2, values_only=True))
        rubros = {f[0]: f[3] for f in filas}
        self.assertEqual(rubros['Ingresos Operacionales'], 'rojo')
        self.assertEqual(rubros['Gastos de Personal'], 'amarillo')
        self.assertEqual(rubros['Materiales'], 'verde')

    def test_hoja_historico_incluye_la_carga_registrada(self):
        from openpyxl import load_workbook

        contenido = generar_excel_presupuesto(self.proyecto, 2026)
        libro = load_workbook(io.BytesIO(contenido))
        hoja = libro['Histórico']
        filas = list(hoja.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(filas), 1)
        self.assertEqual(filas[0][4], 'Procesada')


# ===========================================================================
# 3. CSV — encabezado EXACTO (reuso literal de PlanoFinancieroCsvView)
# ===========================================================================
class CsvPresupuestoTests(TestCase):
    def setUp(self):
        self.proyecto = _crear_proyecto('Proyecto A10 CSV')
        _crear_presupuesto(self.proyecto)

    def test_encabezado_exacto(self):
        from django.http import HttpResponse

        response = HttpResponse(content_type='text/csv; charset=utf-8')
        escribir_csv_presupuesto(response, self.proyecto, 2026)
        contenido = response.content.decode('utf-8-sig')
        filas = list(csv.reader(io.StringIO(contenido)))
        self.assertEqual(
            filas[0],
            ['Código', 'Concepto', 'Centro', 'Proyecto', 'Mes', 'Valor', 'Referencia'],
        )

    def test_filas_agregadas_correctas(self):
        from django.http import HttpResponse

        response = HttpResponse(content_type='text/csv; charset=utf-8')
        escribir_csv_presupuesto(response, self.proyecto, 2026, mes=9)
        contenido = response.content.decode('utf-8-sig')
        filas = list(csv.reader(io.StringIO(contenido)))[1:]
        por_codigo = {f[0]: f for f in filas if f}
        self.assertIn('4135', por_codigo)
        fila_ingresos = por_codigo['4135']
        self.assertEqual(fila_ingresos[1], 'Ingresos Operacionales')
        self.assertEqual(fila_ingresos[2], 'Barranquilla')
        self.assertEqual(fila_ingresos[3], self.proyecto.nombre)
        self.assertEqual(fila_ingresos[4], '09-2026')
        self.assertEqual(fila_ingresos[6], 'Ingresos')

    def test_sin_homologar_cuando_falta_codigo_contable(self):
        from django.http import HttpResponse

        self.proyecto2 = _crear_proyecto('Proyecto A10 CSV sin codigo')
        PresupuestoDetalladoConstruccion.objects.create(
            proyecto=self.proyecto2, anio=2026,
            tipo=PresupuestoDetalladoConstruccion.Tipo.PLANEADO,
            datos={'finv2_bd': {'filas_detalle': [
                {'rubro': 'R1', 'clasificacion': 'Fijo', 'valor': 5.0,
                 'mes': 1, 'anio': 2026, 'ciudad': '', 'codigo_contable': None},
            ]}},
        )
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        escribir_csv_presupuesto(response, self.proyecto2, 2026)
        contenido = response.content.decode('utf-8-sig')
        filas = list(csv.reader(io.StringIO(contenido)))
        self.assertEqual(filas[1][0], 'SIN_HOMOLOGAR')


# ===========================================================================
# 4. PPT — exactamente 7 diapositivas
# ===========================================================================
class PptPresupuestoTests(TestCase):
    def setUp(self):
        self.proyecto = _crear_proyecto('Proyecto A10 PPT')
        _crear_presupuesto(self.proyecto)

    def test_7_diapositivas(self):
        from pptx import Presentation

        contenido = generar_ppt_presupuesto(self.proyecto, 2026, mes=9)
        prs = Presentation(io.BytesIO(contenido))
        self.assertEqual(len(prs.slides), 7)

    def test_portada_incluye_proyecto_y_periodo(self):
        from pptx import Presentation

        contenido = generar_ppt_presupuesto(self.proyecto, 2026, mes=9)
        prs = Presentation(io.BytesIO(contenido))
        portada = prs.slides[0]
        texto_portada = '\n'.join(
            shape.text_frame.text for shape in portada.shapes if shape.has_text_frame
        )
        self.assertIn(self.proyecto.nombre, texto_portada)
        self.assertIn('09/2026', texto_portada)

    def test_sin_datos_no_lanza(self):
        from pptx import Presentation

        otro = _crear_proyecto('Proyecto A10 PPT sin datos')
        contenido = generar_ppt_presupuesto(otro, 2026)
        prs = Presentation(io.BytesIO(contenido))
        self.assertEqual(len(prs.slides), 7)


# ===========================================================================
# 5. Gate de rol (A7) — endpoints reales + botones del template
# ===========================================================================
class GateFormatoReporteIntegracionTests(TestCase):
    def setUp(self):
        _seed_matriz_roles_a7()
        self.addCleanup(_invalidar_cache_roles_a7)
        self.proyecto = _crear_proyecto('Proyecto A10 gate')
        _crear_presupuesto(self.proyecto)
        self.urls = {
            'pdf': reverse('construccion:fin_presupuesto_planeado_pdf', kwargs={'proyecto_id': self.proyecto.pk}),
            'excel': reverse('construccion:fin_presupuesto_planeado_excel', kwargs={'proyecto_id': self.proyecto.pk}),
            'csv': reverse('construccion:fin_presupuesto_planeado_csv', kwargs={'proyecto_id': self.proyecto.pk}),
            'ppt': reverse('construccion:fin_presupuesto_planeado_ppt', kwargs={'proyecto_id': self.proyecto.pk}),
        }

    def _login(self, rol, suffix):
        user = _crear_usuario(rol, suffix)
        client = Client()
        client.force_login(user)
        return client

    def test_contador_solo_excel_200_resto_403(self):
        client = self._login('contador', 'a10_contador')
        self.assertEqual(client.get(self.urls['excel'], {'anio': 2026}).status_code, 200)
        for formato in ('pdf', 'csv', 'ppt'):
            with self.subTest(formato=formato):
                self.assertEqual(client.get(self.urls[formato], {'anio': 2026}).status_code, 403)

    def test_supervisor_los_4_en_403(self):
        client = self._login('supervisor', 'a10_supervisor')
        for formato, url in self.urls.items():
            with self.subTest(formato=formato):
                self.assertEqual(client.get(url, {'anio': 2026}).status_code, 403)

    def test_gerente_financiero_los_4_en_200(self):
        client = self._login('gerente_financiero', 'a10_gf')
        for formato, url in self.urls.items():
            with self.subTest(formato=formato):
                self.assertEqual(client.get(url, {'anio': 2026}).status_code, 200)

    def test_director_pdf_excel_200_ppt_csv_403(self):
        client = self._login('director', 'a10_director')
        self.assertEqual(client.get(self.urls['pdf'], {'anio': 2026}).status_code, 200)
        self.assertEqual(client.get(self.urls['excel'], {'anio': 2026}).status_code, 200)
        self.assertEqual(client.get(self.urls['ppt'], {'anio': 2026}).status_code, 403)
        self.assertEqual(client.get(self.urls['csv'], {'anio': 2026}).status_code, 403)

    def test_content_types_correctos(self):
        client = self._login('gerente_financiero', 'a10_gf_ct')
        self.assertIn('application/pdf', client.get(self.urls['pdf'], {'anio': 2026})['Content-Type'])
        self.assertIn('spreadsheetml', client.get(self.urls['excel'], {'anio': 2026})['Content-Type'])
        self.assertIn('csv', client.get(self.urls['csv'], {'anio': 2026})['Content-Type'])
        self.assertIn('presentationml', client.get(self.urls['ppt'], {'anio': 2026})['Content-Type'])


class BotonesReporteTemplateTests(TestCase):
    """Los botones de descarga en la pestaña Tabla respetan
    ``formatos_reporte_permitidos`` -- ocultar en UI no reemplaza el gate
    server-side (arriba), pero SÍ es parte del contrato de A10 (issue Fase 5:
    los botones deben existir y respetar el rol)."""

    def setUp(self):
        _seed_matriz_roles_a7()
        self.addCleanup(_invalidar_cache_roles_a7)
        self.proyecto = _crear_proyecto('Proyecto A10 botones')
        _crear_presupuesto(self.proyecto)
        self.url = reverse('construccion:fin_presupuesto_planeado', kwargs={'proyecto_id': self.proyecto.pk})

    def test_contador_solo_ve_boton_excel(self):
        user = _crear_usuario('contador', 'a10_btn_contador')
        client = Client()
        client.force_login(user)
        content = client.get(self.url, {'anio': 2026}).content.decode()
        self.assertIn('data-reporte="excel"', content)
        self.assertNotIn('data-reporte="pdf"', content)
        self.assertNotIn('data-reporte="csv"', content)
        self.assertNotIn('data-reporte="ppt"', content)

    def test_supervisor_no_ve_ningun_boton(self):
        user = _crear_usuario('supervisor', 'a10_btn_supervisor')
        client = Client()
        client.force_login(user)
        content = client.get(self.url, {'anio': 2026}).content.decode()
        self.assertNotIn('data-reportes-descargables', content)

    def test_gerente_financiero_ve_los_4_botones(self):
        user = _crear_usuario('gerente_financiero', 'a10_btn_gf')
        client = Client()
        client.force_login(user)
        content = client.get(self.url, {'anio': 2026}).content.decode()
        for formato in ('pdf', 'excel', 'csv', 'ppt'):
            self.assertIn(f'data-reporte="{formato}"', content)


# ===========================================================================
# 6. Edge case: presupuesto inexistente -> 200 con empty-state, nunca 404
# ===========================================================================
class ReportesSinDatosTests(TestCase):
    def setUp(self):
        _seed_matriz_roles_a7()
        self.addCleanup(_invalidar_cache_roles_a7)
        self.proyecto = _crear_proyecto('Proyecto A10 sin presupuesto')

    def test_los_4_formatos_responden_200_sin_presupuesto_cargado(self):
        client = self._login()
        urls = {
            'pdf': reverse('construccion:fin_presupuesto_planeado_pdf', kwargs={'proyecto_id': self.proyecto.pk}),
            'excel': reverse('construccion:fin_presupuesto_planeado_excel', kwargs={'proyecto_id': self.proyecto.pk}),
            'csv': reverse('construccion:fin_presupuesto_planeado_csv', kwargs={'proyecto_id': self.proyecto.pk}),
            'ppt': reverse('construccion:fin_presupuesto_planeado_ppt', kwargs={'proyecto_id': self.proyecto.pk}),
        }
        for formato, url in urls.items():
            with self.subTest(formato=formato):
                resp = client.get(url, {'anio': 2099})
                self.assertEqual(resp.status_code, 200)
                self.assertGreater(len(resp.content), 0)

    def _login(self):
        user = _crear_usuario('gerente_financiero', 'a10_sin_datos')
        client = Client()
        client.force_login(user)
        return client
