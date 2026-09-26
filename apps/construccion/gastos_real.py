"""Instelec#268 — Presupuesto Real (Ejecutado): carga de gastos reales y
comparación contra el Presupuesto Planeado (#267).

Dos piezas:

1. ``procesar_carga_gastos_reales``: importa el Excel de 18 columnas del
   cliente (Auxiliar … Fijo), valida cada fila (NIT contra el maestro de
   proveedores #262, Periodo AAAAMM, Neto positivo, Fecha, Docto) y — solo si
   no hay errores — reemplaza las líneas del proyecto en los períodos que trae
   el archivo (UPSERT por período: re-subir no duplica).

2. ``construir_comparativo``: agrupa lo real por Cuenta Equiv × mes y lo
   compara contra el planeado del mismo proyecto/año. El planeado guarda, bajo
   cada rubro, las mismas "cuentas equivalentes" del contable (Prestaciones
   Sociales, Seguridad Social, CIF…), así que la comparación es cuenta a
   cuenta, con subtotal por rubro (rubro = ``MapeoCtaRubro``, el mismo mapeo
   que usa el importador del planeado).

El presupuesto que se compara es el de los MESES CON EJECUCIÓN cargada (o el
período filtrado): comparar un mes real contra el presupuesto del año entero
dejaría todo en verde.
"""

from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation

from django.db import transaction

from apps.financiero.importers_finv2 import MESES_FISCALES, ContableCompleteImporter
from apps.financiero.importers_finv2_carga import _normalizar
from apps.financiero.models_finv2_facturas import Proveedor

from .models_fin import (
    GastoRealConstruccion,
    HistorialCargaPresupuestoConstruccion,
    PresupuestoDetalladoConstruccion,
)

# Encabezados exactos, en orden (issue #268, "Estructura Exacta y Completa").
COLUMNAS_GASTOS_REALES = [
    "Auxiliar",
    "Desc. auxiliar",
    "Neto",
    "Fecha",
    "Docto.",
    "Periodo",
    "Tercero movto",
    "Razón social tercero movto",
    "Desc. C.O. movto",
    "Usuario creación",
    "C.O. movto",
    "Notas",
    "C.Costo",
    "Desc. C.Costo",
    "Cuenta Equiv",
    "CdeC equiv",
    "Cargo",
    "Fijo",
]
_COLUMNAS_NORM = [_normalizar(c) for c in COLUMNAS_GASTOS_REALES]
_CAMPOS = [
    "auxiliar",
    "desc_auxiliar",
    "neto",
    "fecha",
    "docto",
    "periodo",
    "tercero_nit",
    "tercero_razon_social",
    "desc_co_movto",
    "usuario_creacion",
    "co_movto",
    "notas",
    "centro_costo",
    "desc_centro_costo",
    "cuenta_equiv",
    "cdec_equiv",
    "cargo",
    "fijo",
]

DIAS_RETROACTIVO_MAX = 30
UMBRAL_AMARILLO_PCT = Decimal("10")
MAX_ERRORES_LISTADOS = 50

CLASIFICACION_FIJO = "Fijo"
CLASIFICACION_VARIABLE = "Variable"
_VALORES_FIJO = {"si", "s", "x", "1", "fijo", "true", "verdadero", "yes"}

_MES_KEY_A_NUM = {key: num for key, _label, num in MESES_FISCALES}
MESES_CORTOS = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]
_RUBROS_INGRESO = {"ingresos operacionales", "otros ingresos", "lineas de transmision"}

_RE_PERIODO = re.compile(r"^(\d{4})(0[1-9]|1[0-2])$")


# ===========================================================================
# Utilidades de parseo
# ===========================================================================
def _texto(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def _decimal(valor):
    if valor is None or valor == "":
        return None
    if isinstance(valor, (int, float, Decimal)):
        return Decimal(str(valor))
    texto = str(valor).strip().replace("$", "").replace(" ", "")
    # "1.234.567,89" (es-CO) o "1234567.89"
    if "," in texto:
        texto = texto.replace(".", "").replace(",", ".")
    try:
        return Decimal(texto)
    except InvalidOperation:
        return None


def _fecha(valor):
    """Fecha del Excel: celda fecha, texto DD/MM/YYYY o serial de Excel."""
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    if isinstance(valor, (int, float)) and 20000 < valor < 80000:
        return date(1899, 12, 30) + timedelta(days=int(valor))
    texto = _texto(valor)
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(texto, fmt).date()
        except ValueError:
            continue
    return None


def normalizar_nit(valor) -> str:
    """NIT comparable: sin puntos, espacios ni '.0' de Excel."""
    texto = _texto(valor)
    return re.sub(r"[^0-9A-Za-z-]", "", texto)


def _nit_base(nit: str) -> str:
    """NIT sin dígito de verificación (900123456-7 → 900123456)."""
    return nit.split("-")[0]


def clasificacion_desde_fijo(valor) -> str:
    return CLASIFICACION_FIJO if _normalizar(valor) in _VALORES_FIJO else CLASIFICACION_VARIABLE


def _catalogo_proveedores():
    exacto, base = {}, {}
    for prov in Proveedor.objects.exclude(nit__isnull=True).exclude(nit=""):
        nit = normalizar_nit(prov.nit)
        exacto[nit] = prov
        base.setdefault(_nit_base(nit), prov)
    return exacto, base


def _mapeo_rubros():
    importador = ContableCompleteImporter()
    mapeo = importador._build_mapeo()
    return lambda cuenta: importador._rubro_para(cuenta, mapeo)


# ===========================================================================
# 1. Carga
# ===========================================================================
@dataclass
class ResultadoCarga:
    exito: bool
    errores: list = field(default_factory=list)
    advertencias: list = field(default_factory=list)
    filas: int = 0
    valor_total: Decimal = Decimal("0")
    periodos: list = field(default_factory=list)
    historial: HistorialCargaPresupuestoConstruccion | None = None

    @property
    def mensaje(self) -> str:
        periodos = ", ".join(self.periodos)
        total = f"{self.valor_total:,.0f}".replace(",", ".")
        return f"{self.filas} líneas de gasto real cargadas (período {periodos}), total ${total}."


def _leer_filas(archivo):
    """Devuelve (filas_con_numero, error). Busca el encabezado en las 10
    primeras filas de la primera hoja que lo tenga."""
    from openpyxl import load_workbook

    try:
        wb = load_workbook(archivo, read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 — cualquier archivo no-xlsx
        return None, "El archivo no es un Excel .xlsx válido."

    encontrado = None
    for ws in wb.worksheets:
        filas = list(ws.iter_rows(values_only=True))
        for idx, fila in enumerate(filas[:10]):
            celdas = [_normalizar(c) for c in fila]
            while celdas and not celdas[-1]:
                celdas.pop()
            if celdas[:1] == [_COLUMNAS_NORM[0]] or _COLUMNAS_NORM[2] in celdas:
                encontrado = (filas, idx, celdas)
                break
        if encontrado:
            break
    if not encontrado:
        return None, (
            "No se encontró el encabezado. La primera fila debe tener las 18 columnas: "
            + " | ".join(COLUMNAS_GASTOS_REALES)
            + "."
        )

    filas, idx, celdas = encontrado
    if celdas[:18] != _COLUMNAS_NORM:
        diferencias = []
        for pos, esperado in enumerate(COLUMNAS_GASTOS_REALES):
            real = celdas[pos] if pos < len(celdas) else ""
            if real != _COLUMNAS_NORM[pos]:
                diferencias.append(f'columna {pos + 1}: se esperaba "{esperado}"')
        return None, (
            "Los encabezados no coinciden con el formato de 18 columnas en orden ("
            + "; ".join(diferencias[:6])
            + ")."
        )

    datos = []
    for numero, fila in enumerate(filas[idx + 1 :], start=idx + 2):
        valores = list(fila[:18]) + [None] * (18 - len(fila[:18]))
        if all(v is None or _texto(v) == "" for v in valores):
            continue
        datos.append((numero, dict(zip(_CAMPOS, valores, strict=True))))
    return datos, None


def procesar_carga_gastos_reales(proyecto, archivo, usuario=None) -> ResultadoCarga:
    """Valida el archivo completo; si hay UN error no se guarda nada."""
    nombre_archivo = (getattr(archivo, "name", "") or "")[:255]
    filas, error = _leer_filas(archivo)
    resultado = ResultadoCarga(exito=False)
    if error:
        resultado.errores.append(error)
        return _registrar(resultado, proyecto, usuario, nombre_archivo)
    if not filas:
        resultado.errores.append("El archivo no tiene líneas de gasto debajo del encabezado.")
        return _registrar(resultado, proyecto, usuario, nombre_archivo)

    exacto, base = _catalogo_proveedores()
    rubro_para = _mapeo_rubros()
    lineas, vistos = [], {}
    inactivos_avisados = set()

    for numero, f in filas:
        errores_fila = []
        periodo = _texto(f["periodo"])
        m = _RE_PERIODO.match(periodo)
        if not m:
            errores_fila.append(f'Periodo "{periodo}" inválido (formato AAAAMM, ej. 202602)')
        neto = _decimal(f["neto"])
        if neto is None:
            errores_fila.append(f'Neto "{_texto(f["neto"])}" no es un número')
        elif neto <= 0:
            errores_fila.append(f"Neto {neto} debe ser positivo")
        fecha = _fecha(f["fecha"])
        if fecha is None:
            errores_fila.append(f'Fecha "{_texto(f["fecha"])}" inválida (formato DD/MM/AAAA)')
        elif m:
            inicio = date(int(m.group(1)), int(m.group(2)), 1)
            if fecha < inicio - timedelta(days=DIAS_RETROACTIVO_MAX):
                errores_fila.append(
                    f"Fecha {fecha:%d/%m/%Y} es más de {DIAS_RETROACTIVO_MAX} días anterior "
                    f"al período {periodo}"
                )

        nit = normalizar_nit(f["tercero_nit"])
        proveedor = exacto.get(nit) or base.get(_nit_base(nit)) if nit else None
        if not nit:
            errores_fila.append("Tercero movto (NIT) vacío")
        elif proveedor is None:
            errores_fila.append(
                f"Proveedor {nit} no registrado en el maestro de proveedores (#262)"
            )
        elif not proveedor.activo and proveedor.pk not in inactivos_avisados:
            inactivos_avisados.add(proveedor.pk)
            resultado.advertencias.append(f"Proveedor {proveedor.nombre} ({nit}) está inactivo.")

        docto = _texto(f["docto"])
        if m and docto:
            clave = (periodo, docto, _texto(f["auxiliar"]), nit, neto)
            if clave in vistos:
                errores_fila.append(
                    f"Docto. {docto} duplicado en el período {periodo} (misma línea que la fila {vistos[clave]})"
                )
            else:
                vistos[clave] = numero

        if errores_fila:
            resultado.errores.extend(f"Fila {numero}: {e}" for e in errores_fila)
            continue

        cuenta = _texto(f["cuenta_equiv"])
        lineas.append(
            GastoRealConstruccion(
                proyecto=proyecto,
                proveedor=proveedor,
                periodo=periodo,
                anio=int(m.group(1)),
                mes=int(m.group(2)),
                auxiliar=_texto(f["auxiliar"])[:50],
                desc_auxiliar=_texto(f["desc_auxiliar"])[:255],
                neto=neto,
                fecha=fecha,
                docto=docto[:60],
                tercero_nit=nit[:30],
                tercero_razon_social=_texto(f["tercero_razon_social"])[:255],
                desc_co_movto=_texto(f["desc_co_movto"])[:255],
                usuario_creacion=_texto(f["usuario_creacion"])[:150],
                co_movto=_texto(f["co_movto"])[:30],
                notas=_texto(f["notas"]),
                centro_costo=_texto(f["centro_costo"])[:50],
                desc_centro_costo=_texto(f["desc_centro_costo"])[:255],
                cuenta_equiv=cuenta[:150],
                cdec_equiv=_texto(f["cdec_equiv"])[:150],
                cargo=_texto(f["cargo"])[:150],
                fijo=_texto(f["fijo"])[:30],
                rubro=(rubro_para(cuenta) if cuenta else "")[:150],
                clasificacion=clasificacion_desde_fijo(f["fijo"]),
            )
        )

    if resultado.errores:
        return _registrar(resultado, proyecto, usuario, nombre_archivo)

    periodos = sorted({linea.periodo for linea in lineas})
    resultado.exito = True
    resultado.filas = len(lineas)
    resultado.valor_total = sum((linea.neto for linea in lineas), Decimal("0"))
    resultado.periodos = periodos
    with transaction.atomic():
        _registrar(resultado, proyecto, usuario, nombre_archivo)
        # UPSERT por período: se reemplaza el período completo, nunca se duplica.
        GastoRealConstruccion.objects.filter(proyecto=proyecto, periodo__in=periodos).delete()
        for linea in lineas:
            linea.carga = resultado.historial
        GastoRealConstruccion.objects.bulk_create(lineas, batch_size=500)
    return resultado


def _registrar(resultado, proyecto, usuario, nombre_archivo):
    anio = mes = None
    if len(resultado.periodos) == 1:
        anio, mes = int(resultado.periodos[0][:4]), int(resultado.periodos[0][4:])
    elif resultado.periodos:
        anio = int(resultado.periodos[0][:4])
    Estado = HistorialCargaPresupuestoConstruccion.Estado
    resultado.historial = HistorialCargaPresupuestoConstruccion.objects.create(
        proyecto=proyecto,
        anio=anio,
        mes=mes,
        usuario=usuario if getattr(usuario, "is_authenticated", False) else None,
        filas_procesadas=resultado.filas,
        valor_total=resultado.valor_total,
        estado=Estado.PROCESADA if resultado.exito else Estado.ERROR,
        archivo_nombre=nombre_archivo,
        tipo=PresupuestoDetalladoConstruccion.Tipo.REAL,
        detalle_errores={}
        if resultado.exito
        else {
            "errores": resultado.errores[:MAX_ERRORES_LISTADOS],
            "total_errores": len(resultado.errores),
            "advertencias": resultado.advertencias,
        },
    )
    return resultado


# ===========================================================================
# 2. Comparativo real vs planeado
# ===========================================================================
def semaforo(real, planeado) -> str:
    """🟢 real < planeado · 🟡 0–10 % sobre · 🔴 > 10 % sobre · sin presupuesto."""
    if planeado is None or planeado == 0:
        return "sin_presupuesto"
    if real < planeado:
        return "verde"
    if variacion_pct(real, planeado) <= UMBRAL_AMARILLO_PCT:
        return "amarillo"
    return "rojo"


def variacion_pct(real, planeado):
    if not planeado:
        return None
    return ((real - planeado) / abs(planeado) * Decimal("100")).quantize(Decimal("0.01"))


def _es_ingreso(nombre) -> bool:
    n = _normalizar(nombre)
    return n in _RUBROS_INGRESO or n.startswith("ingreso")


def planeado_por_mes(proyecto, anio):
    """Planeado del año en dos niveles: {cuenta_norm: {...}} y {rubro_norm: {...}},
    cada uno con ``meses`` {1..12: Decimal}. Excluye los rubros de ingreso."""
    presupuesto = PresupuestoDetalladoConstruccion.objects.filter(
        proyecto=proyecto,
        anio=anio,
        tipo=PresupuestoDetalladoConstruccion.Tipo.PLANEADO,
    ).first()
    por_cuenta, por_rubro = {}, {}
    rubros = (((presupuesto.datos if presupuesto else None) or {}).get("finv2_bd") or {}).get(
        "rubros"
    ) or {}
    for rubro, info in rubros.items():
        if _es_ingreso(rubro) or not isinstance(info, dict):
            continue
        meses = defaultdict(Decimal)
        for key, valor in (info.get("meses") or {}).items():
            if key in _MES_KEY_A_NUM:
                meses[_MES_KEY_A_NUM[key]] += _decimal(valor) or Decimal("0")
        por_rubro[_normalizar(rubro)] = {"nombre": rubro, "meses": meses}
        for cuenta in info.get("cuentas") or []:
            nombre = (cuenta or {}).get("cta_equivalente")
            if not nombre:
                continue
            item = por_cuenta.setdefault(
                _normalizar(nombre),
                {"nombre": nombre, "rubro": rubro, "meses": defaultdict(Decimal)},
            )
            for key, valor in (cuenta.get("meses") or {}).items():
                if key in _MES_KEY_A_NUM:
                    item["meses"][_MES_KEY_A_NUM[key]] += _decimal(valor) or Decimal("0")
    return por_cuenta, por_rubro


def filtrar_gastos(proyecto, anio, filtros):
    qs = GastoRealConstruccion.objects.filter(proyecto=proyecto, anio=anio)
    if filtros.get("clasificacion"):
        qs = qs.filter(clasificacion=filtros["clasificacion"])
    if filtros.get("proveedor"):
        qs = qs.filter(tercero_nit=filtros["proveedor"])
    if filtros.get("centro_costo"):
        qs = qs.filter(centro_costo=filtros["centro_costo"])
    if filtros.get("periodo"):
        qs = qs.filter(periodo=filtros["periodo"])
    return qs


def _fila(nombre, rubro, meses_real, plan, es_subtotal=False):
    total = sum(meses_real.values(), Decimal("0"))
    variacion = None if plan is None else total - plan
    return {
        "nombre": nombre,
        "rubro": rubro,
        "es_subtotal": es_subtotal,
        "meses": [meses_real.get(m, Decimal("0")) for m in range(1, 13)],
        "total_real": total,
        "presupuesto": plan,
        "variacion": variacion,
        "variacion_pct": variacion_pct(total, plan) if plan else None,
        "semaforo": semaforo(total, plan) if plan is not None else "sin_presupuesto",
    }


def construir_comparativo(proyecto, anio, filtros=None) -> dict:
    """Matriz Gastos Reales × Meses con Total Real / Presupuesto / Variación,
    subtotales por rubro, KPI cards y relación por proveedor."""
    filtros = filtros or {}
    gastos = list(filtrar_gastos(proyecto, anio, filtros).select_related("proveedor"))
    meses_con_ejecucion = sorted({g.mes for g in gastos})
    if filtros.get("periodo"):
        meses_alcance = [int(filtros["periodo"][4:])]
    else:
        meses_alcance = meses_con_ejecucion

    plan_cuenta, plan_rubro = planeado_por_mes(proyecto, anio)

    def plan_de(item):
        if item is None:
            return None
        return sum((item["meses"].get(m, Decimal("0")) for m in meses_alcance), Decimal("0"))

    real_cuenta = defaultdict(lambda: defaultdict(Decimal))
    rubro_de_cuenta = {}
    for g in gastos:
        clave = _normalizar(g.cuenta_equiv) or "(sin cuenta equiv)"
        real_cuenta[clave][g.mes] += g.neto
        rubro_de_cuenta.setdefault(clave, (g.cuenta_equiv or "(sin cuenta equiv)", g.rubro))

    # Agrupar cuentas por rubro (el rubro del planeado manda si la cuenta existe allí).
    por_rubro = defaultdict(list)
    for clave, (nombre, rubro_real) in rubro_de_cuenta.items():
        rubro = (
            (plan_cuenta.get(clave) or {}).get("rubro") or rubro_real or "Otros / No Clasificado"
        )
        por_rubro[rubro].append(
            _fila(nombre, rubro, real_cuenta[clave], plan_de(plan_cuenta.get(clave)))
        )

    filas = []
    total_real = total_plan = Decimal("0")
    for rubro in sorted(por_rubro, key=_normalizar):
        cuentas = sorted(por_rubro[rubro], key=lambda f: _normalizar(f["nombre"]))
        meses_rubro = defaultdict(Decimal)
        for f in cuentas:
            for idx, valor in enumerate(f["meses"], start=1):
                meses_rubro[idx] += valor
        plan_r = plan_de(plan_rubro.get(_normalizar(rubro)))
        filas.append(_fila(rubro, rubro, meses_rubro, plan_r, es_subtotal=True))
        filas.extend(cuentas)
        total_real += sum(meses_rubro.values(), Decimal("0"))
        total_plan += plan_r or Decimal("0")

    # Rubros planeados sin ejecución en el alcance: también se comparan (real 0).
    rubros_con_real = {_normalizar(r) for r in por_rubro}
    if meses_alcance and not any(
        filtros.get(k) for k in ("clasificacion", "proveedor", "centro_costo")
    ):
        for clave, item in sorted(plan_rubro.items()):
            if clave in rubros_con_real:
                continue
            plan_r = plan_de(item)
            if plan_r:
                filas.append(_fila(item["nombre"], item["nombre"], {}, plan_r, es_subtotal=True))
                total_plan += plan_r

    totales_mes = [
        sum((f["meses"][i] for f in filas if f["es_subtotal"]), Decimal("0")) for i in range(12)
    ]
    variacion_total = total_real - total_plan
    kpi = {
        "planeado": total_plan,
        "real": total_real,
        "variacion": variacion_total,
        "variacion_pct": variacion_pct(total_real, total_plan),
        "cumplimiento_pct": (total_real / total_plan * Decimal("100")).quantize(Decimal("0.1"))
        if total_plan
        else None,
        "semaforo": semaforo(total_real, total_plan),
    }

    # Relación con el maestro de proveedores (#262).
    acumulado_anio = defaultdict(Decimal)
    for nit, neto in GastoRealConstruccion.objects.filter(proyecto=proyecto, anio=anio).values_list(
        "tercero_nit", "neto"
    ):
        acumulado_anio[nit] += neto
    proveedores = {}
    for g in gastos:
        item = proveedores.setdefault(
            g.tercero_nit,
            {
                "nit": g.tercero_nit,
                "proveedor": g.proveedor,
                "nombre": g.proveedor.nombre if g.proveedor else g.tercero_razon_social,
                "activo": g.proveedor.activo if g.proveedor else None,
                "total": Decimal("0"),
                "lineas": 0,
            },
        )
        item["total"] += g.neto
        item["lineas"] += 1
    for item in proveedores.values():
        item["total_anio"] = acumulado_anio.get(item["nit"], Decimal("0"))

    return {
        "filas": filas,
        "totales_mes": totales_mes,
        "total_real": total_real,
        "total_presupuesto": total_plan,
        "kpi": kpi,
        "meses_alcance": meses_alcance,
        "meses_alcance_label": ", ".join(MESES_CORTOS[m - 1] for m in meses_alcance),
        "proveedores": sorted(proveedores.values(), key=lambda p: -p["total"]),
        "n_lineas": len(gastos),
        "tiene_planeado": bool(plan_rubro),
    }


def opciones_filtros(proyecto, anio) -> dict:
    qs = GastoRealConstruccion.objects.filter(proyecto=proyecto, anio=anio)
    proveedores = {}
    for nit, razon, nombre in qs.values_list(
        "tercero_nit", "tercero_razon_social", "proveedor__nombre"
    ):
        proveedores.setdefault(nit, nombre or razon or nit)
    centros = {}
    for cc, desc in qs.values_list("centro_costo", "desc_centro_costo"):
        if cc:
            centros.setdefault(cc, desc)
    return {
        "periodos": sorted(set(qs.values_list("periodo", flat=True))),
        "proveedores": sorted(proveedores.items(), key=lambda p: _normalizar(p[1])),
        "centros_costo": sorted(centros.items()),
        "clasificaciones": [CLASIFICACION_FIJO, CLASIFICACION_VARIABLE],
    }
