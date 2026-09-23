"""Instelec#267 — A2: Importador plano

``Tipo|Proyecto|Rubro|Clasificacion|Valor|mes|año|ciudad`` + validación dura
+ extensión de esquema (Fase 1.1/1.2 del issue).

Dependencia de A1 (commit 0735a8b, ya en esta branch): el POST de
``PresupuestoPlaneadoConstruccionView`` mergea ``datos['finv2_bd']`` vía
``_merge_presupuesto_datos``/``_merge_finv2_bd`` en vez de reemplazarlo. Este
importador NO escribe la BD directo: devuelve ``datos`` para que ese mismo
merge lo persista (precondición explícita de A1: reusar el merge, no volver
a hacer ``update()`` a nivel raíz).

Cobertura:
- Unit (sin BD real más que Homologación, vía ``pytest``-less
  ``django.test.TestCase``): archivo válido con ≥2 rubros del ejemplo
  LITERAL del issue (#267, Fase 1.1), encabezado faltante, Rubro no
  homologado (rechazo total con fila+rubro), Clasificacion inválida, mes/año
  fuera de rango, Tipo≠Presupuesto, Valor no numérico.
- Detector de formato: ``detect_excel_format_construccion`` reconoce el
  formato plano (columnas exactas, SIN columnas de nombre de mes) antes que
  el formato 'presupuesto' legacy.
- Integración (POST real): 2 cargas de meses distintos del formato plano
  coexisten en ``finv2_bd`` (mismo contrato UPSERT que A1); recarga del mismo
  mes reemplaza sin duplicar; una carga contable (BD, con ``cuentas``) previa
  no se pierde al llegar una carga plana de OTRO rubro (interacción nueva:
  ``_merge_finv2_bd`` fue extendido en A2 para reconciliar rubros SIN
  ``cuentas``, ver ``apps/construccion/views_fin.py``).
"""
import datetime
import io
import uuid

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import Client, TestCase
from django.urls import reverse
from openpyxl import Workbook

from apps.construccion.importers import (
    PresupuestoPlanoConstruccionExcelImporter,
    detect_excel_format_construccion,
)
from apps.construccion.models import ProyectoConstruccion
from apps.construccion.models_fin import PresupuestoDetalladoConstruccion
from apps.contratos.models import Contrato
from apps.financiero.models_finv2_carga import HomologacionProjectsContable

User = get_user_model()


# ===========================================================================
# Helpers
# ===========================================================================
_ENCABEZADOS_PLANO = ['Tipo', 'Proyecto', 'Rubro', 'Clasificacion', 'Valor', 'mes', 'año', 'ciudad']


def _xlsx_plano_bytes(filas, encabezados=None):
    """Excel plano en memoria: fila 1 encabezados, resto ``filas`` (list[list])."""
    wb = Workbook()
    ws = wb.active
    ws.append(encabezados or _ENCABEZADOS_PLANO)
    for fila in filas:
        ws.append(fila)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


class _FakeUpload:
    """Mínimo file-like con .name y .size para el importer (sin storage real)."""

    def __init__(self, buf, name='presupuesto.xlsx'):
        self._buf = buf
        self.name = name
        buf.seek(0, io.SEEK_END)
        self.size = buf.tell()
        buf.seek(0)

    def __getattr__(self, item):
        return getattr(self._buf, item)


def _homologacion(rubro, clasificacion, codigo='5100', tipo='PRESUPUESTO'):
    return HomologacionProjectsContable.objects.create(
        tipo=tipo, grupo=clasificacion, concepto=rubro, rubro=rubro,
        codigo_contable=codigo, activo=True,
    )


# Filas literales del ejemplo del issue (#267 Fase 1.1).
FILA_INGRESOS = ['Presupuesto', 'Transelca', 'Ingresos Operacionales', 'Ingresos', -463811660, 1, 2026, 'Barranquilla']
FILA_PERSONAL = ['Presupuesto', 'Transelca', 'Gastos de Personal', 'Fijo', 25625995.04, 1, 2026, 'Barranquilla']
FILA_PARAFISCALES = ['Presupuesto', 'Transelca', 'Aportes Parafiscales', 'Fijo', 8663360.61, 1, 2026, 'Barranquilla']


# ===========================================================================
# Unit — PresupuestoPlanoConstruccionExcelImporter (con BD solo para Homologación)
# ===========================================================================
class PresupuestoPlanoImporterTests(TestCase):

    def setUp(self):
        _homologacion('Ingresos Operacionales', 'Ingresos', codigo='4100')
        _homologacion('Gastos de Personal', 'Fijo', codigo='5105')
        _homologacion('Aportes Parafiscales', 'Fijo', codigo='5110')

    def test_archivo_valido_dos_rubros_del_ejemplo_del_issue(self):
        archivo = _FakeUpload(_xlsx_plano_bytes([FILA_INGRESOS, FILA_PERSONAL]))
        res = PresupuestoPlanoConstruccionExcelImporter().procesar(archivo)

        self.assertTrue(res['exito'], res.get('error'))
        rubros = res['datos']['finv2_bd']['rubros']
        self.assertEqual(set(rubros), {'Ingresos Operacionales', 'Gastos de Personal'})
        # mes=1 → key fiscal 'enero' (MESES_FISCALES, apps/financiero/importers_finv2.py)
        self.assertEqual(rubros['Ingresos Operacionales']['meses'].get('enero'), -463811660.0)
        self.assertEqual(rubros['Gastos de Personal']['meses'].get('enero'), 25625995.04)
        self.assertEqual(res['filas'], 2)

        filas_detalle = res['datos']['finv2_bd']['filas_detalle']
        self.assertEqual(len(filas_detalle), 2)
        fila_personal = next(f for f in filas_detalle if f['rubro'] == 'Gastos de Personal')
        self.assertEqual(fila_personal['clasificacion'], 'Fijo')
        self.assertEqual(fila_personal['ciudad'], 'Barranquilla')
        self.assertEqual(fila_personal['codigo_contable'], '5105')
        self.assertEqual(fila_personal['anio'], 2026)
        self.assertEqual(fila_personal['mes'], 1)

    def test_archivo_valido_tres_rubros_no_pierde_ninguno(self):
        archivo = _FakeUpload(
            _xlsx_plano_bytes([FILA_INGRESOS, FILA_PERSONAL, FILA_PARAFISCALES]))
        res = PresupuestoPlanoConstruccionExcelImporter().procesar(archivo)
        self.assertTrue(res['exito'], res.get('error'))
        self.assertEqual(res['filas'], 3)
        self.assertEqual(len(res['datos']['finv2_bd']['rubros']), 3)

    def test_encabezado_faltante_error_explicito(self):
        encabezados_incompletos = ['Tipo', 'Proyecto', 'Rubro', 'Valor', 'mes', 'año']  # falta Clasificacion
        archivo = _FakeUpload(_xlsx_plano_bytes([], encabezados=encabezados_incompletos))
        res = PresupuestoPlanoConstruccionExcelImporter().procesar(archivo)
        self.assertFalse(res['exito'])
        self.assertIn('clasificacion', res['error'].lower())

    def test_rubro_no_homologado_rechaza_archivo_completo_con_fila_y_rubro(self):
        fila_no_homologada = ['Presupuesto', 'Transelca', 'Rubro Inexistente', 'Fijo', 1000, 1, 2026, 'Bogota']
        archivo = _FakeUpload(_xlsx_plano_bytes([FILA_PERSONAL, fila_no_homologada]))
        res = PresupuestoPlanoConstruccionExcelImporter().procesar(archivo)

        self.assertFalse(res['exito'])
        self.assertIn('fila 3', res['error'])
        self.assertIn('Rubro Inexistente', res['error'])
        self.assertIsNone(res['datos'], 'rechazo total: ni siquiera la fila válida se persiste')

    def test_clasificacion_invalida_error(self):
        fila_invalida = ['Presupuesto', 'Transelca', 'Gastos de Personal', 'Otra', 1000, 1, 2026, 'Bogota']
        archivo = _FakeUpload(_xlsx_plano_bytes([fila_invalida]))
        res = PresupuestoPlanoConstruccionExcelImporter().procesar(archivo)
        self.assertFalse(res['exito'])
        self.assertIn('Clasificacion', res['error'])

    def test_clasificacion_ingresos_es_valida_pese_a_fase_1_2(self):
        """Decisión del PLAN (#267 hallazgo 6): superset {Ingresos,Fijo,Variable}."""
        archivo = _FakeUpload(_xlsx_plano_bytes([FILA_INGRESOS]))
        res = PresupuestoPlanoConstruccionExcelImporter().procesar(archivo)
        self.assertTrue(res['exito'], res.get('error'))

    def test_mes_fuera_de_rango_error(self):
        fila_mes_invalido = ['Presupuesto', 'Transelca', 'Gastos de Personal', 'Fijo', 1000, 13, 2026, 'Bogota']
        archivo = _FakeUpload(_xlsx_plano_bytes([fila_mes_invalido]))
        res = PresupuestoPlanoConstruccionExcelImporter().procesar(archivo)
        self.assertFalse(res['exito'])

    def test_anio_fuera_de_rango_error(self):
        fila_anio_invalido = ['Presupuesto', 'Transelca', 'Gastos de Personal', 'Fijo', 1000, 1, 202, 'Bogota']
        archivo = _FakeUpload(_xlsx_plano_bytes([fila_anio_invalido]))
        res = PresupuestoPlanoConstruccionExcelImporter().procesar(archivo)
        self.assertFalse(res['exito'])
        self.assertIn('año', res['error'].lower())

    def test_tipo_distinto_a_presupuesto_error(self):
        fila_tipo_invalido = ['Real', 'Transelca', 'Gastos de Personal', 'Fijo', 1000, 1, 2026, 'Bogota']
        archivo = _FakeUpload(_xlsx_plano_bytes([fila_tipo_invalido]))
        res = PresupuestoPlanoConstruccionExcelImporter().procesar(archivo)
        self.assertFalse(res['exito'])
        self.assertIn('Tipo', res['error'])

    def test_valor_no_numerico_error(self):
        fila_valor_invalido = ['Presupuesto', 'Transelca', 'Gastos de Personal', 'Fijo', 'abc', 1, 2026, 'Bogota']
        archivo = _FakeUpload(_xlsx_plano_bytes([fila_valor_invalido]))
        res = PresupuestoPlanoConstruccionExcelImporter().procesar(archivo)
        self.assertFalse(res['exito'])

    def test_archivo_sin_filas_de_datos_devuelve_advertencia(self):
        archivo = _FakeUpload(_xlsx_plano_bytes([]))
        res = PresupuestoPlanoConstruccionExcelImporter().procesar(archivo)
        self.assertFalse(res['exito'])
        self.assertIsNotNone(res['advertencia'])

    def test_archivo_no_xlsx_error(self):
        buf = io.BytesIO(b'no es un excel')
        archivo = _FakeUpload(buf, name='presupuesto.csv')
        res = PresupuestoPlanoConstruccionExcelImporter().procesar(archivo)
        self.assertFalse(res['exito'])


# ===========================================================================
# Detector de formato
# ===========================================================================
class DetectorFormatoPlanoTests(TestCase):

    def test_detecta_presupuesto_plano_por_encabezados_exactos(self):
        archivo = _FakeUpload(_xlsx_plano_bytes([FILA_PERSONAL]))
        self.assertEqual(detect_excel_format_construccion(archivo), 'presupuesto_plano')

    def test_no_confunde_plano_con_presupuesto_legacy_de_meses(self):
        # El legacy usa columnas ENERO..DICIEMBRE; el plano usa una columna 'mes'.
        wb_legacy = Workbook()
        ws = wb_legacy.active
        ws.append(['Concepto', 'Enero', 'Febrero'])
        ws.append(['Arriendo', 100, 200])
        buf = io.BytesIO()
        wb_legacy.save(buf)
        buf.seek(0)
        archivo = _FakeUpload(buf)
        self.assertEqual(detect_excel_format_construccion(archivo), 'presupuesto')


# ===========================================================================
# Integración — POST real (UPSERT del formato plano + interacción con A1)
# ===========================================================================
def _crear_proyecto(nombre='Proyecto A2 test #267'):
    contrato = Contrato.objects.create(
        codigo=f"CONS-{uuid.uuid4().hex[:10]}",
        nombre='Contrato test construcción #267 A2',
        unidad_negocio='CONSTRUCCION',
    )
    return ProyectoConstruccion.objects.create(contrato=contrato, nombre=nombre)


def _crear_usuario_admin():
    try:
        return User.objects.create_superuser(
            username=f'qa_267a2_{uuid.uuid4().hex[:6]}',
            email=f'qa_267a2_{uuid.uuid4().hex[:6]}@instelec.com', password='x',
        )
    except TypeError:
        user = User(is_superuser=True, is_staff=True)
        if hasattr(user, 'email'):
            user.email = f'qa_267a2_{uuid.uuid4().hex[:6]}@instelec.com'
        user.set_password('x')
        user.save()
        return user


class PresupuestoPlanoPostUpsertTests(TestCase):
    """B4 POST real: 2 cargas del formato plano + interacción con carga contable."""

    def setUp(self):
        self.proyecto = _crear_proyecto()
        self.user = _crear_usuario_admin()
        self.client = Client()
        self.client.force_login(self.user)
        self.url = reverse('construccion:fin_presupuesto_planeado',
                            kwargs={'proyecto_id': self.proyecto.pk})
        _homologacion('Gastos de Personal', 'Fijo', codigo='5105')

    def _post_plano(self, mes, valor, anio=2026, rubro='Gastos de Personal', clasificacion='Fijo'):
        fila = ['Presupuesto', 'Transelca', rubro, clasificacion, valor, mes, anio, 'Barranquilla']
        buf = _xlsx_plano_bytes([fila])
        archivo = SimpleUploadedFile(
            'plano.xlsx', buf.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        return self.client.post(
            self.url, {'action': 'cargar_bd', 'anio': anio, 'archivo': archivo})

    def _obj(self, anio=2026):
        return PresupuestoDetalladoConstruccion.objects.get(
            proyecto=self.proyecto, anio=anio,
            tipo=PresupuestoDetalladoConstruccion.Tipo.PLANEADO,
        )

    def test_dos_meses_formato_plano_coexisten(self):
        resp1 = self._post_plano(mes=1, valor=100)
        self.assertEqual(resp1.status_code, 302)
        resp2 = self._post_plano(mes=2, valor=200)
        self.assertEqual(resp2.status_code, 302)

        meses = self._obj().datos['finv2_bd']['rubros']['Gastos de Personal']['meses']
        self.assertEqual(meses.get('enero'), 100.0, 'enero no debe desaparecer al cargar febrero')
        self.assertEqual(meses.get('febrero'), 200.0)
        filas_detalle = self._obj().datos['finv2_bd']['filas_detalle']
        self.assertEqual(len(filas_detalle), 2)

    def test_recarga_mismo_mes_formato_plano_reemplaza_no_duplica(self):
        self._post_plano(mes=1, valor=100)
        self._post_plano(mes=1, valor=999)

        obj = self._obj()
        meses = obj.datos['finv2_bd']['rubros']['Gastos de Personal']['meses']
        self.assertEqual(meses, {'enero': 999.0}, 'recarga del mismo mes reemplaza, no acumula')
        filas_detalle = obj.datos['finv2_bd']['filas_detalle']
        self.assertEqual(len(filas_detalle), 1, 'no debe duplicar la fila del mismo período')
        self.assertEqual(filas_detalle[0]['valor'], 999.0)

    def test_carga_contable_previa_no_se_pierde_con_carga_plana_de_otro_rubro(self):
        """Interacción nueva de A2: ``_merge_finv2_bd`` reconcilia rubros SIN
        ``cuentas`` (formato plano) sin pisar rubros CON ``cuentas`` (BD
        contable) ya persistidos — ver docstring de ``_merge_finv2_bd``."""
        wb = Workbook()
        ws = wb.active
        ws.title = 'BD'
        ws.append(['Desc auxiliar', 'Neto', 'Cta equivalente', 'Fecha'])
        ws.append(['mov', 500, 'CTA-A2', datetime.date(2026, 7, 15)])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        archivo_contable = SimpleUploadedFile(
            'BASE DE DATOS.xlsx', buf.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        resp_contable = self.client.post(
            self.url, {'action': 'cargar_bd', 'anio': 2026, 'archivo': archivo_contable})
        self.assertEqual(resp_contable.status_code, 302)

        rubros_tras_contable = self._obj().datos['finv2_bd']['rubros']
        total_tras_contable = self._obj().datos['finv2_bd']['total']
        self.assertGreater(total_tras_contable, 0)

        resp_plano = self._post_plano(mes=1, valor=100)
        self.assertEqual(resp_plano.status_code, 302)

        obj = self._obj()
        rubros_finales = obj.datos['finv2_bd']['rubros']
        # Los rubros que ya existían por la carga contable siguen íntegros.
        for rubro, info in rubros_tras_contable.items():
            self.assertIn(rubro, rubros_finales)
            self.assertEqual(rubros_finales[rubro]['cuentas'], info['cuentas'])
        # El rubro nuevo del formato plano coexiste, sin 'cuentas'.
        self.assertIn('Gastos de Personal', rubros_finales)
        self.assertEqual(rubros_finales['Gastos de Personal']['cuentas'], [])
        self.assertEqual(
            rubros_finales['Gastos de Personal']['meses'].get('enero'), 100.0)
