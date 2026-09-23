"""Instelec#267 — A6: Plantilla XLSX descargable (formato plano de presupuesto).

``DescargarPlantillaPresupuestoPlanoView`` (apps/financiero/views.py, wireada
desde apps/construccion/urls_fin.py como ``construccion:fin_plantilla_presupuesto_plano``)
genera un .xlsx con los encabezados EXACTOS que consume
``PresupuestoPlanoConstruccionExcelImporter`` (A2):
``Tipo|Proyecto|Rubro|Clasificacion|Valor|mes|año|ciudad``, 10 filas de
ejemplo y una hoja "Instrucciones".

Cobertura (tests_requeridos de A6 + edge cases del dominio):
- El archivo generado abre con openpyxl y trae la hoja PRESUPUESTO.
- La hoja Instrucciones está presente y no vacía, y documenta las 3
  clasificaciones válidas (Ingresos/Fijo/Variable).
- Las columnas del header son exactas (mismo contrato que A2).
- 10 filas de ejemplo.
- Edge case 1 (catálogo poblado): cuando hay Homologación activa tipo
  PRESUPUESTO, los ejemplos generados reusan Rubros reales y el archivo
  descargado, re-subido tal cual, PASA la validación dura del importador A2
  (round-trip real, no solo estructura).
- Edge case 2 (catálogo vacío): sin Homologación activa, los ejemplos caen a
  placeholders marcados "EJEMPLO -" y la hoja Instrucciones avisa
  explícitamente que hay que reemplazarlos (documentado como riesgo
  operativo heredado de A2 — "la mayoría de cargas reales será rechazada
  hasta poblar el catálogo").
- Edge case 3 (permisos): un usuario anónimo no descarga la plantilla
  (redirige a login), consistente con LoginRequiredMixin.
- La vista no revienta si se invoca sin ``proyecto_id`` en los kwargs
  (fallback a nombre de proyecto genérico) — defensivo, aunque la URL real
  siempre lo trae.
"""
import io
import uuid

from django.contrib.auth import get_user_model
from django.test import Client, RequestFactory, TestCase
from django.urls import reverse
from openpyxl import load_workbook

from apps.construccion.importers import PresupuestoPlanoConstruccionExcelImporter
from apps.construccion.models import ProyectoConstruccion
from apps.contratos.models import Contrato
from apps.financiero.models_finv2_carga import HomologacionProjectsContable
from apps.financiero.views import DescargarPlantillaPresupuestoPlanoView

User = get_user_model()

_HEADERS_ESPERADOS = ['Tipo', 'Proyecto', 'Rubro', 'Clasificacion', 'Valor', 'mes', 'año', 'ciudad']


class _FakeUpload:
    """Mínimo file-like con .name y .size, espejo del helper de A2."""

    def __init__(self, buf, name='plantilla.xlsx'):
        self._buf = buf
        self.name = name
        buf.seek(0, io.SEEK_END)
        self.size = buf.tell()
        buf.seek(0)

    def __getattr__(self, item):
        return getattr(self._buf, item)


def _crear_proyecto(nombre='Proyecto A6 test #267'):
    contrato = Contrato.objects.create(
        codigo=f"CONS-{uuid.uuid4().hex[:10]}",
        nombre='Contrato test construcción #267 A6',
        unidad_negocio='CONSTRUCCION',
    )
    return ProyectoConstruccion.objects.create(contrato=contrato, nombre=nombre)


def _crear_usuario_admin():
    try:
        return User.objects.create_superuser(
            username=f'qa_267a6_{uuid.uuid4().hex[:6]}',
            email=f'qa_267a6_{uuid.uuid4().hex[:6]}@instelec.com', password='x',
        )
    except TypeError:
        user = User(is_superuser=True, is_staff=True)
        if hasattr(user, 'email'):
            user.email = f'qa_267a6_{uuid.uuid4().hex[:6]}@instelec.com'
        user.set_password('x')
        user.save()
        return user


def _homologacion(rubro, clasificacion, codigo='5100', tipo='PRESUPUESTO'):
    return HomologacionProjectsContable.objects.create(
        tipo=tipo, grupo=clasificacion, concepto=rubro, rubro=rubro,
        codigo_contable=codigo, activo=True,
    )


class DescargarPlantillaPresupuestoPlanoViewTests(TestCase):

    def setUp(self):
        self.proyecto = _crear_proyecto()
        self.user = _crear_usuario_admin()
        self.client = Client()
        self.client.force_login(self.user)
        self.url = reverse(
            'construccion:fin_plantilla_presupuesto_plano',
            kwargs={'proyecto_id': self.proyecto.pk},
        )

    def _descargar_wb(self, anio=None):
        url = self.url if anio is None else f'{self.url}?anio={anio}'
        resp = self.client.get(url)
        self.assertEqual(resp.status_code, 200)
        wb = load_workbook(io.BytesIO(resp.content))
        return resp, wb

    # -- estructura básica (tests_requeridos de A6) --------------------

    def test_archivo_abre_con_openpyxl_y_tiene_hoja_presupuesto(self):
        _resp, wb = self._descargar_wb()
        self.assertIn('PRESUPUESTO', wb.sheetnames)

    def test_hoja_instrucciones_presente_y_no_vacia(self):
        _resp, wb = self._descargar_wb()
        self.assertIn('Instrucciones', wb.sheetnames)
        ws = wb['Instrucciones']
        textos = [c.value for row in ws.iter_rows() for c in row if c.value]
        contenido = ' '.join(str(t) for t in textos)
        self.assertGreater(len(textos), 5)
        # Documenta explícitamente las 3 clasificaciones válidas.
        self.assertIn('Ingresos', contenido)
        self.assertIn('Fijo', contenido)
        self.assertIn('Variable', contenido)

    def test_columnas_exactas_en_header_row(self):
        _resp, wb = self._descargar_wb()
        ws = wb['PRESUPUESTO']
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        self.assertEqual(header, _HEADERS_ESPERADOS)

    def test_diez_filas_de_ejemplo(self):
        _resp, wb = self._descargar_wb()
        ws = wb['PRESUPUESTO']
        filas = [
            row for row in ws.iter_rows(min_row=2, values_only=True)
            if any(v is not None for v in row)
        ]
        self.assertEqual(len(filas), 10)

    def test_filename_incluye_anio(self):
        resp, _wb = self._descargar_wb(anio=2027)
        self.assertIn('2027', resp['Content-Disposition'])
        self.assertIn('.xlsx', resp['Content-Disposition'])

    # -- edge case 1: catálogo poblado → round-trip real contra A2 -----

    def test_con_catalogo_poblado_el_archivo_generado_pasa_el_importador_real(self):
        _homologacion('Gastos de Personal', 'Fijo', codigo='5105')
        _homologacion('Ingresos Operacionales', 'Ingresos', codigo='4100')

        resp, wb = self._descargar_wb()
        ws = wb['PRESUPUESTO']
        rubros_usados = {
            row[2] for row in ws.iter_rows(min_row=2, values_only=True) if row[2]
        }
        # Los ejemplos deben citar Rubros REALES del catálogo, no placeholders.
        self.assertTrue(rubros_usados & {'Gastos de Personal', 'Ingresos Operacionales'})
        self.assertFalse(any(r.startswith('EJEMPLO -') for r in rubros_usados))

        archivo = _FakeUpload(io.BytesIO(resp.content))
        res = PresupuestoPlanoConstruccionExcelImporter().procesar(archivo)
        self.assertTrue(res['exito'], res.get('error'))

    # -- edge case 2: catálogo vacío → placeholders explícitos ----------

    def test_sin_catalogo_usa_placeholders_marcados_y_avisa_en_instrucciones(self):
        self.assertFalse(
            HomologacionProjectsContable.objects.filter(
                tipo__iexact='PRESUPUESTO', activo=True,
            ).exists()
        )
        _resp, wb = self._descargar_wb()
        ws = wb['PRESUPUESTO']
        rubros_usados = [
            row[2] for row in ws.iter_rows(min_row=2, values_only=True) if row[2]
        ]
        self.assertTrue(all(r.startswith('EJEMPLO -') for r in rubros_usados))

        instrucciones = wb['Instrucciones']
        textos = ' '.join(
            str(c.value) for row in instrucciones.iter_rows() for c in row if c.value
        )
        self.assertIn('AVISO', textos)
        self.assertIn('PLACEHOLDERS', textos)

    # -- edge case 3: permisos -------------------------------------------

    def test_usuario_anonimo_no_descarga_redirige_a_login(self):
        anon = Client()
        resp = anon.get(self.url)
        self.assertEqual(resp.status_code, 302)
        self.assertIn('login', resp.url.lower())


class DescargarPlantillaPresupuestoPlanoViewSinProyectoIdTests(TestCase):
    """La vista no debe reventar si se invoca sin proyecto_id (defensivo)."""

    def setUp(self):
        self.user = _crear_usuario_admin()
        self.factory = RequestFactory()

    def test_get_sin_proyecto_id_no_revienta_usa_nombre_generico(self):
        request = self.factory.get('/financiero/plantilla-presupuesto-plano/')
        request.user = self.user
        resp = DescargarPlantillaPresupuestoPlanoView.as_view()(request)
        self.assertEqual(resp.status_code, 200)
        wb = load_workbook(io.BytesIO(resp.content))
        ws = wb['PRESUPUESTO']
        primera_fila = next(ws.iter_rows(min_row=2, max_row=2, values_only=True))
        self.assertEqual(primera_fila[1], 'Mi Proyecto')
