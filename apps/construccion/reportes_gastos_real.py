"""Instelec#268 — Reportes PDF / Excel / CSV del Presupuesto Real (gastos reales).

Mismos números que la pestaña "Real vs Planeado" de la UI: todos salen de
``gastos_real.construir_comparativo`` con los mismos filtros del querystring
(anio, periodo, clasificacion, proveedor, centro_costo). El gate por formato
(contador → solo Excel, supervisor → ninguno…) es el de #267
(``permissions_fin.user_puede_descargar_reporte``), heredado de
``_ReportePresupuestoBaseView``.
"""

from __future__ import annotations

import csv
from decimal import Decimal
from html import escape
from io import BytesIO

from django.http import HttpResponse

from . import gastos_real
from .permissions_fin import FORMATO_CSV, FORMATO_EXCEL, FORMATO_PDF
from .reportes_presupuesto import _leer_anio_querystring, _ReportePresupuestoBaseView

_SEMAFORO_TXT = {
    "verde": "Verde (bajo presupuesto)",
    "amarillo": "Amarillo (0-10% sobre)",
    "rojo": "Rojo (>10% sobre)",
    "sin_presupuesto": "Sin presupuesto",
}


def _money(valor) -> str:
    if valor is None:
        return "—"
    return "$" + f"{Decimal(valor):,.0f}".replace(",", ".")


def _pct(valor) -> str:
    return "—" if valor is None else f"{valor}%"


def filtros_desde_request(request) -> dict:
    g = request.GET
    periodo = (g.get("periodo") or "").strip()
    clasificacion = (g.get("clasificacion") or "").strip()
    return {
        "clasificacion": clasificacion
        if clasificacion in (gastos_real.CLASIFICACION_FIJO, gastos_real.CLASIFICACION_VARIABLE)
        else "",
        "proveedor": (g.get("proveedor") or "").strip(),
        "centro_costo": (g.get("centro_costo") or "").strip(),
        "periodo": periodo if len(periodo) == 6 and periodo.isdigit() else "",
    }


def _anio(request, filtros) -> int:
    return int(filtros["periodo"][:4]) if filtros["periodo"] else _leer_anio_querystring(request)


def _alcance(anio, filtros, comp) -> str:
    partes = [f"Año {anio}"]
    if filtros["periodo"]:
        partes.append(f"período {filtros['periodo']}")
    elif comp["meses_alcance_label"]:
        partes.append(f"meses con ejecución: {comp['meses_alcance_label']}")
    for clave, etiqueta in (
        ("clasificacion", "clasificación"),
        ("proveedor", "proveedor"),
        ("centro_costo", "centro de costo"),
    ):
        if filtros[clave]:
            partes.append(f"{etiqueta} {filtros[clave]}")
    return " · ".join(partes)


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def generar_excel_gastos_real(proyecto, anio, filtros) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    comp = gastos_real.construir_comparativo(proyecto, anio, filtros)
    wb = Workbook()
    ws = wb.active
    ws.title = "Real vs Planeado"
    ws.append([f"Presupuesto Real vs Planeado — {proyecto.nombre}"])
    ws.append([_alcance(anio, filtros, comp)])
    kpi = comp["kpi"]
    ws.append([])
    ws.append(["Planeado", "Real", "Variación ($)", "Variación (%)", "Cumplimiento (%)"])
    ws.append(
        [
            float(kpi["planeado"]),
            float(kpi["real"]),
            float(kpi["variacion"]),
            float(kpi["variacion_pct"]) if kpi["variacion_pct"] is not None else None,
            float(kpi["cumplimiento_pct"]) if kpi["cumplimiento_pct"] is not None else None,
        ]
    )
    ws.append([])
    ws.append(
        [
            "Rubro / Cuenta Equiv",
            *gastos_real.MESES_CORTOS,
            "Total Real",
            "Presupuesto",
            "Variación ($)",
            "Variación (%)",
            "Semáforo",
        ]
    )
    for fila in comp["filas"]:
        ws.append(
            [
                fila["nombre"] if fila["es_subtotal"] else f"   {fila['nombre']}",
                *[float(v) for v in fila["meses"]],
                float(fila["total_real"]),
                float(fila["presupuesto"]) if fila["presupuesto"] is not None else None,
                float(fila["variacion"]) if fila["variacion"] is not None else None,
                float(fila["variacion_pct"]) if fila["variacion_pct"] is not None else None,
                _SEMAFORO_TXT[fila["semaforo"]],
            ]
        )
        if fila["es_subtotal"]:
            for celda in ws[ws.max_row]:
                celda.font = Font(bold=True)
    for fila_idx in (1, 4, 7):
        for celda in ws[fila_idx]:
            celda.font = Font(bold=True)

    wp = wb.create_sheet("Proveedores")
    wp.append(["NIT", "Proveedor", "Estado", "Líneas", "Total (filtro)", "Acumulado año"])
    for p in comp["proveedores"]:
        estado = "—" if p["activo"] is None else ("Activo" if p["activo"] else "Inactivo")
        wp.append(
            [p["nit"], p["nombre"], estado, p["lineas"], float(p["total"]), float(p["total_anio"])]
        )

    wd = wb.create_sheet("Detalle gastos")
    wd.append(gastos_real.COLUMNAS_GASTOS_REALES + ["Rubro", "Clasificación"])
    for g in gastos_real.filtrar_gastos(proyecto, anio, filtros).order_by(
        "periodo", "cuenta_equiv", "fecha"
    ):
        wd.append(
            [
                g.auxiliar,
                g.desc_auxiliar,
                float(g.neto),
                g.fecha.strftime("%d/%m/%Y") if g.fecha else "",
                g.docto,
                g.periodo,
                g.tercero_nit,
                g.tercero_razon_social,
                g.desc_co_movto,
                g.usuario_creacion,
                g.co_movto,
                g.notas,
                g.centro_costo,
                g.desc_centro_costo,
                g.cuenta_equiv,
                g.cdec_equiv,
                g.cargo,
                g.fijo,
                g.rubro,
                g.clasificacion,
            ]
        )
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()


# ---------------------------------------------------------------------------
# CSV
# ---------------------------------------------------------------------------
def escribir_csv_gastos_real(destino, proyecto, anio, filtros) -> None:
    """CSV plano de las líneas de gasto real (con rubro y clasificación)."""
    destino.write("﻿")  # BOM: Excel abre bien acentos
    writer = csv.writer(destino, delimiter=";")
    writer.writerow(gastos_real.COLUMNAS_GASTOS_REALES + ["Rubro", "Clasificación"])
    for g in gastos_real.filtrar_gastos(proyecto, anio, filtros).order_by(
        "periodo", "cuenta_equiv", "fecha"
    ):
        writer.writerow(
            [
                g.auxiliar,
                g.desc_auxiliar,
                f"{g.neto:.2f}",
                g.fecha.strftime("%d/%m/%Y") if g.fecha else "",
                g.docto,
                g.periodo,
                g.tercero_nit,
                g.tercero_razon_social,
                g.desc_co_movto,
                g.usuario_creacion,
                g.co_movto,
                g.notas,
                g.centro_costo,
                g.desc_centro_costo,
                g.cuenta_equiv,
                g.cdec_equiv,
                g.cargo,
                g.fijo,
                g.rubro,
                g.clasificacion,
            ]
        )


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def construir_html_pdf_gastos_real(proyecto, anio, filtros, usuario=None) -> str:
    comp = gastos_real.construir_comparativo(proyecto, anio, filtros)
    kpi = comp["kpi"]
    filas = (
        "".join(
            f"<tr class='{f['semaforo']}{' sub' if f['es_subtotal'] else ''}'>"
            f"<td>{'' if f['es_subtotal'] else '&nbsp;&nbsp;'}{escape(str(f['nombre']))}</td>"
            f"<td>{_money(f['total_real'])}</td><td>{_money(f['presupuesto'])}</td>"
            f"<td>{_money(f['variacion'])}</td><td>{_pct(f['variacion_pct'])}</td>"
            f"<td>{_SEMAFORO_TXT[f['semaforo']]}</td></tr>"
            for f in comp["filas"]
        )
        or "<tr><td colspan='6'>Sin gastos reales cargados para el alcance.</td></tr>"
    )
    provs = (
        "".join(
            f"<tr><td>{escape(p['nit'])}</td><td>{escape(str(p['nombre']))}</td>"
            f"<td>{'—' if p['activo'] is None else ('Activo' if p['activo'] else 'Inactivo')}</td>"
            f"<td>{_money(p['total'])}</td><td>{_money(p['total_anio'])}</td></tr>"
            for p in comp["proveedores"][:40]
        )
        or "<tr><td colspan='5'>Sin proveedores.</td></tr>"
    )
    return f"""<!DOCTYPE html>
<html lang="es"><head><meta charset="utf-8"><title>Presupuesto Real vs Planeado</title>
<style>
  @page {{ size: A4 landscape; margin: 1.5cm; }}
  body {{ font-family: Helvetica, Arial, sans-serif; color: #1f2937; font-size: 10pt; }}
  h1 {{ font-size: 18pt; margin: 0; }} h2 {{ color: #1e40af; font-size: 12pt; margin-top: 18px; }}
  .sub {{ font-weight: bold; background: #f3f4f6; }}
  table {{ width: 100%; border-collapse: collapse; }}
  th, td {{ text-align: left; padding: 4px 6px; border-bottom: 1px solid #e5e7eb; font-size: 9pt; }}
  th {{ background: #e5e7eb; }}
  tr.rojo td {{ color: #b91c1c; }} tr.amarillo td {{ color: #92400e; }} tr.verde td:last-child {{ color: #166534; }}
  .kpi {{ display: flex; gap: 14px; margin-top: 8px; }}
  .kpi div {{ border: 1px solid #e5e7eb; border-radius: 6px; padding: 8px 12px; }}
</style></head><body>
  <h1>Presupuesto Real vs Planeado</h1>
  <p>{escape(proyecto.nombre)} · {escape(_alcance(anio, filtros, comp))}</p>
  <div class="kpi">
    <div><strong>Planeado</strong><br>{_money(kpi["planeado"])}</div>
    <div><strong>Real</strong><br>{_money(kpi["real"])}</div>
    <div><strong>Variación total</strong><br>{_money(kpi["variacion"])} ({_pct(kpi["variacion_pct"])})</div>
    <div><strong>Cumplimiento</strong><br>{_pct(kpi["cumplimiento_pct"])}</div>
  </div>
  <h2>Rubros y cuentas</h2>
  <table><thead><tr><th>Rubro / Cuenta Equiv</th><th>Total Real</th><th>Presupuesto</th>
  <th>Variación ($)</th><th>%</th><th>Semáforo</th></tr></thead><tbody>{filas}</tbody></table>
  <h2>Proveedores</h2>
  <table><thead><tr><th>NIT</th><th>Proveedor</th><th>Estado</th><th>Total</th><th>Acumulado año</th></tr></thead>
  <tbody>{provs}</tbody></table>
  <p style="color:#6b7280;font-size:8pt;margin-top:20px;">Generado por el sistema financiero de Construcción
  Instelec (Indunnova S.A.S.){f" por {escape(str(usuario))}" if usuario else ""}.</p>
</body></html>"""


# ---------------------------------------------------------------------------
# Vistas
# ---------------------------------------------------------------------------
class _ReporteGastosRealBaseView(_ReportePresupuestoBaseView):
    active_subtab = "presupuesto_real"

    def get(self, request, *args, **kwargs):
        self._filtros = filtros_desde_request(request)
        return super().get(request, *args, **kwargs)

    def _generar_response(self, request, proyecto, anio, mes):
        anio = _anio(request, self._filtros)
        response = self._respuesta(request, proyecto, anio, self._filtros)
        sufijo = self._filtros["periodo"] or str(anio)
        response["Content-Disposition"] = (
            f'attachment; filename="Presupuesto_Real_{proyecto.pk}_{sufijo}.{self.extension}"'
        )
        return response


class ReporteGastosRealPdfView(_ReporteGastosRealBaseView):
    formato_reporte = FORMATO_PDF
    content_type = "application/pdf"
    extension = "pdf"

    def _respuesta(self, request, proyecto, anio, filtros):
        from weasyprint import HTML

        html = construir_html_pdf_gastos_real(proyecto, anio, filtros, usuario=request.user)
        return HttpResponse(HTML(string=html).write_pdf(), content_type=self.content_type)


class ReporteGastosRealExcelView(_ReporteGastosRealBaseView):
    formato_reporte = FORMATO_EXCEL
    content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    extension = "xlsx"

    def _respuesta(self, request, proyecto, anio, filtros):
        return HttpResponse(
            generar_excel_gastos_real(proyecto, anio, filtros), content_type=self.content_type
        )


class ReporteGastosRealCsvView(_ReporteGastosRealBaseView):
    formato_reporte = FORMATO_CSV
    content_type = "text/csv; charset=utf-8"
    extension = "csv"

    def _respuesta(self, request, proyecto, anio, filtros):
        response = HttpResponse(content_type=self.content_type)
        escribir_csv_gastos_real(response, proyecto, anio, filtros)
        return response
