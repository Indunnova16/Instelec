"""Instelec#267 A10 — 4 Reportes descargables del Presupuesto Planeado de
Construcción (Fase 5 del issue): PDF ejecutivo, Excel 4 hojas, CSV contable,
PPT 7 diapositivas.

Contrato con A7 (``permissions_fin.py``): cada vista de acá llama a
``user_puede_descargar_reporte`` server-side ANTES de generar el archivo —
ocultar el botón en el template (A10, templates) NO alcanza por sí solo.

Contrato con A2/A5/A8: reusa ``build_rubro_display_rows``/
``build_rubro_matrix_rows`` (``apps.financiero.importers_finv2``, ya
extendidos por A4 con ``pct``/``semaforo``) y ``_kpi_cards_finv2_bd``
(``views_fin.py``, A5) — los 4 reportes ven EXACTAMENTE los mismos números
que la pestaña "Tabla" de la UI, nunca un cálculo paralelo. Cuando la
descarga trae ``?mes=``, se restringe ``filas_detalle`` (A2) a ese
mes/año y se reagrega con ``_construir_finv2_bd_desde_filas_planas``
(``apps.construccion.importers``, el MISMO agregador que A9 usa para su
respuesta puntual por período) — sin ``?mes=`` el reporte es del año fiscal
completo, igual que la tabla.

Dependencias de generación:
- **PDF**: WeasyPrint (``requirements/base.txt``, YA declarada — mismo
  patrón que ``apps/financiero/reports_finv2_ejecutivo_pdf.py``: construir
  un HTML string y convertirlo). NO se testea extrayendo texto del PDF
  binario: WeasyPrint subsetea y reindexa las fuentes embebidas, así que el
  content stream usa glyph IDs (``<0035>0<0048>...`` en el operador ``TJ``),
  NO los bytes ASCII originales — verificado empíricamente (2026-09-23):
  ni siquiera un fragmento corto como "Ingreso" aparece literal en el PDF
  final, aunque SÍ estuvo en el HTML fuente. Extraer texto real requeriría
  una librería de parsing PDF (pypdf/pdfplumber) que NO está autorizada como
  dependencia nueva solo para tests (a diferencia de python-pptx, que SÍ es
  necesaria en runtime). Por eso ``_construir_html_pdf_presupuesto`` es una
  función separada y pública: los tests verifican el HTML EXACTO que
  WeasyPrint va a convertir (mismo dato, sin la capa de fuentes/glyphs que
  lo vuelve opaco), más un smoke test de que el PDF final es válido
  (cabecera ``%PDF-`` + tamaño mínimo realista).
- **Excel**: openpyxl (YA declarada).
- **CSV**: reusa LITERAL el encabezado/orden de columnas de
  ``apps.financiero.views_finv2_carga.PlanoFinancieroCsvView``
  (``Código,Concepto,Centro,Proyecto,Mes,Valor,Referencia``) — pero la
  FUENTE de datos es distinta (JSON ``filas_detalle`` del formato plano A2,
  no líneas de ``CargaFinanciera`` con FK a homologación), así que el
  mapeo de columnas es una interpretación documentada, no un reuso 1:1 de
  campos:
    * Código    ← ``codigo_contable`` de la fila (o ``'SIN_HOMOLOGAR'``)
    * Concepto  ← ``rubro`` (no hay un campo "concepto" propio en el plano)
    * Centro    ← ``ciudad`` (el plano no tiene centro de costo; ciudad es
                  la dimensión geográfica más cercana)
    * Proyecto  ← nombre del proyecto de Construcción
    * Mes       ← ``MM-AAAA``
    * Valor     ← suma de ``valor`` del grupo
    * Referencia ← ``clasificacion`` (Fijo/Variable/Ingresos) — el plano no
                  trae un campo de trazabilidad "referencia" propio;
                  Clasificación es el único dato adicional por fila que
                  vale la pena preservar en esa columna. A confirmar con el
                  cliente si "CSV para Contabilidad" espera otra cosa ahí
                  (ver ``riesgos_operativos_no_bloqueantes`` de F2 sobre CSV
                  fuera de la tabla literal Rol×Reportes del issue).
- **PPT**: ``python-pptx`` — dependencia NUEVA, agregada a
  ``requirements/base.txt`` (sección "Reports", junto a WeasyPrint/openpyxl)
  porque el endpoint la usa en runtime, no solo en tests. Verificado ANTES
  de asumir que faltaba: no estaba en ningún ``requirements/*.txt`` del
  repo ni instalada en el venv del worktree.
"""

from __future__ import annotations

import csv
from collections import defaultdict
from datetime import date
from decimal import Decimal
from html import escape
from io import BytesIO
from typing import Any

from django.core.exceptions import PermissionDenied
from django.http import HttpResponse
from django.views import View

from apps.financiero.importers_finv2 import (
    build_rubro_display_rows,
    build_rubro_matrix_rows,
)

from .importers import _construir_finv2_bd_desde_filas_planas
from .models_fin import (
    HistorialCargaPresupuestoConstruccion,
    PresupuestoDetalladoConstruccion,
)
from .permissions_fin import (
    FORMATO_CSV,
    FORMATO_EXCEL,
    FORMATO_PDF,
    FORMATO_PPT,
    user_puede_descargar_reporte,
)
from .views_fin import ProyectoFinMixin, _kpi_cards_finv2_bd, _to_decimal

ZERO = Decimal("0.00")


def _money(valor: Any) -> str:
    try:
        return f"${Decimal(str(valor)).quantize(Decimal('0.01')):,.2f}"
    except Exception:
        return str(valor)


def _pct1(valor: Any) -> str:
    try:
        return f"{Decimal(str(valor)).quantize(Decimal('0.1'))}%"
    except Exception:
        return f"{valor}%"


# ===========================================================================
# Contexto compartido — mismos números que la pestaña "Tabla" de la UI
# ===========================================================================
def _leer_mes_querystring(request) -> int | None:
    """``?mes=`` opcional — ``None`` = reporte del año fiscal completo
    (mismo default que la pestaña Tabla/Matriz). Nunca lanza."""
    mes_raw = request.GET.get("mes")
    if mes_raw in (None, ""):
        return None
    try:
        mes = int(mes_raw)
    except (TypeError, ValueError):
        return None
    return mes if 1 <= mes <= 12 else None


def _leer_anio_querystring(request) -> int:
    try:
        return int(request.GET.get("anio", date.today().year))
    except (TypeError, ValueError):
        return date.today().year


def construir_contexto_reporte(proyecto, anio: int, mes: int | None = None) -> dict:
    """Datos base para los 4 reportes, restringidos a ``mes`` si se pide.

    Reusa el MISMO patrón que A9 (``api.py::presupuesto_periodo``): filtra
    ``filas_detalle`` (A2) por (mes, anio) y reagrega con
    ``_construir_finv2_bd_desde_filas_planas`` — nunca un cálculo paralelo.
    Presupuesto inexistente o sin filas para el período pedido → contexto
    "sin datos" (reportes generados igual, con el mismo empty-state que la
    UI — nunca un 404: una descarga que el usuario pidió a propósito no
    debe romperse por falta de dato, solo decirlo).
    """
    presupuesto = PresupuestoDetalladoConstruccion.objects.filter(
        proyecto=proyecto, anio=anio, tipo=PresupuestoDetalladoConstruccion.Tipo.PLANEADO
    ).first()
    datos = presupuesto.datos if presupuesto and isinstance(presupuesto.datos, dict) else {}
    filas_detalle_anio = ((datos.get("finv2_bd") or {}).get("filas_detalle")) or []

    if mes is not None:
        filas_efectivas = [
            f
            for f in filas_detalle_anio
            if isinstance(f, dict) and f.get("mes") == mes and f.get("anio") == anio
        ]
        datos_periodo = (
            {"finv2_bd": _construir_finv2_bd_desde_filas_planas(filas_efectivas)}
            if filas_efectivas
            else {}
        )
    else:
        filas_efectivas = filas_detalle_anio
        datos_periodo = datos

    rubro_rows, rubro_total = build_rubro_display_rows(datos_periodo)
    matrix_rows, totales_columna, meses_fiscales, matrix_total = build_rubro_matrix_rows(
        datos_periodo
    )
    kpi_cards = _kpi_cards_finv2_bd(datos_periodo)

    return {
        "presupuesto": presupuesto,
        "rubro_rows": rubro_rows,
        "rubro_total": rubro_total,
        "matrix_rows": matrix_rows,
        "totales_columna": totales_columna,
        "meses_fiscales": meses_fiscales,
        "matrix_total": matrix_total,
        "kpi_cards": kpi_cards,
        "filas_detalle": filas_efectivas,
        "tiene_datos": bool(filas_efectivas) or bool(rubro_rows),
    }


def _periodo_label(anio: int, mes: int | None) -> str:
    return f"{mes:02d}/{anio}" if mes else f"Año fiscal {anio}"


# ===========================================================================
# 1. PDF ejecutivo (Fase 5.1) — WeasyPrint
# ===========================================================================
def _construir_html_pdf_presupuesto(
    proyecto,
    anio: int,
    mes: int | None,
    ctx: dict,
    generado_en=None,
    usuario=None,
) -> str:
    """HTML fuente del PDF ejecutivo — función separada y pública a propósito
    (ver docstring del módulo: es lo que los tests verifican, no el binario)."""
    periodo = _periodo_label(anio, mes)
    kpi = ctx["kpi_cards"]
    rubro_rows = ctx["rubro_rows"]

    filas_rubros = (
        "".join(
            f"<tr class='{r['semaforo']}'>"
            f"<td>{escape(str(r['rubro']))}</td>"
            f"<td>{_money(r['total'])}</td>"
            f"<td>{_pct1(r['pct'])}</td>"
            "</tr>"
            for r in rubro_rows
        )
        or "<tr><td colspan='3'>Sin rubros cargados para el período.</td></tr>"
    )

    alertas_rojo = [r for r in rubro_rows if r["semaforo"] == "rojo"]
    filas_alertas = (
        "".join(
            f"<li>{escape(str(r['rubro']))}: {_pct1(r['pct'])} sobre |Total Año| "
            f"({_money(r['total'])})</li>"
            for r in alertas_rojo
        )
        or "<li>Ningún rubro supera el 100% de |Total Año| en el período.</li>"
    )

    html = f"""<!DOCTYPE html>
<html lang="es">
<head><meta charset="utf-8"><title>Presupuesto Planeado — Reporte Ejecutivo</title>
<style>
  @page {{ size: A4; margin: 2cm; }}
  body {{ font-family: 'Helvetica', 'Arial', sans-serif; color: #1f2937; font-size: 11pt; }}
  h1 {{ color: #111827; font-size: 22pt; margin-bottom: 0; }}
  h2 {{ color: #1e40af; font-size: 14pt; margin-top: 24px; border-bottom: 1px solid #d1d5db; padding-bottom: 4px; }}
  .subtitulo {{ color: #6b7280; font-size: 12pt; margin-top: 4px; }}
  table {{ width: 100%; border-collapse: collapse; margin-top: 10px; }}
  th, td {{ text-align: left; padding: 6px 8px; border-bottom: 1px solid #e5e7eb; font-size: 10pt; }}
  th {{ background: #f3f4f6; }}
  tr.rojo td {{ color: #b91c1c; font-weight: bold; }}
  tr.amarillo td {{ color: #92400e; }}
  .resumen {{ display: flex; gap: 24px; margin-top: 12px; flex-wrap: wrap; }}
  .resumen div {{ background: #f9fafb; border: 1px solid #e5e7eb; border-radius: 6px; padding: 10px 16px; }}
  .firma {{ margin-top: 40px; padding-top: 12px; border-top: 1px solid #d1d5db; color: #6b7280; font-size: 9pt; }}
  ul {{ margin: 8px 0; padding-left: 20px; }}
</style>
</head>
<body>
  <h1>Presupuesto Planeado — Reporte Ejecutivo</h1>
  <p class="subtitulo">{escape(proyecto.nombre)} &middot; Período {periodo}</p>

  <h2>KPI Ejecutivos</h2>
  <div class="resumen">
    <div><strong>Ingreso</strong><br>{_money(kpi["ingreso"])}</div>
    <div><strong>Costos Fijos</strong><br>{_money(kpi["costos_fijos"])}</div>
    <div><strong>Costos Variables</strong><br>{_money(kpi["costos_variables"])}</div>
    <div><strong>Resultado</strong><br>{_money(kpi["resultado"])}</div>
  </div>

  <h2>Rubros ({len(rubro_rows)})</h2>
  <table>
    <thead><tr><th>Rubro</th><th>Total</th><th>% sobre |Total Año|</th></tr></thead>
    <tbody>{filas_rubros}</tbody>
  </table>
  <p style="font-size:9pt;color:#6b7280;">Total general: {_money(ctx["rubro_total"])}</p>

  <h2>Alertas (semáforo rojo, &gt;100%)</h2>
  <ul>{filas_alertas}</ul>

  <div class="firma">
    Generado automáticamente por el sistema financiero de Construcción Instelec (Indunnova S.A.S.)
    {f"el {generado_en:%d/%m/%Y %H:%M}" if generado_en else ""}
    {f"por {escape(str(usuario))}" if usuario else ""}.
  </div>
</body>
</html>"""
    return html


def generar_pdf_presupuesto(
    proyecto, anio: int, mes: int | None = None, generado_en=None, usuario=None
) -> bytes:
    from weasyprint import HTML

    ctx = construir_contexto_reporte(proyecto, anio, mes)
    html = _construir_html_pdf_presupuesto(
        proyecto, anio, mes, ctx, generado_en=generado_en, usuario=usuario
    )
    return HTML(string=html).write_pdf()


# ===========================================================================
# 2. Excel 4 hojas (Fase 5.2) — openpyxl: Resumen / Matriz / Rubros / Histórico
# ===========================================================================
def generar_excel_presupuesto(proyecto, anio: int, mes: int | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    ctx = construir_contexto_reporte(proyecto, anio, mes)
    periodo = _periodo_label(anio, mes)
    kpi = ctx["kpi_cards"]

    libro = Workbook()

    # --- Hoja 1: Resumen -----------------------------------------------
    hoja_resumen = libro.active
    hoja_resumen.title = "Resumen"
    hoja_resumen.append(["Presupuesto Planeado — Reporte Ejecutivo"])
    hoja_resumen["A1"].font = Font(bold=True, size=14)
    hoja_resumen.append(["Proyecto", proyecto.nombre])
    hoja_resumen.append(["Período", periodo])
    hoja_resumen.append([])
    hoja_resumen.append(["Indicador", "Valor"])
    hoja_resumen.append(["Ingreso", float(kpi["ingreso"])])
    hoja_resumen.append(["Costos Fijos", float(kpi["costos_fijos"])])
    hoja_resumen.append(["Costos Variables", float(kpi["costos_variables"])])
    hoja_resumen.append(["Resultado", float(kpi["resultado"])])
    hoja_resumen.append(["Total Rubros", ctx["rubro_total"]])

    # --- Hoja 2: Matriz (rubro × 12 meses fiscales) ---------------------
    hoja_matriz = libro.create_sheet("Matriz")
    encabezados = (
        ["Rubro"]
        + [label for _key, label, _num in ctx["meses_fiscales"]]
        + ["Total", "% Total", "Semáforo"]
    )
    hoja_matriz.append(encabezados)
    for fila in ctx["matrix_rows"]:
        hoja_matriz.append(
            [fila["rubro"], *fila["meses"], fila["total"], fila["pct"], fila["semaforo"]]
        )
    hoja_matriz.append(["TOTAL", *ctx["totales_columna"], ctx["matrix_total"], "", ""])

    # --- Hoja 3: Rubros --------------------------------------------------
    hoja_rubros = libro.create_sheet("Rubros")
    hoja_rubros.append(["Rubro", "Total", "% sobre |Total Año|", "Semáforo"])
    for fila in ctx["rubro_rows"]:
        hoja_rubros.append([fila["rubro"], fila["total"], fila["pct"], fila["semaforo"]])

    # --- Hoja 4: Histórico (A3 — últimas 50 cargas del proyecto) --------
    hoja_hist = libro.create_sheet("Histórico")
    hoja_hist.append(["Fecha", "Usuario", "Filas", "Valor total", "Estado", "Período"])
    cargas = (
        HistorialCargaPresupuestoConstruccion.objects.filter(proyecto=proyecto)
        .select_related("usuario")
        .order_by("-fecha")[:50]
    )
    for carga in cargas:
        usuario_nombre = "—"
        if carga.usuario:
            usuario_nombre = carga.usuario.get_full_name() or carga.usuario.username
        hoja_hist.append(
            [
                carga.fecha.strftime("%d/%m/%Y %H:%M:%S") if carga.fecha else "",
                usuario_nombre,
                carga.filas_procesadas,
                float(carga.valor_total or 0),
                carga.get_estado_display(),
                carga.periodo_display,
            ]
        )

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()


# ===========================================================================
# 3. CSV contable (Fase 5.4) — reusa LITERAL el layout de PlanoFinancieroCsvView
# ===========================================================================
def escribir_csv_presupuesto(
    response: HttpResponse, proyecto, anio: int, mes: int | None = None
) -> None:
    """Escribe el CSV contable en ``response`` (ya abierto por el caller con
    ``content_type``/``Content-Disposition`` seteados) — mismo patrón
    (BOM + ``csv.writer`` + agrupar/sumar + orden alfabético) que
    ``PlanoFinancieroCsvView``, ver mapeo de columnas en el docstring del
    módulo.

    Fallback (hallazgo del validador-cierre, 2026-09-23): ``filas_detalle``
    solo existe para presupuestos cargados por el importador plano (A2) — un
    presupuesto cargado por el importador CONTABLE legacy (el caso real hoy
    en prod) no tiene esa llave y antes producía un CSV con SOLO encabezado,
    sin aviso. Si no hay ``filas_detalle`` pero SÍ hay ``rubro_rows``
    (dato legacy real), se emite una fila anual por rubro con el MISMO total
    que ya muestran la pestaña Tabla / hoja "Rubros" del Excel (A10) —
    ``build_rubro_display_rows``, sin inventar un desglose mensual que este
    formato de dato no tiene."""
    ctx = construir_contexto_reporte(proyecto, anio, mes)
    filas = ctx["filas_detalle"]

    response.write("﻿")
    writer = csv.writer(response)
    writer.writerow(["Código", "Concepto", "Centro", "Proyecto", "Mes", "Valor", "Referencia"])

    if not filas:
        for fila in ctx["rubro_rows"]:
            writer.writerow(
                [
                    "SIN_HOMOLOGAR",
                    fila["rubro"],
                    "",
                    proyecto.nombre,
                    _periodo_label(anio, mes),
                    fila["total"],
                    "",
                ]
            )
        return

    agrupadas: dict = defaultdict(lambda: {"valor": ZERO, "referencias": set()})
    for f in filas:
        if not isinstance(f, dict):
            continue
        codigo = f.get("codigo_contable") or "SIN_HOMOLOGAR"
        concepto = f.get("rubro") or ""
        centro = f.get("ciudad") or ""
        mes_f = f.get("mes")
        anio_f = f.get("anio")
        llave = (codigo, concepto, centro, mes_f, anio_f)
        agrupadas[llave]["valor"] += _to_decimal(f.get("valor"))
        clasificacion = f.get("clasificacion")
        if clasificacion:
            agrupadas[llave]["referencias"].add(str(clasificacion))

    for (codigo, concepto, centro, mes_f, anio_f), info in sorted(
        agrupadas.items(), key=lambda kv: (kv[0][0] or "", kv[0][1] or "", kv[0][2] or "")
    ):
        mes_label = f"{mes_f:02d}-{anio_f}" if mes_f and anio_f else str(anio_f or anio)
        writer.writerow(
            [
                codigo,
                concepto,
                centro,
                proyecto.nombre,
                mes_label,
                info["valor"],
                " | ".join(sorted(info["referencias"])),
            ]
        )


# ===========================================================================
# 4. PPT 7 diapositivas (Fase 5.3) — python-pptx (dependencia NUEVA)
# ===========================================================================
def generar_ppt_presupuesto(proyecto, anio: int, mes: int | None = None) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    ctx = construir_contexto_reporte(proyecto, anio, mes)
    periodo = _periodo_label(anio, mes)
    kpi = ctx["kpi_cards"]
    rubro_rows = ctx["rubro_rows"]
    alertas_rojo = [r for r in rubro_rows if r["semaforo"] == "rojo"]

    prs = Presentation()
    layout_titulo = prs.slide_layouts[0]
    layout_contenido = prs.slide_layouts[1]
    layout_blanco = prs.slide_layouts[6]

    def _slide_tabla(titulo, encabezados, filas, vacio_msg):
        slide = prs.slides.add_slide(layout_blanco)
        caja_titulo = slide.shapes.add_textbox(Inches(0.4), Inches(0.3), Inches(9), Inches(0.6))
        caja_titulo.text_frame.text = titulo
        caja_titulo.text_frame.paragraphs[0].font.size = Pt(24)
        caja_titulo.text_frame.paragraphs[0].font.bold = True

        n_filas = len(filas) if filas else 1
        tabla_shape = slide.shapes.add_table(
            n_filas + 1,
            len(encabezados),
            Inches(0.4),
            Inches(1.1),
            Inches(9),
            Inches(0.4 * (n_filas + 1)),
        )
        tabla = tabla_shape.table
        for c, encabezado in enumerate(encabezados):
            tabla.cell(0, c).text = str(encabezado)
        if filas:
            for r, fila in enumerate(filas, start=1):
                for c, valor in enumerate(fila):
                    tabla.cell(r, c).text = str(valor)
        else:
            tabla.cell(1, 0).text = vacio_msg
        return slide

    # Slide 1: Portada
    slide1 = prs.slides.add_slide(layout_titulo)
    slide1.shapes.title.text = "Presupuesto Planeado — Reporte Ejecutivo"
    slide1.placeholders[1].text = f"{proyecto.nombre}\nPeríodo {periodo}"

    # Slide 2: KPI Ejecutivos
    slide2 = prs.slides.add_slide(layout_contenido)
    slide2.shapes.title.text = "KPI Ejecutivos"
    cuerpo = slide2.placeholders[1].text_frame
    cuerpo.text = f"Ingreso: {_money(kpi['ingreso'])}"
    for etiqueta, valor in (
        ("Costos Fijos", kpi["costos_fijos"]),
        ("Costos Variables", kpi["costos_variables"]),
        ("Resultado", kpi["resultado"]),
    ):
        parrafo = cuerpo.add_paragraph()
        parrafo.text = f"{etiqueta}: {_money(valor)}"

    # Slide 3: Rubros (top 10 por total)
    _slide_tabla(
        "Rubros (top 10)",
        ["Rubro", "Total", "% sobre |Total Año|", "Semáforo"],
        [(r["rubro"], _money(r["total"]), _pct1(r["pct"]), r["semaforo"]) for r in rubro_rows[:10]],
        "Sin rubros cargados para el período.",
    )

    # Slide 4: Alertas (semáforo rojo)
    _slide_tabla(
        "Alertas — Rubros en rojo (>100%)",
        ["Rubro", "Total", "% sobre |Total Año|"],
        [(r["rubro"], _money(r["total"]), _pct1(r["pct"])) for r in alertas_rojo],
        "Ningún rubro supera el 100% de |Total Año| en el período.",
    )

    # Slide 5: Matriz mensual (totales por columna)
    _slide_tabla(
        "Totales por mes (matriz)",
        [label for _key, label, _num in ctx["meses_fiscales"]],
        [tuple(_money(v) for v in ctx["totales_columna"])] if ctx["matrix_rows"] else [],
        "Sin datos de matriz mensual para el período.",
    )

    # Slide 6: Historial de cargas (últimas 5)
    cargas = list(
        HistorialCargaPresupuestoConstruccion.objects.filter(proyecto=proyecto)
        .select_related("usuario")
        .order_by("-fecha")[:5]
    )
    _slide_tabla(
        "Historial de Cargas (últimas 5)",
        ["Fecha", "Usuario", "Filas", "Valor total", "Estado"],
        [
            (
                c.fecha.strftime("%d/%m/%Y") if c.fecha else "",
                (c.usuario.get_full_name() or c.usuario.username) if c.usuario else "—",
                c.filas_procesadas,
                _money(c.valor_total),
                c.get_estado_display(),
            )
            for c in cargas
        ],
        "Aún no hay cargas registradas para este proyecto.",
    )

    # Slide 7: Cierre / notas
    slide7 = prs.slides.add_slide(layout_contenido)
    slide7.shapes.title.text = "Notas del reporte"
    cuerpo7 = slide7.placeholders[1].text_frame
    cuerpo7.text = f"Fuente: Presupuesto Planeado {periodo}, proyecto {proyecto.nombre}."
    p2 = cuerpo7.add_paragraph()
    p2.text = (
        'Los valores reflejan la misma agregación que la pestaña "Tabla" del módulo financiero.'
    )
    p3 = cuerpo7.add_paragraph()
    p3.text = "Generado automáticamente por el sistema financiero de Construcción Instelec (Indunnova S.A.S.)."

    salida = BytesIO()
    prs.save(salida)
    return salida.getvalue()


# ===========================================================================
# Vistas — gate de FORMATO server-side (A7) antes de generar cada archivo
# ===========================================================================
class _ReportePresupuestoBaseView(ProyectoFinMixin, View):
    """Base común a los 4 endpoints de descarga.

    ``ProyectoFinMixin`` (``LoginRequiredMixin`` + ``RoleRequiredMixin`` con
    ``required_submodulo='FINANCIERO'``) ya gatea el acceso de LECTURA al
    módulo (nivel ``ver``/``ver_editar`` — mismo nivel que abrir la pestaña
    Tabla, GET). El gate ESPECÍFICO por formato (contador→excel,
    supervisor→ninguno, etc., A7) se aplica acá encima, server-side, con
    ``permissions_fin.user_puede_descargar_reporte`` — NUNCA basta con
    ocultar el botón en el template.
    """

    active_subtab = "presupuesto_planeado"
    formato_reporte: str | None = None
    content_type: str = "application/octet-stream"
    extension: str = "bin"

    def _nombre_archivo(self, proyecto, anio, mes):
        """``ProyectoConstruccion`` no tiene un campo ``codigo`` corto (solo
        ``nombre``/``pk`` UUID) — se usa el ``pk`` como identificador estable
        y sin espacios/acentos para el nombre de archivo."""
        periodo = f"{mes:02d}_{anio}" if mes else str(anio)
        return f"Presupuesto_Planeado_{proyecto.pk}_{periodo}.{self.extension}"

    def get(self, request, *args, **kwargs):
        if not user_puede_descargar_reporte(request.user, self.formato_reporte):
            raise PermissionDenied(
                "Su rol no tiene permiso para descargar reportes en formato "
                f"{self.formato_reporte!r}."
            )
        proyecto = self.get_proyecto()
        anio = _leer_anio_querystring(request)
        mes = _leer_mes_querystring(request)
        return self._generar_response(request, proyecto, anio, mes)

    def _generar_response(self, request, proyecto, anio, mes):
        raise NotImplementedError


class ReportePresupuestoPdfView(_ReportePresupuestoBaseView):
    formato_reporte = FORMATO_PDF
    content_type = "application/pdf"
    extension = "pdf"

    def _generar_response(self, request, proyecto, anio, mes):
        contenido = generar_pdf_presupuesto(
            proyecto,
            anio,
            mes,
            generado_en=None,
            usuario=request.user if request.user.is_authenticated else None,
        )
        response = HttpResponse(contenido, content_type=self.content_type)
        response["Content-Disposition"] = (
            f'attachment; filename="{self._nombre_archivo(proyecto, anio, mes)}"'
        )
        return response


class ReportePresupuestoExcelView(_ReportePresupuestoBaseView):
    formato_reporte = FORMATO_EXCEL
    content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    extension = "xlsx"

    def _generar_response(self, request, proyecto, anio, mes):
        contenido = generar_excel_presupuesto(proyecto, anio, mes)
        response = HttpResponse(contenido, content_type=self.content_type)
        response["Content-Disposition"] = (
            f'attachment; filename="{self._nombre_archivo(proyecto, anio, mes)}"'
        )
        return response


class ReportePresupuestoCsvView(_ReportePresupuestoBaseView):
    formato_reporte = FORMATO_CSV
    content_type = "text/csv; charset=utf-8"
    extension = "csv"

    def _generar_response(self, request, proyecto, anio, mes):
        response = HttpResponse(content_type=self.content_type)
        response["Content-Disposition"] = (
            f'attachment; filename="{self._nombre_archivo(proyecto, anio, mes)}"'
        )
        escribir_csv_presupuesto(response, proyecto, anio, mes)
        return response


class ReportePresupuestoPptView(_ReportePresupuestoBaseView):
    formato_reporte = FORMATO_PPT
    content_type = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    extension = "pptx"

    def _generar_response(self, request, proyecto, anio, mes):
        contenido = generar_ppt_presupuesto(proyecto, anio, mes)
        response = HttpResponse(contenido, content_type=self.content_type)
        response["Content-Disposition"] = (
            f'attachment; filename="{self._nombre_archivo(proyecto, anio, mes)}"'
        )
        return response
