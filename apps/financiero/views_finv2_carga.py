"""Workflow de carga financiera, homologación y plano contable (#246, #247)."""
from __future__ import annotations

import csv
from io import BytesIO
from collections import defaultdict
from decimal import Decimal

from django.contrib import messages
from django.contrib.auth.mixins import LoginRequiredMixin
from django.db.models import Q
from django.http import Http404, HttpResponse
from django.shortcuts import redirect
from django.utils import timezone
from django.views.generic import TemplateView

from apps.contratos.models import Contrato
from apps.core.mixins import RoleRequiredMixin
from apps.core.models_roles import RoleModuloPermiso
from apps.core.permissions import SUBMODULO_FIN_HOMOLOGACION, user_nivel_acceso_submodulo

from openpyxl import Workbook

from .forms_finv2_carga import CargaFinancieraForm, ImportarTablaMaestraForm
from .importers_finv2_carga import confirmar_tabla_maestra, previsualizar_tabla_maestra, procesar_carga_financiera
from .models_finv2_carga import (
    CargaFinanciera,
    HomologacionProjectsContable,
    LineaCargaFinanciera,
    VersionHomologacionProjectsContable,
)
from .services_finv2_indicadores_integracion import construir_contexto_dashboard_integrado


ZERO = Decimal('0.00')


def _porcentaje(numerador, denominador):
    """Evita falsos indicadores cuando aún no existe una base comparable."""
    if not denominador:
        return None
    return (numerador / denominador * 100).quantize(Decimal('0.01'))


class CargaFinancieraView(LoginRequiredMixin, RoleRequiredMixin, TemplateView):
    """Carga un libro por proyecto/período y conserva su auditoría inmutable."""

    template_name = 'financiero/carga_financiera.html'
    required_submodulo = SUBMODULO_FIN_HOMOLOGACION
    admin_bypass = False

    def _proyecto_seleccionado(self):
        proyecto_id = self.request.GET.get('proyecto') or self.request.POST.get('proyecto')
        if not proyecto_id:
            return None
        return Contrato.objects.filter(pk=proyecto_id, estado=Contrato.Estado.ACTIVO).first()

    def _periodo(self):
        periodo = self.request.GET.get('periodo') or self.request.POST.get('periodo')
        if periodo:
            try:
                anio, mes = int(periodo[:4]), int(periodo[4:])
            except (TypeError, ValueError):
                return timezone.now().year, timezone.now().month, False
        else:
            anio = self.request.GET.get('anio') or self.request.POST.get('anio') or timezone.now().year
            mes = self.request.GET.get('mes') or self.request.POST.get('mes') or timezone.now().month
        try:
            anio, mes = int(anio), int(mes)
        except (TypeError, ValueError):
            return timezone.now().year, timezone.now().month, False
        return anio, mes, 1 <= mes <= 12

    def _filtros(self):
        """Valores GET validados después de período y proyecto."""
        return (
            (self.request.GET.get('tipo') or self.request.POST.get('tipo') or '').strip(),
            (self.request.GET.get('centro_costo') or self.request.POST.get('centro_costo') or '').strip(),
        )

    @staticmethod
    def _base_lineas(carga, *, tipo_operacional='', centro_costo=''):
        """Aplica los predicados Tipo/CdeC antes de cualquier agregación (#261)."""
        lineas = carga.lineas.all()
        if tipo_operacional:
            lineas = lineas.filter(tipo_operacional=tipo_operacional)
        if centro_costo:
            lineas = lineas.filter(centro_costo=centro_costo)
        return lineas

    @classmethod
    def _totales(cls, carga, *, tipo_operacional='', centro_costo=''):
        """Costo real ejecutado (`BD Real`) y costo presupuestado (`BD Ppto`).

        El libro TRANSELCA no trae una hoja de facturación/ingresos — ``BD
        Real`` es costo real incurrido, no venta. Por eso los seis
        indicadores se calculan sobre bases de costo, nunca inventando una
        facturación que la fuente no tiene (#246 Sprint B, plan
        PLAN_2026-09-15_indicadores_246.md sección "Sprint B").
        """
        if not carga:
            return ZERO, ZERO
        lineas = cls._base_lineas(carga, tipo_operacional=tipo_operacional, centro_costo=centro_costo)
        costo_real = sum((linea.valor for linea in lineas.filter(tipo=LineaCargaFinanciera.Tipo.REAL)), ZERO)
        costo_presupuestado = sum((linea.valor for linea in lineas.filter(tipo=LineaCargaFinanciera.Tipo.PRESUPUESTO)), ZERO)
        return costo_real, costo_presupuestado

    @classmethod
    def _carga_periodo_anterior(cls, carga):
        """Última carga vigente de un período estrictamente anterior al mismo proyecto."""
        anteriores = CargaFinanciera.objects.filter(
            proyecto=carga.proyecto, vigente=True,
        ).exclude(pk=carga.pk).filter(
            Q(anio__lt=carga.anio) | Q(anio=carga.anio, mes__lt=carga.mes)
        )
        return anteriores.order_by('-anio', '-mes').first()

    @classmethod
    def _indicadores(cls, carga, *, tipo_operacional='', centro_costo=''):
        if not carga:
            return []
        costo_real, costo_presupuestado = cls._totales(carga, tipo_operacional=tipo_operacional, centro_costo=centro_costo)
        margen = costo_presupuestado - costo_real  # positivo = se gastó menos de lo presupuestado
        anterior = cls._carga_periodo_anterior(carga)
        costo_real_anterior = None
        if anterior:
            costo_real_anterior, _ = cls._totales(anterior, tipo_operacional=tipo_operacional, centro_costo=centro_costo)

        sin_base_facturacion = 'Sin base comparable: el libro cargado no trae una hoja de facturación/ingresos, sólo costo real y presupuestado.'
        variacion_mes = _porcentaje(costo_real - costo_real_anterior, costo_real_anterior) if costo_real_anterior is not None else None

        return [
            {
                'nombre': 'Facturación vs meta',
                'valor': None,
                'unidad': '%',
                'alerta': False,
                'sin_base': True,
                'causa': sin_base_facturacion,
            },
            {
                'nombre': 'Margen bruto',
                'valor': margen,
                'unidad': '$',
                'alerta': margen < ZERO,
                'sin_base': False,
                'causa': 'Costo real superó el presupuesto del período.' if margen < ZERO else '',
            },
            {
                'nombre': 'Ratio costos/facturación',
                'valor': None,
                'unidad': '%',
                'alerta': False,
                'sin_base': True,
                'causa': sin_base_facturacion,
            },
            {
                'nombre': 'Rentabilidad operativa',
                'valor': None,
                'unidad': '%',
                'alerta': False,
                'sin_base': True,
                'causa': sin_base_facturacion,
            },
            {
                'nombre': 'Cumplimiento presupuesto',
                'valor': _porcentaje(costo_real, costo_presupuestado),
                'unidad': '%',
                'alerta': bool(costo_presupuestado and costo_real > costo_presupuestado),
                'sin_base': not costo_presupuestado,
                'causa': (
                    'Sin base comparable: el período no tiene presupuesto cargado (BD Ppto vacía).'
                    if not costo_presupuestado
                    else (f'Costo real {_porcentaje(costo_real, costo_presupuestado)}% del presupuesto — sobre lo aprobado.' if costo_presupuestado and costo_real > costo_presupuestado else '')
                ),
            },
            {
                'nombre': 'Variación mes a mes',
                'valor': variacion_mes,
                'unidad': '%',
                'alerta': bool(costo_real_anterior and costo_real > costo_real_anterior),
                'sin_base': costo_real_anterior is None,
                'causa': (
                    'Sin base comparable: no hay carga vigente de un período anterior para este proyecto/filtro.'
                    if costo_real_anterior is None
                    else (f'Costo real subió {variacion_mes}% frente al período anterior.' if costo_real_anterior and costo_real > costo_real_anterior else '')
                ),
            },
        ]

    @classmethod
    def _desglose_costos(cls, carga, *, tipo_operacional='', centro_costo=''):
        """Desglose por grupo cuya suma reconcilia EXACTO con el costo real total (B2)."""
        if not carga:
            return []
        lineas = cls._base_lineas(carga, tipo_operacional=tipo_operacional, centro_costo=centro_costo).filter(
            tipo=LineaCargaFinanciera.Tipo.REAL,
        )
        acumulado = defaultdict(lambda: ZERO)
        for linea in lineas:
            clave = linea.grupo or 'Sin grupo'
            acumulado[clave] += linea.valor
        total = sum(acumulado.values(), ZERO)
        desglose = []
        for grupo, valor in sorted(acumulado.items(), key=lambda item: item[1], reverse=True):
            desglose.append({
                'grupo': grupo,
                'valor': valor,
                'porcentaje': _porcentaje(valor, total) if total else None,
            })
        return desglose

    @classmethod
    def _tendencia_6_meses(cls, carga, *, tipo_operacional='', centro_costo=''):
        """Costo real/presupuestado de las últimas 6 cargas vigentes del proyecto (B2)."""
        if not carga:
            return []
        cargas = list(
            CargaFinanciera.objects.filter(proyecto=carga.proyecto, vigente=True)
            .filter(Q(anio__lt=carga.anio) | Q(anio=carga.anio, mes__lte=carga.mes))
            .order_by('-anio', '-mes')[:6]
        )
        cargas.reverse()
        serie = []
        for item in cargas:
            costo_real, costo_presupuestado = cls._totales(item, tipo_operacional=tipo_operacional, centro_costo=centro_costo)
            serie.append({
                'periodo': f'{item.mes:02d}/{item.anio}',
                'costo_real': costo_real,
                'costo_presupuestado': costo_presupuestado,
                'es_actual': item.pk == carga.pk,
            })
        return serie

    @classmethod
    def _resumen_totales(cls, carga, *, tipo_operacional='', centro_costo=''):
        if not carga:
            return None
        costo_real, costo_presupuestado = cls._totales(carga, tipo_operacional=tipo_operacional, centro_costo=centro_costo)
        return {'costo_real': costo_real, 'costo_presupuestado': costo_presupuestado}

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        anio, mes, periodo_valido = self._periodo()
        proyecto = self._proyecto_seleccionado()
        tipo_operacional, centro_costo = self._filtros()
        cargas_vigentes = CargaFinanciera.objects.filter(vigente=True)
        periodos_disponibles = list(cargas_vigentes.order_by('-anio', '-mes').values_list('anio', 'mes'))
        proyectos_disponibles = Contrato.objects.filter(estado=Contrato.Estado.ACTIVO)
        if periodo_valido:
            proyectos_disponibles = proyectos_disponibles.filter(
                cargas_financieras__vigente=True, cargas_financieras__anio=anio, cargas_financieras__mes=mes,
            ).distinct()
        carga = None
        if proyecto and periodo_valido:
            carga = CargaFinanciera.objects.filter(
                proyecto=proyecto, anio=anio, mes=mes, vigente=True,
            ).first()
        estado = self.request.GET.get('estado', 'vigentes')
        homologaciones = HomologacionProjectsContable.objects.all()
        if estado != 'archivadas':
            homologaciones = homologaciones.filter(activo=True)
        elif estado == 'archivadas':
            homologaciones = homologaciones.filter(activo=False)
        buscar = (self.request.GET.get('buscar') or '').strip()
        tipo = (self.request.GET.get('tipo') or '').strip()
        if buscar:
            homologaciones = homologaciones.filter(concepto__icontains=buscar)
        if tipo:
            homologaciones = homologaciones.filter(tipo__iexact=tipo)
        base_lineas = carga.lineas.all() if carga else LineaCargaFinanciera.objects.none()
        tipos_disponibles = list(base_lineas.order_by('tipo_operacional').values_list('tipo_operacional', flat=True).distinct())
        if tipo_operacional and tipo_operacional not in tipos_disponibles:
            tipo_operacional = ''
        lineas_por_tipo = base_lineas.filter(tipo_operacional=tipo_operacional) if tipo_operacional else base_lineas
        centros_disponibles = list(lineas_por_tipo.exclude(centro_costo='').order_by('centro_costo').values_list('centro_costo', flat=True).distinct())
        if centro_costo and centro_costo not in centros_disponibles:
            centro_costo = ''
        context.update({
            'carga_form': kwargs.get('carga_form') or CargaFinancieraForm(),
            'proyectos': Contrato.objects.filter(estado=Contrato.Estado.ACTIVO),
            'proyectos_disponibles': proyectos_disponibles,
            'periodos_disponibles': periodos_disponibles,
            'proyecto_seleccionado': proyecto,
            'anio': anio,
            'mes': mes,
            'periodo_valido': periodo_valido,
            'carga_actual': carga,
            'tipo_seleccionado': tipo_operacional,
            'centro_costo_seleccionado': centro_costo,
            'tipos_disponibles': tipos_disponibles,
            'centros_disponibles': centros_disponibles,
            'historial_cargas': CargaFinanciera.objects.select_related('proyecto', 'usuario').all()[:20],
            'homologaciones': homologaciones[:100],
            'buscar_homologacion': buscar,
            'tipo_homologacion': tipo,
            'estado_homologacion': estado,
            'puede_editar_homologacion': user_nivel_acceso_submodulo(
                self.request.user, SUBMODULO_FIN_HOMOLOGACION
            ) == RoleModuloPermiso.VER_EDITAR,
            'versiones_homologacion': VersionHomologacionProjectsContable.objects.select_related('autor').all()[:20],
            'preview_homologacion': self.request.session.get('homologacion_preview'),
            'importar_homologacion_form': ImportarTablaMaestraForm(),
            'indicadores': (indicadores_calculados := self._indicadores(carga, tipo_operacional=tipo_operacional, centro_costo=centro_costo)),
            'desglose_costos': self._desglose_costos(carga, tipo_operacional=tipo_operacional, centro_costo=centro_costo),
            'tendencia_6_meses': self._tendencia_6_meses(carga, tipo_operacional=tipo_operacional, centro_costo=centro_costo),
            'resumen_alertas_count': sum(1 for i in indicadores_calculados if i['alerta']),
            'resumen_sin_base_count': sum(1 for i in indicadores_calculados if i['sin_base']),
            'resumen_totales': self._resumen_totales(carga, tipo_operacional=tipo_operacional, centro_costo=centro_costo),
        })
        # #246 Sprint C: nómina/gastos/ingresos/clientes reales de las 3
        # fuentes desplegadas (#252/#248/#249/#261), independiente de si hay
        # una CargaFinanciera (libro Excel) cargada para el período -- por
        # eso NO se gatea con `carga` como el resto de este bloque.
        context.update(self._contexto_integrado(anio, mes, proyecto, periodo_valido))
        return context

    @staticmethod
    def _contexto_integrado(anio, mes, proyecto, periodo_valido):
        """Envuelve ``construir_contexto_dashboard_integrado`` (B3) con un
        estado vacío coherente cuando el período todavía no es válido (mes
        fuera de 1-12), para no romper ``calendar.monthrange``."""
        if periodo_valido:
            return construir_contexto_dashboard_integrado(anio, mes, contrato=proyecto)
        vacio_lista = {'con_datos': False, 'cantidad': 0, 'detalle': []}
        return {
            'integracion_nomina': {
                'con_datos': False, 'costo_nomina': ZERO, 'horas_trabajadas': ZERO,
                'cantidad_producciones': 0, 'producciones': [],
            },
            'integracion_gastos': {**vacio_lista, 'total': ZERO, 'pendiente_pago': ZERO},
            'integracion_ingresos': {**vacio_lista, 'total': ZERO, 'cobrado': ZERO},
            'integracion_clientes': {
                'con_datos': False, 'activos_total': 0, 'facturados_periodo': 0, 'detalle': [],
            },
        }

    def post(self, request, *args, **kwargs):
        accion = request.POST.get('accion')
        if accion == 'homologar':
            return self._crear_homologacion()
        if accion == 'editar_homologacion':
            return self._editar_homologacion()
        if accion == 'archivar_homologacion':
            return self._archivar_homologacion()
        if accion == 'previsualizar_tabla_maestra':
            return self._previsualizar_tabla_maestra()
        if accion == 'confirmar_tabla_maestra':
            return self._confirmar_tabla_maestra()
        if accion == 'cancelar_tabla_maestra':
            self.request.session.pop('homologacion_preview', None)
            messages.info(self.request, 'Importación cancelada; no se modificó la tabla maestra.')
            return self._redirect_periodo(self._proyecto_seleccionado(), *self._periodo()[:2])
        if accion == 'restaurar_version_tabla_maestra':
            return self._restaurar_version_tabla_maestra()
        return self._cargar_archivo()

    def _redirect_periodo(self, proyecto=None, anio=None, mes=None, tipo_operacional=None, centro_costo=None):
        params = []
        if proyecto:
            params.append(f'proyecto={proyecto.pk}')
        if anio:
            params.append(f'anio={anio}')
        if mes:
            params.append(f'mes={mes}')
        tipo_operacional = tipo_operacional if tipo_operacional is not None else self._filtros()[0]
        centro_costo = centro_costo if centro_costo is not None else self._filtros()[1]
        if tipo_operacional:
            params.append(f'tipo={tipo_operacional}')
        if centro_costo:
            params.append(f'centro_costo={centro_costo}')
        return redirect(f'{self.request.path}?{"&".join(params)}')

    def _cargar_archivo(self):
        proyecto = self._proyecto_seleccionado()
        anio, mes, periodo_valido = self._periodo()
        form = CargaFinancieraForm(self.request.POST, self.request.FILES)
        if not proyecto:
            messages.error(self.request, 'Seleccione un proyecto activo antes de cargar el archivo.')
        elif not periodo_valido:
            messages.error(self.request, 'Seleccione un mes entre 1 y 12.')
        elif form.is_valid():
            resultado = procesar_carga_financiera(
                form.cleaned_data['archivo'], proyecto=proyecto, anio=anio, mes=mes, usuario=self.request.user,
            )
            if resultado.exito:
                messages.success(self.request, f'Carga procesada: {resultado.resumen["lineas_reales"] + resultado.resumen["lineas_presupuesto"]} líneas.')
                if resultado.resumen['lineas_no_mapeadas']:
                    messages.warning(self.request, f"Advertencia: {resultado.resumen['lineas_no_mapeadas']} líneas sin homologación; el plano las identifica como SIN_HOMOLOGAR.")
            else:
                messages.error(self.request, resultado.error or 'No fue posible procesar el archivo.')
        else:
            messages.error(self.request, form.errors.get('archivo', ['Archivo inválido.'])[0])
        return self._redirect_periodo(proyecto, anio, mes)

    def _crear_homologacion(self):
        datos = {campo: (self.request.POST.get(campo) or '').strip() for campo in ('tipo', 'grupo', 'concepto', 'rubro', 'codigo_contable', 'centro_costo')}
        if not datos['tipo'] or not datos['concepto'] or not datos['codigo_contable']:
            messages.error(self.request, 'Tipo, concepto y código contable son obligatorios para homologar.')
        else:
            obj, creada = HomologacionProjectsContable.objects.update_or_create(
                tipo=datos['tipo'], grupo=datos['grupo'], concepto=datos['concepto'], rubro=datos['rubro'],
                defaults={'codigo_contable': datos['codigo_contable'], 'centro_costo': datos['centro_costo'], 'activo': True},
            )
            messages.success(self.request, 'Homologación creada.' if creada else f'Homologación {obj.concepto} actualizada.')
        return self._redirect_periodo(self._proyecto_seleccionado(), *self._periodo()[:2])

    def _editar_homologacion(self):
        obj = HomologacionProjectsContable.objects.filter(pk=self.request.POST.get('pk'), activo=True).first()
        if not obj:
            messages.error(self.request, 'La homologación solicitada no está vigente.')
        else:
            for campo in ('codigo_contable', 'centro_costo'):
                setattr(obj, campo, (self.request.POST.get(campo) or '').strip())
            if not obj.codigo_contable:
                messages.error(self.request, 'El código contable es obligatorio.')
            else:
                obj.save(update_fields=['codigo_contable', 'centro_costo', 'updated_at'])
                messages.success(self.request, f'Homologación {obj.concepto} actualizada.')
        return self._redirect_periodo(self._proyecto_seleccionado(), *self._periodo()[:2])

    def _archivar_homologacion(self):
        archivada = HomologacionProjectsContable.objects.filter(
            pk=self.request.POST.get('pk'), activo=True
        ).update(activo=False)
        if archivada:
            messages.success(self.request, 'Homologación archivada; el historial y sus referencias se conservan.')
        else:
            messages.error(self.request, 'La homologación solicitada no está vigente.')
        return self._redirect_periodo(self._proyecto_seleccionado(), *self._periodo()[:2])

    def _previsualizar_tabla_maestra(self):
        form = ImportarTablaMaestraForm(self.request.POST, self.request.FILES)
        if not form.is_valid():
            messages.error(self.request, form.errors.get('archivo', ['Archivo inválido.'])[0])
        else:
            preview = previsualizar_tabla_maestra(form.cleaned_data['archivo'])
            self.request.session['homologacion_preview'] = preview
            if preview['errores']:
                messages.error(self.request, f"Preview listo: {len(preview['filas'])} filas válidas y {len(preview['errores'])} errores. No se guardó nada.")
            else:
                messages.success(self.request, f"Preview listo: {len(preview['filas'])} filas válidas. Confirme para reemplazar la tabla maestra.")
        return self._redirect_periodo(self._proyecto_seleccionado(), *self._periodo()[:2])

    def _confirmar_tabla_maestra(self):
        preview = self.request.session.get('homologacion_preview')
        if not preview:
            messages.error(self.request, 'No hay un preview pendiente para confirmar.')
        elif preview['errores']:
            messages.error(self.request, 'Corrija los errores del preview antes de confirmar.')
        elif not preview['filas']:
            messages.error(self.request, 'El preview no contiene filas válidas.')
        else:
            version = confirmar_tabla_maestra(preview['filas'], usuario=self.request.user)
            self.request.session.pop('homologacion_preview', None)
            messages.success(self.request, f'Tabla maestra reemplazada atómicamente en versión {version.numero}.')
        return self._redirect_periodo(self._proyecto_seleccionado(), *self._periodo()[:2])

    def _restaurar_version_tabla_maestra(self):
        version = VersionHomologacionProjectsContable.objects.filter(pk=self.request.POST.get('version')).first()
        if not version:
            messages.error(self.request, 'La versión solicitada no existe.')
        else:
            filas = list(version.homologaciones.values('tipo', 'grupo', 'concepto', 'rubro', 'codigo_contable', 'centro_costo'))
            nueva = confirmar_tabla_maestra(filas, usuario=self.request.user, origen=f'RESTAURACION:{version.numero}')
            messages.success(self.request, f'Versión {version.numero} restaurada como nueva versión {nueva.numero}.')
        return self._redirect_periodo(self._proyecto_seleccionado(), *self._periodo()[:2])


class PlanoFinancieroCsvView(LoginRequiredMixin, RoleRequiredMixin, TemplateView):
    """Entrega el plano contable definitivo, agregado por código y trazable."""

    required_submodulo = SUBMODULO_FIN_HOMOLOGACION
    admin_bypass = False

    def get(self, request, carga_id, *args, **kwargs):
        carga = CargaFinanciera.objects.select_related('proyecto').filter(pk=carga_id).first()
        if not carga:
            raise Http404('La carga solicitada no existe.')
        response = HttpResponse(content_type='text/csv; charset=utf-8')
        nombre = f'Plano_{carga.proyecto.codigo}_{carga.anio}_{carga.mes:02d}.csv'
        response['Content-Disposition'] = f'attachment; filename="{nombre}"'
        response.write('\ufeff')
        writer = csv.writer(response)
        writer.writerow(['Código', 'Concepto', 'Centro', 'Proyecto', 'Mes', 'Valor', 'Referencia'])
        agrupadas = defaultdict(lambda: {'valor': ZERO, 'referencias': set()})
        for linea in carga.lineas.select_related('homologacion'):
            homologacion = linea.homologacion
            llave = (
                homologacion.codigo_contable if homologacion else 'SIN_HOMOLOGAR',
                homologacion.concepto if homologacion else linea.concepto,
                homologacion.centro_costo if homologacion else '',
            )
            agrupadas[llave]['valor'] += linea.valor
            if linea.referencia:
                agrupadas[llave]['referencias'].add(linea.referencia)
        for (codigo, concepto, centro), datos in sorted(agrupadas.items()):
            writer.writerow([codigo, concepto, centro, carga.proyecto.nombre, f'{carga.mes:02d}-{carga.anio}', datos['valor'], ' | '.join(sorted(datos['referencias']))])
        return response


class DescargarVersionHomologacionXlsxView(LoginRequiredMixin, RoleRequiredMixin, TemplateView):
    """Descarga un snapshot, sin reconstruirlo desde el catálogo vigente."""
    required_submodulo = SUBMODULO_FIN_HOMOLOGACION
    admin_bypass = False

    def get(self, request, version_id, *args, **kwargs):
        version = VersionHomologacionProjectsContable.objects.filter(pk=version_id).first()
        if not version:
            raise Http404('La versión solicitada no existe.')
        libro = Workbook()
        libro.remove(libro.active)
        for nombre, filas in {'INGRESOS': version.homologaciones.filter(grupo__icontains='ingreso'), 'GASTOS': version.homologaciones.exclude(grupo__icontains='ingreso')}.items():
            hoja = libro.create_sheet(nombre)
            hoja.append(['Tipo', 'Grupo', 'Concepto', 'Rubro', 'Código contable', 'Centro de costo'])
            for fila in filas.order_by('grupo', 'concepto'):
                hoja.append([fila.tipo, fila.grupo, fila.concepto, fila.rubro, fila.codigo_contable, fila.centro_costo])
        salida = BytesIO()
        libro.save(salida)
        response = HttpResponse(salida.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="Tabla_maestra_version_{version.numero}.xlsx"'
        return response


class ExportarTablaMaestraXlsxView(LoginRequiredMixin, RoleRequiredMixin, TemplateView):
    """Exporta únicamente el catálogo vigente, nunca los registros archivados."""
    required_submodulo = SUBMODULO_FIN_HOMOLOGACION
    admin_bypass = False

    def get(self, request, *args, **kwargs):
        libro = Workbook()
        hoja = libro.active
        hoja.title = 'TABLA_MAESTRA_VIGENTE'
        hoja.append(['Tipo', 'Grupo', 'Concepto', 'Rubro', 'Código contable', 'Centro de costo'])
        for fila in HomologacionProjectsContable.objects.filter(activo=True).order_by('tipo', 'grupo', 'concepto'):
            hoja.append([fila.tipo, fila.grupo, fila.concepto, fila.rubro, fila.codigo_contable, fila.centro_costo])
        salida = BytesIO()
        libro.save(salida)
        response = HttpResponse(salida.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="Tabla_maestra_vigente.xlsx"'
        return response
