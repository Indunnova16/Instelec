import io
from datetime import datetime

import pytest
from openpyxl import Workbook

from apps.contratos.models import Contrato
from apps.financiero.importers_finv2_carga import procesar_carga_financiera
from apps.financiero.models_finv2_carga import CargaFinanciera, LineaCargaFinanciera
from apps.financiero.models_finv2_facturas import Proveedor


def _libro(*, incluir_homologacion=True, neto=100, tercero_movto=None, razon_social=None):
    libro = Workbook()
    real = libro.active
    real.title = "BD Real"
    real.append(
        [
            "Auxiliar",
            "Desc. auxiliar",
            "Neto",
            "Fecha",
            "Docto.",
            "Periodo",
            "Tercero movto",
            "Razón social tercero movto",
            "x",
            "x",
            "x",
            "x",
            "x",
            "x",
            "Cuenta Equiv",
            "CdeC equiv",
        ]
    )
    real.append(
        [
            1,
            "Compra cable",
            neto,
            datetime(2026, 1, 31),
            "DOC-1",
            202601,
            tercero_movto,
            razon_social,
            None,
            None,
            None,
            None,
            None,
            None,
            "Materiales",
            "TRANSELCA",
        ]
    )
    real.append(
        [
            2,
            "Otro periodo",
            999,
            datetime(2026, 2, 1),
            "DOC-2",
            202602,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            "Materiales",
            "TRANSELCA",
        ]
    )
    ppto = libro.create_sheet("BD Ppto")
    ppto.append(["Tipo", "Proyecto", "Rubro", "Clasificacion", "Valor", "mes", "año", "ciudad"])
    ppto.append(["Presupuesto", "Transelca", "Materiales", "Costos", 120, 1, 2026, "Barranquilla"])
    if incluir_homologacion:
        homologacion = libro.create_sheet("Homologacion")
        homologacion.append(["tipo", "Grupo", "Concepto", "Rubro"])
        homologacion.append(["REAL", "Materiales", "Compra cable", "TRANSELCA"])
    salida = io.BytesIO()
    libro.save(salida)
    salida.seek(0)
    salida.name = "TRANSELCA.xlsx"
    return salida


@pytest.fixture
def proyecto(db):
    return Contrato.objects.create(
        codigo="TRAN-246", nombre="Transelca", unidad_negocio="MANTENIMIENTO"
    )


@pytest.mark.django_db
def test_carga_transelca_versiona_periodo_sin_borrar_historico(proyecto):
    primero = procesar_carga_financiera(
        _libro(neto=100), proyecto=proyecto, anio=2026, mes=1, usuario=None
    )
    assert primero.exito
    assert LineaCargaFinanciera.objects.filter(carga=primero.carga).count() == 2

    segundo = procesar_carga_financiera(
        _libro(neto=250), proyecto=proyecto, anio=2026, mes=1, usuario=None
    )
    assert segundo.exito
    cargas = CargaFinanciera.objects.filter(proyecto=proyecto, anio=2026, mes=1)
    assert cargas.count() == 2
    assert cargas.get(pk=primero.carga.pk).vigente is False
    assert segundo.carga.vigente is True
    assert segundo.carga.version == 2
    lineas = LineaCargaFinanciera.objects.filter(carga=segundo.carga)
    assert lineas.count() == 2
    assert lineas.get(tipo="REAL").valor == 250
    assert segundo.resumen["lineas_homologacion_origen"] == 1


@pytest.mark.django_db
def test_carga_rechaza_hoja_faltante(proyecto):
    resultado = procesar_carga_financiera(
        _libro(incluir_homologacion=False), proyecto=proyecto, anio=2026, mes=1, usuario=None
    )
    assert not resultado.exito
    assert resultado.carga is None
    assert "Homologacion" in resultado.error
    assert CargaFinanciera.objects.count() == 0


@pytest.mark.django_db
def test_carga_rechaza_periodo_sin_lineas(proyecto):
    resultado = procesar_carga_financiera(
        _libro(), proyecto=proyecto, anio=2026, mes=3, usuario=None
    )
    assert not resultado.exito
    assert "no contiene líneas" in resultado.error


@pytest.mark.django_db
def test_carga_resuelve_proveedor_por_nit_tercero_movto(proyecto):
    """Hallazgo #248 (2026-09-23): 'Tercero movto.' (NIT) nunca se leía -- el
    FK `proveedor` de LineaCargaFinanciera (#268/A1) quedaba SIEMPRE None."""
    Proveedor.objects.create(nombre="Sanmartin Blanco Pedro Antonio", nit="98687561")
    resultado = procesar_carga_financiera(
        _libro(tercero_movto="98687561", razon_social="Sanmartin Blanco Pedro Antonio"),
        proyecto=proyecto,
        anio=2026,
        mes=1,
        usuario=None,
    )
    assert resultado.exito
    linea_con_tercero = LineaCargaFinanciera.objects.get(
        carga=resultado.carga, tipo="REAL", referencia="DOC-1"
    )
    assert linea_con_tercero.proveedor is not None
    assert linea_con_tercero.proveedor.nit == "98687561"
    assert linea_con_tercero.datos_origen["tercero_movto"] == "98687561"


@pytest.mark.django_db
def test_carga_nit_sin_proveedor_registrado_no_falla(proyecto):
    """Un NIT presente en el Excel pero sin Proveedor creado en #262 no debe
    romper la carga -- solo queda sin resolver (mismo comportamiento que hoy
    para cualquier centro de costo no homologado)."""
    resultado = procesar_carga_financiera(
        _libro(tercero_movto="999999999-NO-EXISTE"),
        proyecto=proyecto,
        anio=2026,
        mes=1,
        usuario=None,
    )
    assert resultado.exito
    linea = LineaCargaFinanciera.objects.get(carga=resultado.carga, tipo="REAL", referencia="DOC-1")
    assert linea.proveedor is None
