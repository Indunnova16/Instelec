import io
from datetime import datetime

import pytest
from openpyxl import Workbook

from apps.contratos.models import Contrato
from apps.financiero.models_finv2_carga import CargaFinanciera, HomologacionProjectsContable


def _libro():
    libro = Workbook()
    real = libro.active
    real.title = 'BD Real'
    real.append(['Auxiliar', 'Desc. auxiliar', 'Neto', 'Fecha', 'Docto.', 'Periodo', 'x', 'x', 'x', 'x', 'x', 'x', 'x', 'x', 'Cuenta Equiv', 'CdeC equiv'])
    real.append([1, 'Compra cable', 250, datetime(2026, 1, 31), 'DOC-1', 202601, None, None, None, None, None, None, None, None, 'Materiales', 'CC-01'])
    presupuesto = libro.create_sheet('BD Ppto')
    presupuesto.append(['Tipo', 'Proyecto', 'Rubro', 'Clasificacion', 'Valor', 'mes', 'año', 'ciudad'])
    presupuesto.append(['Presupuesto', 'Transelca', 'Materiales', 'Costos', 120, 1, 2026, 'Barranquilla'])
    homologacion = libro.create_sheet('Homologacion')
    homologacion.append(['tipo', 'Grupo', 'Concepto', 'Rubro'])
    homologacion.append(['REAL', 'Materiales', 'Compra cable', 'CC-01'])
    salida = io.BytesIO()
    libro.save(salida)
    salida.seek(0)
    salida.name = 'TRANSELCA.xlsx'
    return salida


@pytest.fixture
def proyecto(db):
    return Contrato.objects.create(codigo='TRAN-VIEW', nombre='Transelca', unidad_negocio='MANTENIMIENTO')


@pytest.mark.django_db
def test_carga_muestra_indicadores(client, admin_user, proyecto):
    client.force_login(admin_user)
    respuesta = client.post('/financiero/carga-financiera/', {
        'proyecto': proyecto.pk, 'anio': 2026, 'mes': 1, 'archivo': _libro(),
    }, follow=True)
    assert respuesta.status_code == 200
    assert CargaFinanciera.objects.filter(proyecto=proyecto, anio=2026, mes=1).exists()
    contenido = respuesta.content.decode()
    assert 'Facturación vs meta' in contenido
    assert 'Margen bruto' in contenido
    assert 'Ratio costos/facturación' in contenido
    assert 'Variación mes a mes' in contenido


@pytest.mark.django_db
def test_plano_csv_homologado(client, admin_user, proyecto):
    client.force_login(admin_user)
    HomologacionProjectsContable.objects.create(
        tipo='REAL', grupo='Materiales', concepto='Compra cable', rubro='CC-01', codigo_contable='5130', centro_costo='CC-001',
    )
    client.post('/financiero/carga-financiera/', {
        'proyecto': proyecto.pk, 'anio': 2026, 'mes': 1, 'archivo': _libro(),
    })
    carga = CargaFinanciera.objects.get(proyecto=proyecto, anio=2026, mes=1)
    respuesta = client.get(f'/financiero/carga-financiera/{carga.pk}/plano.csv')
    assert respuesta.status_code == 200
    assert respuesta['Content-Type'].startswith('text/csv')
    assert 'Plano_TRAN-VIEW_2026_01.csv' in respuesta['Content-Disposition']
    contenido = respuesta.content.decode('utf-8-sig')
    assert 'Código,Concepto,Centro,Proyecto,Mes,Valor,Referencia' in contenido
    assert '5130,Compra cable,CC-001,Transelca,01-2026,250.00,DOC-1' in contenido


@pytest.mark.django_db
def test_carga_rechaza_proyecto_inactivo(client, admin_user, proyecto):
    client.force_login(admin_user)
    proyecto.estado = Contrato.Estado.FINALIZADO
    proyecto.save(update_fields=['estado'])
    respuesta = client.post('/financiero/carga-financiera/', {
        'proyecto': proyecto.pk, 'anio': 2026, 'mes': 1, 'archivo': _libro(),
    }, follow=True)
    assert respuesta.status_code == 200
    assert not CargaFinanciera.objects.exists()
    assert 'Seleccione un proyecto activo' in respuesta.content.decode()


@pytest.mark.django_db
def test_carga_rechaza_mes_invalido(client, admin_user, proyecto):
    client.force_login(admin_user)
    respuesta = client.post('/financiero/carga-financiera/', {
        'proyecto': proyecto.pk, 'anio': 2026, 'mes': 13, 'archivo': _libro(),
    }, follow=True)
    assert respuesta.status_code == 200
    assert not CargaFinanciera.objects.exists()
    assert 'mes entre 1 y 12' in respuesta.content.decode()
