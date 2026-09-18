"""E2E de #246 Sprint D: reportes PDF/Excel/PPT + RBAC de descarga (B4).

Cubre (protocolo de issues, tests_e2e del BLUEPRINT):
- `descarga_pdf_excel_ppt_usuario_gerente_bytes_no_vacios`
- `supervisor_sin_boton_descarga_y_403_url_directa`
más la matriz completa (Contador solo Excel, Coordinador sin descarga) y los
2 edge cases exigidos por F3: usuario sin permiso (403) y período sin datos
de alguna fuente (el reporte igual se genera, sin crashear).

Nota `--nomigrations` (pyproject.toml, addopts globales de pytest): las
migraciones de datos (`core/migrations/0008_..._248.py`, que siembra los
roles `contador`/`gerente_financiero` en prod) NO corren en este modo de
test -- se siembran los `Role`/`RoleModuloPermiso` necesarios explícitamente
en los fixtures de abajo, mismo patrón que
`test_issue_249_v2.py::_rol_consulta`.
"""

from decimal import Decimal
from io import BytesIO
from zipfile import ZipFile

import pytest
from django.core.cache import cache
from openpyxl import load_workbook

from apps.contratos.models import Contrato
from apps.core.models_roles import Role, RoleModuloPermiso
from apps.financiero.models_finv2_carga import CargaFinanciera, LineaCargaFinanciera

FIN_HOMOLOGACION = "FIN_HOMOLOGACION"
MANTENIMIENTO = "MANTENIMIENTO"


def _asegurar_rol(codigo, nombre, nivel_acceso):
    """Crea (si falta) el `Role` + su `RoleModuloPermiso` sobre
    FIN_HOMOLOGACION -- necesario porque `--nomigrations` no corre el seed
    real de S1/legacy. `nivel_acceso=None` dej sin fila -> SIN_ACCESO."""
    role, _ = Role.objects.get_or_create(
        codigo=codigo, defaults={"nombre": nombre, "nivel": "operario"}
    )
    if nivel_acceso:
        RoleModuloPermiso.objects.update_or_create(
            role=role,
            modulo=MANTENIMIENTO,
            submodulo=FIN_HOMOLOGACION,
            defaults={"nivel_acceso": nivel_acceso},
        )
    cache.clear()
    return role


@pytest.fixture
def contrato(db):
    return Contrato.objects.create(
        codigo="RPT246-1",
        nombre="Proyecto reportes #246",
        unidad_negocio="MANTENIMIENTO",
    )


@pytest.fixture
def carga_con_datos(contrato):
    carga = CargaFinanciera.objects.create(proyecto=contrato, anio=2026, mes=9, vigente=True)
    LineaCargaFinanciera.objects.create(
        carga=carga,
        tipo=LineaCargaFinanciera.Tipo.REAL,
        concepto="Materiales",
        grupo="Materiales",
        valor=Decimal("800.00"),
        periodo=202609,
    )
    LineaCargaFinanciera.objects.create(
        carga=carga,
        tipo=LineaCargaFinanciera.Tipo.PRESUPUESTO,
        concepto="Materiales",
        grupo="Materiales",
        valor=Decimal("1000.00"),
        periodo=202609,
    )
    return carga


def _usuario(django_user_model, email, rol):
    return django_user_model.objects.create_user(email=email, password="x", rol=rol)


def _login(client, django_user_model, email, rol_codigo, rol_nombre, nivel_acceso):
    _asegurar_rol(rol_codigo, rol_nombre, nivel_acceso)
    user = _usuario(django_user_model, email, rol_codigo)
    client.force_login(user)
    return user


# ---------------------------------------------------------------------------
# descarga_pdf_excel_ppt_usuario_gerente_bytes_no_vacios
# ---------------------------------------------------------------------------


@pytest.mark.django_db
def test_gerente_financiero_descarga_pdf_bytes_no_vacios(
    client, django_user_model, contrato, carga_con_datos
):
    _login(
        client,
        django_user_model,
        "gerente-246@example.test",
        "gerente_financiero",
        "Gerente Financiero",
        "ver_editar",
    )
    respuesta = client.get(
        "/financiero/carga-financiera/exportar/pdf/",
        {"proyecto": contrato.pk, "anio": 2026, "mes": 9},
    )
    assert respuesta.status_code == 200
    assert respuesta["Content-Type"] == "application/pdf"
    contenido = respuesta.content
    assert len(contenido) > 0
    assert contenido.startswith(b"%PDF")


@pytest.mark.django_db
def test_gerente_financiero_descarga_excel_bytes_no_vacios_y_parseable(
    client, django_user_model, contrato, carga_con_datos
):
    _login(
        client,
        django_user_model,
        "gerente-246-xlsx@example.test",
        "gerente_financiero",
        "Gerente Financiero",
        "ver_editar",
    )
    respuesta = client.get(
        "/financiero/carga-financiera/exportar/excel/",
        {"proyecto": contrato.pk, "anio": 2026, "mes": 9},
    )
    assert respuesta.status_code == 200
    assert (
        respuesta["Content-Type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    contenido = respuesta.content
    assert len(contenido) > 0
    libro = load_workbook(BytesIO(contenido))
    assert libro.sheetnames == ["Resumen", "Facturación", "Costos", "Proyectos", "Gráficos"]


@pytest.mark.django_db
def test_gerente_financiero_descarga_ppt_bytes_no_vacios_y_estructura_valida(
    client, django_user_model, contrato, carga_con_datos
):
    _login(
        client,
        django_user_model,
        "gerente-246-ppt@example.test",
        "gerente_financiero",
        "Gerente Financiero",
        "ver_editar",
    )
    respuesta = client.get(
        "/financiero/carga-financiera/exportar/ppt/",
        {"proyecto": contrato.pk, "anio": 2026, "mes": 9},
    )
    assert respuesta.status_code == 200
    assert (
        respuesta["Content-Type"]
        == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )
    contenido = respuesta.content
    assert len(contenido) > 0
    zf = ZipFile(BytesIO(contenido))
    nombres = zf.namelist()
    assert "ppt/presentation.xml" in nombres
    assert "[Content_Types].xml" in nombres
    slides = [
        n
        for n in nombres
        if n.startswith("ppt/slides/slide") and n.endswith(".xml") and "_rels" not in n
    ]
    assert len(slides) == 9  # portada + resumen + 6 indicadores + conclusiones


@pytest.mark.django_db
def test_admin_legacy_tambien_descarga_los_3_formatos(
    client, django_user_model, contrato, carga_con_datos
):
    """`admin` (rol legacy, no solo `gerente_financiero`) también está en la
    matriz 'Gerente/Admin todos los reportes' del issue."""
    _login(
        client,
        django_user_model,
        "admin-246@example.test",
        "admin",
        "Administrador (legacy)",
        "ver_editar",
    )
    for url in (
        "/financiero/carga-financiera/exportar/pdf/",
        "/financiero/carga-financiera/exportar/excel/",
        "/financiero/carga-financiera/exportar/ppt/",
    ):
        respuesta = client.get(url, {"proyecto": contrato.pk, "anio": 2026, "mes": 9})
        assert respuesta.status_code == 200, url


# ---------------------------------------------------------------------------
# supervisor_sin_boton_descarga_y_403_url_directa
# ---------------------------------------------------------------------------


@pytest.mark.django_db
def test_supervisor_403_en_los_3_endpoints_de_descarga(
    client, django_user_model, contrato, carga_con_datos
):
    """Supervisor tiene `ver` (lectura) sobre FIN_HOMOLOGACION -- ve el
    dashboard -- pero la matriz de descarga de B4 lo excluye explícitamente:
    un GET directo a cualquiera de los 3 endpoints debe responder 403."""
    _login(
        client,
        django_user_model,
        "supervisor-246@example.test",
        "supervisor",
        "Supervisor de Cuadrilla (legacy)",
        "ver",
    )
    for url in (
        "/financiero/carga-financiera/exportar/pdf/",
        "/financiero/carga-financiera/exportar/excel/",
        "/financiero/carga-financiera/exportar/ppt/",
    ):
        respuesta = client.get(url, {"proyecto": contrato.pk, "anio": 2026, "mes": 9})
        assert respuesta.status_code == 403, url


@pytest.mark.django_db
def test_supervisor_no_ve_los_botones_de_descarga_en_el_dashboard(
    client, django_user_model, contrato, carga_con_datos
):
    _login(
        client,
        django_user_model,
        "supervisor-246-ui@example.test",
        "supervisor",
        "Supervisor de Cuadrilla (legacy)",
        "ver",
    )
    respuesta = client.get(
        "/financiero/carga-financiera/",
        {"proyecto": contrato.pk, "anio": 2026, "mes": 9},
    )
    assert respuesta.status_code == 200
    html = respuesta.content.decode()
    assert 'id="btn-exportar-pdf"' not in html
    assert 'id="btn-exportar-excel"' not in html
    assert 'id="btn-exportar-ppt"' not in html
    assert "no tiene autorización para descargar" in html.lower()


@pytest.mark.django_db
def test_gerente_financiero_ve_los_3_botones_en_el_dashboard(
    client, django_user_model, contrato, carga_con_datos
):
    _login(
        client,
        django_user_model,
        "gerente-246-ui@example.test",
        "gerente_financiero",
        "Gerente Financiero",
        "ver_editar",
    )
    respuesta = client.get(
        "/financiero/carga-financiera/",
        {"proyecto": contrato.pk, "anio": 2026, "mes": 9},
    )
    assert respuesta.status_code == 200
    html = respuesta.content.decode()
    assert 'id="btn-exportar-pdf"' in html
    assert 'id="btn-exportar-excel"' in html
    assert 'id="btn-exportar-ppt"' in html


# ---------------------------------------------------------------------------
# Matriz completa: Coordinador sin descarga, Contador solo Excel
# ---------------------------------------------------------------------------


@pytest.mark.django_db
def test_coordinador_lectura_sin_descarga(client, django_user_model, contrato, carga_con_datos):
    """Coordinador tiene `ver_editar` heredado del seed legacy (#247) sobre
    FIN_HOMOLOGACION -- edita homologaciones -- pero la matriz de descarga de
    negocio (#246 Sprint D) lo excluye igual que a Supervisor: 403 en los 3
    formatos aunque su nivel de acceso genérico sea más alto."""
    _login(
        client,
        django_user_model,
        "coordinador-246@example.test",
        "coordinador",
        "Coordinador (legacy)",
        "ver_editar",
    )
    for url in (
        "/financiero/carga-financiera/exportar/pdf/",
        "/financiero/carga-financiera/exportar/excel/",
        "/financiero/carga-financiera/exportar/ppt/",
    ):
        respuesta = client.get(url, {"proyecto": contrato.pk, "anio": 2026, "mes": 9})
        assert respuesta.status_code == 403, url


@pytest.mark.django_db
def test_contador_solo_descarga_excel_pdf_y_ppt_403(
    client, django_user_model, contrato, carga_con_datos
):
    _login(
        client, django_user_model, "contador-246@example.test", "contador", "Contador", "ver_editar"
    )
    respuesta_excel = client.get(
        "/financiero/carga-financiera/exportar/excel/",
        {"proyecto": contrato.pk, "anio": 2026, "mes": 9},
    )
    assert respuesta_excel.status_code == 200

    for url in (
        "/financiero/carga-financiera/exportar/pdf/",
        "/financiero/carga-financiera/exportar/ppt/",
    ):
        respuesta = client.get(url, {"proyecto": contrato.pk, "anio": 2026, "mes": 9})
        assert respuesta.status_code == 403, url


@pytest.mark.django_db
def test_contador_ve_solo_el_boton_de_excel_en_el_dashboard(
    client, django_user_model, contrato, carga_con_datos
):
    _login(
        client,
        django_user_model,
        "contador-246-ui@example.test",
        "contador",
        "Contador",
        "ver_editar",
    )
    respuesta = client.get(
        "/financiero/carga-financiera/",
        {"proyecto": contrato.pk, "anio": 2026, "mes": 9},
    )
    html = respuesta.content.decode()
    assert 'id="btn-exportar-excel"' in html
    assert 'id="btn-exportar-pdf"' not in html
    assert 'id="btn-exportar-ppt"' not in html


# ---------------------------------------------------------------------------
# Edge cases obligatorios de F3
# ---------------------------------------------------------------------------


@pytest.mark.django_db
def test_usuario_sin_ningun_rol_de_descarga_403(
    client, django_user_model, contrato, carga_con_datos
):
    """Un rol no listado en la matriz (deny-by-default) -- p.ej. un rol
    operativo nuevo sin fila en `FORMATOS_DESCARGA_POR_ROL` -- no descarga
    nada aunque tenga `ver_editar` en FIN_HOMOLOGACION."""
    _login(
        client,
        django_user_model,
        "otro-rol-246@example.test",
        "director",
        "Director de Proyecto (legacy)",
        "ver_editar",
    )
    respuesta = client.get(
        "/financiero/carga-financiera/exportar/pdf/",
        {"proyecto": contrato.pk, "anio": 2026, "mes": 9},
    )
    assert respuesta.status_code == 403


@pytest.mark.django_db
def test_usuario_sin_acceso_al_submodulo_403(client, django_user_model, contrato, carga_con_datos):
    """Sin ninguna fila RoleModuloPermiso para FIN_HOMOLOGACION -> SIN_ACCESO
    -- 403 aunque el rol esté (hipotéticamente) en la matriz de formatos."""
    _login(
        client,
        django_user_model,
        "sin-acceso-246@example.test",
        "gerente_financiero",
        "Gerente Financiero",
        None,
    )
    respuesta = client.get(
        "/financiero/carga-financiera/exportar/pdf/",
        {"proyecto": contrato.pk, "anio": 2026, "mes": 9},
    )
    assert respuesta.status_code == 403


@pytest.mark.django_db
def test_periodo_sin_carga_financiera_igual_genera_los_3_reportes_sin_crashear(
    client, django_user_model, contrato
):
    """Edge case obligatorio: período/proyecto sin `CargaFinanciera` (ninguna
    fuente cargada para ese mes) -- los 3 generadores deben producir un
    archivo válido en vez de lanzar, con los indicadores/tablas vacíos."""
    _login(
        client,
        django_user_model,
        "gerente-246-vacio@example.test",
        "gerente_financiero",
        "Gerente Financiero",
        "ver_editar",
    )
    for url, content_type in (
        ("/financiero/carga-financiera/exportar/pdf/", "application/pdf"),
        (
            "/financiero/carga-financiera/exportar/excel/",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        ),
        (
            "/financiero/carga-financiera/exportar/ppt/",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        ),
    ):
        respuesta = client.get(url, {"proyecto": contrato.pk, "anio": 2026, "mes": 3})
        assert respuesta.status_code == 200, url
        assert respuesta["Content-Type"] == content_type
        assert len(respuesta.content) > 0


@pytest.mark.django_db
def test_sin_proyecto_ni_periodo_seleccionado_no_crashea(client, django_user_model):
    """Edge case: `?proyecto=` ausente y período por default (hoy) -- sin
    ningún filtro activo, `proyecto=None` agrega el portafolio completo
    (mismo criterio que B3) y el reporte se genera igual."""
    _login(
        client,
        django_user_model,
        "gerente-246-sinfiltro@example.test",
        "gerente_financiero",
        "Gerente Financiero",
        "ver_editar",
    )
    respuesta = client.get("/financiero/carga-financiera/exportar/excel/")
    assert respuesta.status_code == 200
    assert len(respuesta.content) > 0


# ---------------------------------------------------------------------------
# Filtros activos se respetan en la descarga (tipo/centro de costo/proyecto)
# ---------------------------------------------------------------------------


@pytest.mark.django_db
def test_descarga_respeta_filtro_tipo_operacional_activo(
    client, django_user_model, contrato, carga_con_datos
):
    """Agrega una 2da línea REAL con otro `tipo_operacional` y confirma que
    el Excel, filtrado por `tipo=Construcción`, sólo refleja esa porción --
    igual que ya hace el dashboard HTML (`_base_lineas`)."""
    LineaCargaFinanciera.objects.create(
        carga=carga_con_datos,
        tipo=LineaCargaFinanciera.Tipo.REAL,
        concepto="Nómina",
        grupo="Nómina",
        tipo_operacional="Mantenimiento",
        valor=Decimal("300.00"),
        periodo=202609,
    )
    # Backfill el tipo_operacional de las líneas de materiales para que el
    # filtro tenga 2 valores distintos entre los que discriminar.
    LineaCargaFinanciera.objects.filter(carga=carga_con_datos, concepto="Materiales").update(
        tipo_operacional="Construcción"
    )
    _login(
        client,
        django_user_model,
        "gerente-246-filtro@example.test",
        "gerente_financiero",
        "Gerente Financiero",
        "ver_editar",
    )

    respuesta_filtrada = client.get(
        "/financiero/carga-financiera/exportar/excel/",
        {"proyecto": contrato.pk, "anio": 2026, "mes": 9, "tipo": "Construcción"},
    )
    assert respuesta_filtrada.status_code == 200
    libro = load_workbook(BytesIO(respuesta_filtrada.content))
    hoja_resumen = libro["Resumen"]
    costo_real_filtrado = next(
        fila[1].value for fila in hoja_resumen.iter_rows() if fila[0].value == "Costo real"
    )
    # Filtrado por Construcción sólo debe contar la línea de Materiales
    # (800.00 REAL), no los 300.00 REAL de Mantenimiento.
    assert costo_real_filtrado == 800.00
