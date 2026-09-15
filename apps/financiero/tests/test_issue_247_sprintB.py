import io

import pytest
from openpyxl import Workbook, load_workbook

from apps.financiero.importers_finv2_carga import confirmar_tabla_maestra, previsualizar_tabla_maestra
from apps.financiero.models_finv2_carga import HomologacionProjectsContable, VersionHomologacionProjectsContable


def _tabla(*, invalida=False, codigo_gasto='6100'):
    libro = Workbook()
    ingresos = libro.active
    ingresos.title = 'INGRESOS'
    gastos = libro.create_sheet('GASTOS')
    headers = ['Tipo', 'Grupo', 'Concepto', 'Rubro', 'Código contable', 'Centro de costo']
    ingresos.append(headers)
    ingresos.append(['Fijo', 'Ingresos QA_E2E_247', 'QA_E2E_247 Venta', 'Servicios', '5100', 'CC-01'])
    gastos.append(headers)
    gastos.append(['Variable' if not invalida else 'Otro', 'Gastos QA_E2E_247', 'QA_E2E_247 Insumo', 'Materiales', codigo_gasto, 'CC-02'])
    salida = io.BytesIO()
    libro.save(salida)
    salida.seek(0)
    salida.name = 'QA_E2E_247_tabla_maestra.xlsx'
    return salida


@pytest.mark.django_db
def test_247_b_preview_no_persiste_y_valida_tipo_rango_y_duplicado():
    preview = previsualizar_tabla_maestra(_tabla(invalida=True, codigo_gasto='5100'))
    assert len(preview['filas']) == 1
    assert any('Fijo o Variable' in error for error in preview['errores'])
    assert any('duplicado' in error for error in preview['errores'])
    assert not HomologacionProjectsContable.objects.exists()


@pytest.mark.django_db
def test_247_b_confirmar_versiona_y_preserva_homologacion_legacy(admin_user):
    legacy = HomologacionProjectsContable.objects.create(
        tipo='REAL', grupo='Legacy', concepto='Dato legacy', rubro='R', codigo_contable='5001',
    )
    preview = previsualizar_tabla_maestra(_tabla())
    version_1 = confirmar_tabla_maestra(preview['filas'], usuario=admin_user)
    assert version_1.numero == 1
    assert version_1.homologaciones.count() == 2
    legacy.refresh_from_db()
    assert legacy.pk and not legacy.activo  # no se borra: FKs históricas continúan válidas.

    filas = list(version_1.homologaciones.values('tipo', 'grupo', 'concepto', 'rubro', 'codigo_contable', 'centro_costo'))
    version_2 = confirmar_tabla_maestra(filas, usuario=admin_user, origen='RESTAURACION:1')
    assert version_2.numero == 2
    assert version_2.homologaciones.count() == version_1.homologaciones.count()
    assert VersionHomologacionProjectsContable.objects.count() == 2
    with pytest.raises(ValueError, match='inmutables'):
        version_1.origen = 'ALTERADA'
        version_1.save()


@pytest.mark.django_db
def test_247_b_preview_confirmacion_historial_y_descarga(client, admin_user):
    client.force_login(admin_user)
    respuesta = client.post('/financiero/carga-financiera/', {
        'accion': 'previsualizar_tabla_maestra', 'archivo': _tabla(),
    }, follow=True)
    assert 'Preview de importación' in respuesta.content.decode()
    assert not HomologacionProjectsContable.objects.exists()
    respuesta = client.post('/financiero/carga-financiera/', {'accion': 'confirmar_tabla_maestra'}, follow=True)
    assert 'versión 1' in respuesta.content.decode()
    version = VersionHomologacionProjectsContable.objects.get(numero=1)
    descarga = client.get(f'/financiero/carga-financiera/homologacion/version/{version.pk}/xlsx/')
    assert descarga.status_code == 200
    assert descarga['Content-Type'].startswith('application/vnd.openxmlformats')
    libro = load_workbook(io.BytesIO(descarga.content), read_only=True)
    assert set(libro.sheetnames) == {'INGRESOS', 'GASTOS'}
