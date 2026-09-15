"""Regresiones Sprint A #246 con el libro TRANSELCA real."""
import io
import os
from pathlib import Path

import pytest
from openpyxl import Workbook

from apps.contratos.models import Contrato
from apps.financiero.importers_finv2_carga import procesar_carga_financiera
from apps.financiero.models_finv2_carga import CargaFinanciera, LineaCargaFinanciera


ARCHIVO_REAL = os.environ.get('TRANSELCA_REAL_XLSX')


def _libro_con_ramas():
    libro = Workbook()
    real = libro.active
    real.title = 'BD Real'
    real.append(['Desc. auxiliar', 'Neto', 'Periodo', 'Cuenta Equiv', 'CdeC equiv', 'C.Costo', 'Desc. C.O. movto.'])
    real.append(['Cable construcción', 100, 202503, 'Materiales', 'TRANSELCA', '400102', 'CONSTRUCCION LINEA'])
    real.append(['Cable mantenimiento', 200, 202503, 'Materiales', 'TRANSELCA', '400103', 'MANTENIMIENTO LINEA'])
    ppto = libro.create_sheet('BD Ppto')
    ppto.append(['Tipo', 'Proyecto', 'Rubro', 'Clasificacion', 'Valor', 'mes', 'año'])
    ppto.append(['Presupuesto', 'Transelca', 'Materiales', 'Costos', 90, 3, 2025])
    homologacion = libro.create_sheet('Homologacion')
    homologacion.append(['tipo', 'Grupo', 'Concepto', 'Rubro'])
    salida = io.BytesIO()
    libro.save(salida)
    salida.seek(0)
    salida.name = 'ramas.xlsx'
    return salida


@pytest.fixture
def proyecto(db):
    return Contrato.objects.create(codigo='TRAN-246-A', nombre='Transelca', unidad_negocio='MANTENIMIENTO')


@pytest.mark.django_db
def test_workbook_real_legacy_se_escanea_completo_y_persiste_procedencia(proyecto):
    """El período 202503 obliga a recorrer las 14.132 filas reales del adjunto."""
    if not ARCHIVO_REAL:
        pytest.skip('Defina TRANSELCA_REAL_XLSX con el adjunto TRANSELCA para la prueba de escala.')
    with Path(ARCHIVO_REAL).open('rb') as archivo:
        resultado = procesar_carga_financiera(archivo, proyecto=proyecto, anio=2025, mes=3, usuario=None)

    assert resultado.exito, resultado.error
    assert resultado.resumen['lineas_reales'] == 62
    assert resultado.resumen['lineas_presupuesto'] == 0
    assert resultado.resumen['lineas_homologacion_origen'] == 70
    assert resultado.resumen['duracion_validacion_segundos'] < 30
    assert resultado.resumen['legacy_sin_tipo'] is True
    linea = resultado.carga.lineas.get(fila_origen=2)
    assert linea.periodo == 202503
    assert linea.tipo_operacional == LineaCargaFinanciera.TipoOperacional.MANTENIMIENTO
    assert linea.cdec_equiv == 'TRANSELCA'
    assert linea.centro_costo == '400102'
    assert linea.datos_origen['tipo_operacional']['fila_fuente'] == 2


@pytest.mark.django_db
def test_filtros_tipo_y_centro_son_encadenados_y_preservados(client, admin_user, proyecto):
    resultado = procesar_carga_financiera(_libro_con_ramas(), proyecto=proyecto, anio=2025, mes=3, usuario=admin_user)
    assert resultado.exito, resultado.error
    client.force_login(admin_user)

    respuesta = client.get(
        f'/financiero/carga-financiera/?proyecto={proyecto.pk}&anio=2025&mes=3'
        '&tipo=Construcci%C3%B3n&centro_costo=400102'
    )
    assert respuesta.status_code == 200
    contenido = respuesta.content.decode()
    assert 'name="tipo"' in contenido
    assert 'sin_clasificar' not in contenido  # el libro sintético sólo tiene las dos ramas explícitas
    assert '400102' in contenido
    assert '400103' not in contenido

    respuesta = client.post('/financiero/carga-financiera/', {
        'proyecto': proyecto.pk, 'anio': 2025, 'mes': 3, 'tipo': 'Construcción',
        'centro_costo': '400102', 'archivo': _libro_con_ramas(),
    })
    assert 'tipo=Construcci%C3%B3n' in respuesta.url
    assert 'centro_costo=400102' in respuesta.url
    assert CargaFinanciera.objects.filter(proyecto=proyecto, anio=2025, mes=3).count() == 2
