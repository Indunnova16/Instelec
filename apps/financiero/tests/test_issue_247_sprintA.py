import io
from pathlib import Path

import pytest
from openpyxl import load_workbook

from apps.contratos.models import Contrato
from apps.financiero.importers_finv2_carga import (
    normalizar_headers,
    procesar_carga_financiera,
)
from apps.financiero.models_finv2_carga import LineaCargaFinanciera


FIXTURE_REAL = Path(__file__).resolve().parents[3] / 'fixtures' / 'Instelec' / 'QA_E2E_247_transelca_real_80_filas.xlsx'


@pytest.fixture
def proyecto_247(db):
    return Contrato.objects.create(codigo='QA-E2E-247', nombre='Transelca', unidad_negocio='MANTENIMIENTO')


def _archivo_real():
    archivo = io.BytesIO(FIXTURE_REAL.read_bytes())
    archivo.name = FIXTURE_REAL.name
    return archivo


@pytest.mark.django_db
def test_muestra_real_80_filas_carga_dos_periodos_y_tipo_trazable(proyecto_247):
    for mes in (6, 7):
        resultado = procesar_carga_financiera(
            _archivo_real(), proyecto=proyecto_247, anio=2026, mes=mes, usuario=None,
        )
        assert resultado.exito, resultado.error
        assert resultado.resumen['lineas_homologacion_origen'] == 70
        assert resultado.resumen['lineas_reales'] == 2
        assert resultado.resumen['lineas_presupuesto'] == 3

        real = resultado.carga.lineas.filter(tipo=LineaCargaFinanciera.Tipo.REAL).first()
        assert real.tipo_operacional == LineaCargaFinanciera.TipoOperacional.MANTENIMIENTO
        assert real.datos_origen['tipo_operacional'] == {
            'valor_fuente': 'MANTENIMIENTO LINEA TRANSELCA',
            'columna_fuente': 'Desc. C.O. movto.',
            'hoja_fuente': 'BD Real',
            'fila_fuente': real.fila_origen,
            'regla': 'descripcion_operacional',
        }
        presupuesto = resultado.carga.lineas.filter(tipo=LineaCargaFinanciera.Tipo.PRESUPUESTO).first()
        assert presupuesto.tipo_operacional == 'Presupuesto'
        assert presupuesto.datos_origen['tipo_operacional']['columna_fuente'] == 'Tipo'
        assert presupuesto.datos_origen['tipo_operacional']['fila_fuente'] == presupuesto.fila_origen


@pytest.mark.django_db
def test_headers_reales_aceptan_aliases_y_acentos(proyecto_247):
    libro = load_workbook(FIXTURE_REAL)
    real = libro['BD Real']
    real['O1'] = 'CTA EQUIVALENTE'
    real['P1'] = ' C de C Equiv '
    ppto = libro['BD Ppto']
    ppto['G1'] = 'AÑO'
    assert normalizar_headers(real)['cuenta equiv'] == 15
    assert normalizar_headers(real)['cdec equiv'] == 16
    salida = io.BytesIO()
    libro.save(salida)
    salida.seek(0)
    salida.name = 'headers-mixtos.xlsx'
    resultado = procesar_carga_financiera(salida, proyecto=proyecto_247, anio=2026, mes=6, usuario=None)
    assert resultado.exito, resultado.error


@pytest.mark.django_db
def test_periodo_invalido_reporta_hoja_fila_y_columna(proyecto_247):
    libro = load_workbook(FIXTURE_REAL)
    libro['BD Real']['F2'] = 202699
    salida = io.BytesIO()
    libro.save(salida)
    salida.seek(0)
    salida.name = 'periodo-invalido.xlsx'
    resultado = procesar_carga_financiera(salida, proyecto=proyecto_247, anio=2026, mes=6, usuario=None)
    assert not resultado.exito
    assert "BD Real, fila 2, columna 'Periodo'" in resultado.error


@pytest.mark.django_db
def test_linea_historica_sin_senal_queda_sin_clasificar_explicito(proyecto_247):
    libro = load_workbook(FIXTURE_REAL)
    libro['BD Real']['I2'] = None
    salida = io.BytesIO()
    libro.save(salida)
    salida.seek(0)
    salida.name = 'historico-sin-senal.xlsx'
    resultado = procesar_carga_financiera(salida, proyecto=proyecto_247, anio=2026, mes=6, usuario=None)
    assert resultado.exito
    linea = resultado.carga.lineas.filter(tipo=LineaCargaFinanciera.Tipo.REAL).first()
    assert linea.tipo_operacional == 'sin_clasificar'
    assert linea.datos_origen['tipo_operacional']['valor_fuente'] is None
