"""Regresiones Sprint B #246 — seis indicadores + dashboard histórico.

No reusa fixtures de test_finv2_carga.py / test_finv2_carga_views.py: cada
prueba construye su propio libro para no depender de mutaciones de otro
archivo (política del repo, ver CLAUDE.md).
"""

import io
from datetime import datetime
from decimal import Decimal

import pytest
from openpyxl import Workbook

from apps.contratos.models import Contrato
from apps.financiero.importers_finv2_carga import procesar_carga_financiera
from apps.financiero.models_finv2_carga import CargaFinanciera, LineaCargaFinanciera
from apps.financiero.views_finv2_carga import CargaFinancieraView


def _libro(*, neto=100, ppto=120, mes=1, anio=2026, grupo="Materiales"):
    libro = Workbook()
    real = libro.active
    real.title = "BD Real"
    real.append(
        [
            "Auxiliar",
            "Desc. auxiliar",
            "Neto",
            "Fecha",
            "Docto.",
            "Periodo",
            "x",
            "x",
            "x",
            "x",
            "x",
            "x",
            "x",
            "x",
            "Cuenta Equiv",
            "CdeC equiv",
        ]
    )
    real.append(
        [
            1,
            "Compra cable",
            neto,
            datetime(anio, mes, 28),
            "DOC-1",
            anio * 100 + mes,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            grupo,
            "TRANSELCA",
        ]
    )
    presupuesto = libro.create_sheet("BD Ppto")
    presupuesto.append(
        ["Tipo", "Proyecto", "Rubro", "Clasificacion", "Valor", "mes", "año", "ciudad"]
    )
    if ppto is not None:
        presupuesto.append(
            ["Presupuesto", "Transelca", grupo, "Costos", ppto, mes, anio, "Barranquilla"]
        )
    homologacion = libro.create_sheet("Homologacion")
    homologacion.append(["tipo", "Grupo", "Concepto", "Rubro"])
    salida = io.BytesIO()
    libro.save(salida)
    salida.seek(0)
    salida.name = "TRANSELCA.xlsx"
    return salida


def _libro_sin_real(*, ppto=120, mes=1, anio=2026):
    """Libro sin líneas REAL (solo BD Ppto) — cubre costo real = 0."""
    libro = Workbook()
    real = libro.active
    real.title = "BD Real"
    real.append(
        [
            "Auxiliar",
            "Desc. auxiliar",
            "Neto",
            "Fecha",
            "Docto.",
            "Periodo",
            "x",
            "x",
            "x",
            "x",
            "x",
            "x",
            "x",
            "x",
            "Cuenta Equiv",
            "CdeC equiv",
        ]
    )
    presupuesto = libro.create_sheet("BD Ppto")
    presupuesto.append(
        ["Tipo", "Proyecto", "Rubro", "Clasificacion", "Valor", "mes", "año", "ciudad"]
    )
    presupuesto.append(
        ["Presupuesto", "Transelca", "Materiales", "Costos", ppto, mes, anio, "Barranquilla"]
    )
    homologacion = libro.create_sheet("Homologacion")
    homologacion.append(["tipo", "Grupo", "Concepto", "Rubro"])
    salida = io.BytesIO()
    libro.save(salida)
    salida.seek(0)
    salida.name = "SOLO_PPTO.xlsx"
    return salida


@pytest.fixture
def proyecto(db):
    return Contrato.objects.create(
        codigo="TRAN-246-B", nombre="Transelca", unidad_negocio="MANTENIMIENTO"
    )


# --- B1: fórmulas sobre bases financieras correctas -------------------------


@pytest.mark.django_db
def test_indicadores_usan_costo_real_y_presupuestado_no_facturacion_inventada(proyecto):
    """Facturación/Rentabilidad quedan explícitamente sin base: la fuente no trae ingresos."""
    resultado = procesar_carga_financiera(
        _libro(neto=100, ppto=120), proyecto=proyecto, anio=2026, mes=1, usuario=None
    )
    assert resultado.exito, resultado.error

    indicadores = CargaFinancieraView._indicadores(resultado.carga)
    por_nombre = {i["nombre"]: i for i in indicadores}

    assert por_nombre["Facturación vs meta"]["valor"] is None
    assert por_nombre["Facturación vs meta"]["sin_base"] is True
    assert "no trae" in por_nombre["Facturación vs meta"]["causa"]

    assert por_nombre["Ratio costos/facturación"]["valor"] is None
    assert por_nombre["Ratio costos/facturación"]["sin_base"] is True

    assert por_nombre["Rentabilidad operativa"]["valor"] is None
    assert por_nombre["Rentabilidad operativa"]["sin_base"] is True

    # Margen bruto y Cumplimiento presupuesto SÍ tienen base (costo real vs presupuesto).
    assert por_nombre["Margen bruto"]["valor"] == Decimal("20.00")  # 120 ppto - 100 real
    assert por_nombre["Margen bruto"]["sin_base"] is False
    assert por_nombre["Margen bruto"]["alerta"] is False

    cumplimiento = por_nombre["Cumplimiento presupuesto"]
    assert cumplimiento["valor"] == Decimal("83.33")  # 100/120 * 100
    assert cumplimiento["sin_base"] is False
    assert cumplimiento["alerta"] is False


@pytest.mark.django_db
def test_cumplimiento_presupuesto_alerta_cuando_costo_real_excede_ppto(proyecto):
    resultado = procesar_carga_financiera(
        _libro(neto=150, ppto=120), proyecto=proyecto, anio=2026, mes=1, usuario=None
    )
    assert resultado.exito, resultado.error

    indicadores = CargaFinancieraView._indicadores(resultado.carga)
    por_nombre = {i["nombre"]: i for i in indicadores}

    assert por_nombre["Cumplimiento presupuesto"]["alerta"] is True
    assert "sobre lo aprobado" in por_nombre["Cumplimiento presupuesto"]["causa"]
    assert por_nombre["Margen bruto"]["valor"] == Decimal("-30.00")
    assert por_nombre["Margen bruto"]["alerta"] is True


@pytest.mark.django_db
def test_cumplimiento_presupuesto_sin_base_cuando_no_hay_presupuesto_cargado(proyecto):
    """Denominador cero (sin BD Ppto para el período) → estado explícito, nunca inventado."""
    resultado = procesar_carga_financiera(
        _libro(neto=100, ppto=None), proyecto=proyecto, anio=2026, mes=1, usuario=None
    )
    assert resultado.exito, resultado.error

    indicadores = CargaFinancieraView._indicadores(resultado.carga)
    por_nombre = {i["nombre"]: i for i in indicadores}

    assert por_nombre["Cumplimiento presupuesto"]["valor"] is None
    assert por_nombre["Cumplimiento presupuesto"]["sin_base"] is True
    assert "BD Ppto vacía" in por_nombre["Cumplimiento presupuesto"]["causa"]


@pytest.mark.django_db
def test_variacion_mes_a_mes_sin_base_sin_periodo_anterior(proyecto):
    resultado = procesar_carga_financiera(
        _libro(neto=100, ppto=120, mes=1), proyecto=proyecto, anio=2026, mes=1, usuario=None
    )
    assert resultado.exito, resultado.error

    indicadores = CargaFinancieraView._indicadores(resultado.carga)
    variacion = next(i for i in indicadores if i["nombre"] == "Variación mes a mes")
    assert variacion["valor"] is None
    assert variacion["sin_base"] is True
    assert "Sin base comparable" in variacion["causa"]


@pytest.mark.django_db
def test_variacion_mes_a_mes_calcula_contra_periodo_anterior_vigente(proyecto):
    anterior = procesar_carga_financiera(
        _libro(neto=100, ppto=100, mes=1), proyecto=proyecto, anio=2026, mes=1, usuario=None
    )
    assert anterior.exito, anterior.error
    actual = procesar_carga_financiera(
        _libro(neto=150, ppto=100, mes=2), proyecto=proyecto, anio=2026, mes=2, usuario=None
    )
    assert actual.exito, actual.error

    indicadores = CargaFinancieraView._indicadores(actual.carga)
    variacion = next(i for i in indicadores if i["nombre"] == "Variación mes a mes")
    assert variacion["sin_base"] is False
    assert variacion["valor"] == Decimal("50.00")  # (150-100)/100 * 100
    assert variacion["alerta"] is True
    assert "subió" in variacion["causa"]


@pytest.mark.django_db
def test_variacion_mes_a_mes_ignora_version_no_vigente_del_periodo_anterior(proyecto):
    """El período anterior debe leerse de su carga VIGENTE, no de una versión reemplazada."""
    primera = procesar_carga_financiera(
        _libro(neto=999, ppto=100, mes=1), proyecto=proyecto, anio=2026, mes=1, usuario=None
    )
    assert primera.exito, primera.error
    segunda = procesar_carga_financiera(
        _libro(neto=100, ppto=100, mes=1), proyecto=proyecto, anio=2026, mes=1, usuario=None
    )
    assert segunda.exito, segunda.error
    assert (
        CargaFinanciera.objects.filter(proyecto=proyecto, anio=2026, mes=1, vigente=True).count()
        == 1
    )

    actual = procesar_carga_financiera(
        _libro(neto=150, ppto=100, mes=2), proyecto=proyecto, anio=2026, mes=2, usuario=None
    )
    assert actual.exito, actual.error

    indicadores = CargaFinancieraView._indicadores(actual.carga)
    variacion = next(i for i in indicadores if i["nombre"] == "Variación mes a mes")
    # Debe comparar contra la versión vigente de enero (100), no la reemplazada (999).
    assert variacion["valor"] == Decimal("50.00")


@pytest.mark.django_db
def test_indicadores_sin_carga_devuelve_lista_vacia():
    assert CargaFinancieraView._indicadores(None) == []


@pytest.mark.django_db
def test_indicadores_respetan_filtro_tipo_operacional(proyecto):
    libro = Workbook()
    real = libro.active
    real.title = "BD Real"
    real.append(
        [
            "Desc. auxiliar",
            "Neto",
            "Periodo",
            "Cuenta Equiv",
            "CdeC equiv",
            "C.Costo",
            "Desc. C.O. movto.",
        ]
    )
    real.append(
        [
            "Cable construcción",
            100,
            202601,
            "Materiales",
            "TRANSELCA",
            "400102",
            "CONSTRUCCION LINEA",
        ]
    )
    real.append(
        [
            "Cable mantenimiento",
            300,
            202601,
            "Materiales",
            "TRANSELCA",
            "400103",
            "MANTENIMIENTO LINEA",
        ]
    )
    ppto = libro.create_sheet("BD Ppto")
    ppto.append(["Tipo", "Proyecto", "Rubro", "Clasificacion", "Valor", "mes", "año"])
    ppto.append(["Construcción", "Transelca", "Materiales", "Costos", 500, 1, 2026])
    ppto.append(["Mantenimiento", "Transelca", "Materiales", "Costos", 500, 1, 2026])
    homologacion = libro.create_sheet("Homologacion")
    homologacion.append(["tipo", "Grupo", "Concepto", "Rubro"])
    salida = io.BytesIO()
    libro.save(salida)
    salida.seek(0)
    salida.name = "ramas.xlsx"

    resultado = procesar_carga_financiera(salida, proyecto=proyecto, anio=2026, mes=1, usuario=None)
    assert resultado.exito, resultado.error

    indicadores_construccion = CargaFinancieraView._indicadores(
        resultado.carga,
        tipo_operacional=LineaCargaFinanciera.TipoOperacional.CONSTRUCCION,
    )
    margen_construccion = next(i for i in indicadores_construccion if i["nombre"] == "Margen bruto")
    assert margen_construccion["valor"] == Decimal(
        "400.00"
    )  # 500 ppto construcción - 100 real construcción

    indicadores_mantenimiento = CargaFinancieraView._indicadores(
        resultado.carga,
        tipo_operacional=LineaCargaFinanciera.TipoOperacional.MANTENIMIENTO,
    )
    margen_mantenimiento = next(
        i for i in indicadores_mantenimiento if i["nombre"] == "Margen bruto"
    )
    assert margen_mantenimiento["valor"] == Decimal("200.00")  # 500 ppto - 300 real mantenimiento


# --- B2: dashboard histórico — desglose reconciliado y tendencia 6 meses ----


@pytest.mark.django_db
def test_desglose_costos_suma_exacto_el_total_filtrado(proyecto):
    libro = Workbook()
    real = libro.active
    real.title = "BD Real"
    real.append(["Desc. auxiliar", "Neto", "Periodo", "Cuenta Equiv", "CdeC equiv"])
    real.append(["Cable", 100, 202601, "Materiales", "TRANSELCA"])
    real.append(["Viáticos hotel", Decimal("33.33"), 202601, "Viáticos", "TRANSELCA"])
    real.append(["Viáticos comida", Decimal("16.67"), 202601, "Viáticos", "TRANSELCA"])
    ppto = libro.create_sheet("BD Ppto")
    ppto.append(["Tipo", "Proyecto", "Rubro", "Clasificacion", "Valor", "mes", "año"])
    ppto.append(["Presupuesto", "Transelca", "Materiales", "Costos", 200, 1, 2026])
    homologacion = libro.create_sheet("Homologacion")
    homologacion.append(["tipo", "Grupo", "Concepto", "Rubro"])
    salida = io.BytesIO()
    libro.save(salida)
    salida.seek(0)
    salida.name = "desglose.xlsx"

    resultado = procesar_carga_financiera(salida, proyecto=proyecto, anio=2026, mes=1, usuario=None)
    assert resultado.exito, resultado.error

    desglose = CargaFinancieraView._desglose_costos(resultado.carga)
    total_desglosado = sum((fila["valor"] for fila in desglose), Decimal("0.00"))
    costo_real_total, _ = CargaFinancieraView._totales(resultado.carga)
    assert total_desglosado == costo_real_total == Decimal("150.00")
    grupos = {fila["grupo"]: fila["valor"] for fila in desglose}
    assert grupos["Materiales"] == Decimal("100.00")
    assert grupos["Viáticos"] == Decimal("50.00")


@pytest.mark.django_db
def test_desglose_costos_sin_lineas_reales_es_lista_vacia(proyecto):
    resultado = procesar_carga_financiera(
        _libro_sin_real(ppto=100), proyecto=proyecto, anio=2026, mes=1, usuario=None
    )
    assert resultado.exito, resultado.error
    assert CargaFinancieraView._desglose_costos(resultado.carga) == []


@pytest.mark.django_db
def test_tendencia_6_meses_incluye_solo_cargas_vigentes_del_proyecto_hasta_el_periodo_actual(
    proyecto,
):
    for mes in range(1, 8):  # 7 meses cargados; sólo deben aparecer los últimos 6 hasta el actual
        resultado = procesar_carga_financiera(
            _libro(neto=10 * mes, ppto=10 * mes, mes=mes),
            proyecto=proyecto,
            anio=2026,
            mes=mes,
            usuario=None,
        )
        assert resultado.exito, resultado.error
    carga_julio = CargaFinanciera.objects.get(proyecto=proyecto, anio=2026, mes=7, vigente=True)

    tendencia = CargaFinancieraView._tendencia_6_meses(carga_julio)
    assert len(tendencia) == 6
    assert [p["periodo"] for p in tendencia] == [
        "02/2026",
        "03/2026",
        "04/2026",
        "05/2026",
        "06/2026",
        "07/2026",
    ]
    assert tendencia[-1]["es_actual"] is True
    assert tendencia[-1]["costo_real"] == Decimal("70.00")


@pytest.mark.django_db
def test_tendencia_6_meses_sin_carga_devuelve_lista_vacia():
    assert CargaFinancieraView._tendencia_6_meses(None) == []


@pytest.mark.django_db
def test_dashboard_view_expone_resumen_alertas_desglose_y_tendencia(client, admin_user, proyecto):
    client.force_login(admin_user)
    resultado = procesar_carga_financiera(
        _libro(neto=150, ppto=120, mes=1), proyecto=proyecto, anio=2026, mes=1, usuario=admin_user
    )
    assert resultado.exito, resultado.error

    respuesta = client.get(f"/financiero/carga-financiera/?proyecto={proyecto.pk}&anio=2026&mes=1")
    assert respuesta.status_code == 200
    contenido = respuesta.content.decode()

    assert "Resumen ejecutivo" in contenido
    assert "Desglose de costos por grupo" in contenido
    assert "Tendencia últimos 6 meses" in contenido
    assert (
        "alertas activas" in contenido
    )  # Cumplimiento presupuesto + Margen bruto en rojo (neto 150 > ppto 120)
    assert "Sin base comparable" in contenido  # Facturación vs meta / Ratio / Rentabilidad
    assert "Materiales" in contenido  # fila del desglose
    assert "01/2026" in contenido  # punto de la tendencia
