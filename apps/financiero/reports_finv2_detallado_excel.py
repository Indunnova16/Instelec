"""Excel detallado del dashboard financiero integrado (#246 Sprint D).

Consume el mismo ``payload`` que ``reports_finv2_ejecutivo_pdf`` -- ver su
docstring. openpyxl ya está declarado en ``requirements/base.txt`` (y en
``requirements/dev_lite.txt``), mismo patrón que
``views_finv2_carga.DescargarVersionHomologacionXlsxView``/
``ExportarTablaMaestraXlsxView``.

5 hojas (Sprint D, versión 1.0 completa): Resumen / Facturación / Costos /
Proyectos / Gráficos. Todos los valores monetarios a 2 decimales
(``NUM_FORMAT``).
"""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font

NUM_FORMAT = "#,##0.00"
ZERO = Decimal("0.00")


def _num(valor: Any) -> float:
    if valor is None:
        return 0.0
    try:
        return float(Decimal(valor).quantize(Decimal("0.01")))
    except Exception:
        return 0.0


def _encabezado(hoja, columnas: list[str]) -> None:
    hoja.append(columnas)
    for celda in hoja[hoja.max_row]:
        celda.font = Font(bold=True)


def _hoja_resumen(libro: Workbook, payload: dict) -> None:
    hoja = libro.active
    hoja.title = "Resumen"
    proyecto = payload.get("proyecto")
    hoja.append(["Reporte ejecutivo financiero detallado"])
    hoja["A1"].font = Font(bold=True, size=14)
    hoja.append(["Proyecto", str(proyecto) if proyecto else "Todos los proyectos"])
    hoja.append(["Período", f"{payload['mes']:02d}/{payload['anio']}"])
    hoja.append([])

    resumen_totales = payload.get("resumen_totales") or {}
    hoja.append(["Costo real", _num(resumen_totales.get("costo_real"))])
    hoja[f"B{hoja.max_row}"].number_format = NUM_FORMAT
    hoja.append(["Costo presupuestado", _num(resumen_totales.get("costo_presupuestado"))])
    hoja[f"B{hoja.max_row}"].number_format = NUM_FORMAT
    hoja.append([])

    _encabezado(hoja, ["Indicador", "Valor", "Unidad", "Estado"])
    for indicador in payload.get("indicadores") or []:
        fila = hoja.max_row + 1
        valor = indicador.get("valor")
        hoja.append(
            [
                indicador.get("nombre", ""),
                "Sin base comparable"
                if indicador.get("sin_base") or valor is None
                else _num(valor),
                indicador.get("unidad", ""),
                "Alerta" if indicador.get("alerta") else "OK",
            ]
        )
        if not (indicador.get("sin_base") or valor is None):
            hoja[f"B{fila}"].number_format = NUM_FORMAT
    for columna, ancho in zip("ABCD", (34, 20, 10, 10), strict=True):
        hoja.column_dimensions[columna].width = ancho


def _hoja_facturacion(libro: Workbook, payload: dict) -> None:
    hoja = libro.create_sheet("Facturación")
    ingresos = payload.get("integracion_ingresos") or {}
    hoja.append(["Ingresos del período — total: ", _num(ingresos.get("total"))])
    hoja["B1"].number_format = NUM_FORMAT
    hoja.append([])
    _encabezado(hoja, ["Cliente", "Factura", "Fecha", "Total", "Estado"])
    for factura in ingresos.get("detalle") or []:
        fila = hoja.max_row + 1
        hoja.append(
            [
                factura.get("cliente", ""),
                factura.get("numero_factura", ""),
                factura.get("fecha_factura"),
                _num(factura.get("total")),
                factura.get("estado_display", ""),
            ]
        )
        hoja[f"D{fila}"].number_format = NUM_FORMAT
        if factura.get("fecha_factura"):
            hoja[f"C{fila}"].number_format = "DD/MM/YYYY"
    if not ingresos.get("detalle"):
        hoja.append(["Sin facturas de ingreso emitidas en este período/proyecto."])
    for columna, ancho in zip("ABCDE", (28, 16, 14, 14, 18), strict=True):
        hoja.column_dimensions[columna].width = ancho


def _hoja_costos(libro: Workbook, payload: dict) -> None:
    hoja = libro.create_sheet("Costos")
    gastos = payload.get("integracion_gastos") or {}
    hoja.append(["Gastos del período — total:", _num(gastos.get("total"))])
    hoja["B1"].number_format = NUM_FORMAT
    hoja.append(["Pendiente de pago:", _num(gastos.get("pendiente_pago"))])
    hoja["B2"].number_format = NUM_FORMAT
    hoja.append([])

    _encabezado(hoja, ["Desglose por grupo (costo real)", "Valor", "% del total"])
    for grupo in payload.get("desglose_costos") or []:
        fila = hoja.max_row + 1
        hoja.append([grupo.get("grupo", ""), _num(grupo.get("valor")), grupo.get("porcentaje")])
        hoja[f"B{fila}"].number_format = NUM_FORMAT
    if not payload.get("desglose_costos"):
        hoja.append(["Sin desglose: no hay carga financiera vigente para el período/proyecto."])
    hoja.append([])

    _encabezado(hoja, ["Proveedor", "Documento", "Fecha", "Categoría", "Total", "Estado"])
    for gasto in gastos.get("detalle") or []:
        fila = hoja.max_row + 1
        hoja.append(
            [
                gasto.get("proveedor", ""),
                gasto.get("numero_documento", ""),
                gasto.get("fecha"),
                gasto.get("categoria", ""),
                _num(gasto.get("total")),
                gasto.get("estado_display", ""),
            ]
        )
        hoja[f"E{fila}"].number_format = NUM_FORMAT
        if gasto.get("fecha"):
            hoja[f"C{fila}"].number_format = "DD/MM/YYYY"
    if not gastos.get("detalle"):
        hoja.append(["Sin facturas de gasto en este período/proyecto."])
    for columna, ancho in zip("ABCDEF", (28, 16, 14, 18, 14, 18), strict=True):
        hoja.column_dimensions[columna].width = ancho


def _hoja_proyectos(libro: Workbook, payload: dict) -> None:
    hoja = libro.create_sheet("Proyectos")
    proyecto = payload.get("proyecto")
    clientes = payload.get("integracion_clientes") or {}
    hoja.append(["Proyecto filtrado", str(proyecto) if proyecto else "Todos los proyectos"])
    hoja.append(["Clientes activos (maestro #261)", clientes.get("activos_total", 0)])
    hoja.append(["Clientes facturados en el período", clientes.get("facturados_periodo", 0)])
    hoja.append([])

    nomina = payload.get("integracion_nomina") or {}
    hoja.append(
        ["Nómina (producción diaria) — costo del período:", _num(nomina.get("costo_nomina"))]
    )
    hoja[f"B{hoja.max_row}"].number_format = NUM_FORMAT
    hoja.append(["Horas trabajadas:", _num(nomina.get("horas_trabajadas"))])
    hoja.append([])

    _encabezado(hoja, ["Cliente", "Total facturado", "Cantidad de facturas"])
    for cliente in clientes.get("detalle") or []:
        fila = hoja.max_row + 1
        hoja.append(
            [
                cliente.get("cliente", ""),
                _num(cliente.get("total_facturado")),
                cliente.get("cantidad_facturas", 0),
            ]
        )
        hoja[f"B{fila}"].number_format = NUM_FORMAT
    if not clientes.get("detalle"):
        hoja.append(["Ningún cliente facturado en este período/proyecto."])
    for columna, ancho in zip("ABC", (30, 18, 20), strict=True):
        hoja.column_dimensions[columna].width = ancho


def _hoja_graficos(libro: Workbook, payload: dict) -> None:
    hoja = libro.create_sheet("Gráficos")
    tendencia = payload.get("tendencia_6_meses") or []
    _encabezado(hoja, ["Período", "Costo real", "Costo presupuestado"])
    for punto in tendencia:
        fila = hoja.max_row + 1
        hoja.append(
            [punto["periodo"], _num(punto["costo_real"]), _num(punto["costo_presupuestado"])]
        )
        hoja[f"B{fila}"].number_format = NUM_FORMAT
        hoja[f"C{fila}"].number_format = NUM_FORMAT

    if not tendencia:
        hoja.append(
            ["Sin histórico de cargas previas para este proyecto -- sin datos para graficar."]
        )
        return

    grafico = BarChart()
    grafico.title = "Costo real vs presupuestado (tendencia)"
    grafico.y_axis.title = "Valor"
    grafico.x_axis.title = "Período"
    filas_datos = len(tendencia)
    datos = Reference(hoja, min_col=2, max_col=3, min_row=1, max_row=1 + filas_datos)
    categorias = Reference(hoja, min_col=1, min_row=2, max_row=1 + filas_datos)
    grafico.add_data(datos, titles_from_data=True)
    grafico.set_categories(categorias)
    hoja.add_chart(grafico, "E2")


def generar_excel_detallado(payload: dict) -> bytes:
    """5 hojas Resumen/Facturación/Costos/Proyectos/Gráficos -- versión 1.0
    completa (Sprint D), 2 decimales en todo valor monetario."""
    libro = Workbook()
    _hoja_resumen(libro, payload)
    _hoja_facturacion(libro, payload)
    _hoja_costos(libro, payload)
    _hoja_proyectos(libro, payload)
    _hoja_graficos(libro, payload)

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()
