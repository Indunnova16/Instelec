"""Issue #271 Sprint A — A3: import/export Excel de Colaboradores con las
2 columnas nuevas (fecha_firma_contrato / fecha_ingreso_proyecto), en vez
de la antigua columna única `Fecha Ingreso`.

Cubre:
- `PersonalCuadrillaUploadView`: el encabezado legado 'FECHA INGRESO' (y su
  posición legacy sin encabezados) sigue siendo aceptado, pero ahora puebla
  `fecha_firma_contrato` -- ya NO el `fecha_ingreso` deprecado (A1).
  'FECHA INGRESO PROYECTO' es la columna nueva y opcional, sin posición
  legacy, resuelta solo por encabezado.
- `ColaboradorExportView`: la columna única 'Fecha Ingreso' se reemplaza por
  'Fecha Firma Contrato' + 'Fecha Ingreso Proyecto'.
- Round-trip export -> import contra un registro con dato LEGACY (backfillado
  por A1), no solo fixtures nuevas del sub-item.
"""
from datetime import date
from io import BytesIO

import openpyxl
import pytest
from django.contrib.auth import get_user_model
from django.test import Client
from django.urls import reverse

from apps.cuadrillas.models import Cargo, PersonalCuadrilla

pytestmark = pytest.mark.django_db

Usuario = get_user_model()


@pytest.fixture(autouse=True)
def _seed_cargo_liniero():
    """Este módulo vive en tests/cuadrillas/ (NO apps/cuadrillas/), fuera
    del alcance del conftest autouse que siembra el catálogo completo de
    Cargo (apps/cuadrillas/conftest.py). `rol_cuadrilla` es FK requerida
    (sin blank=True) -- sin un Cargo real, el import descarta toda fila."""
    Cargo.objects.get_or_create(
        codigo="LINIERO_I", defaults={"nombre": "Liniero I", "activo": True}
    )


def _crear_admin():
    return Usuario.objects.create_user(
        email="admin_271_a3@test.com",
        password="testpass123!",
        first_name="Admin",
        last_name="Test271A3",
        rol="admin",
        is_staff=True,
        is_superuser=True,
    )


@pytest.fixture
def admin_client():
    client = Client()
    client.force_login(_crear_admin())
    return client


def _build_workbook(headers, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    archivo = BytesIO()
    wb.save(archivo)
    archivo.seek(0)
    archivo.name = "colaboradores_271_a3.xlsx"
    return archivo


# ---------------------------------------------------------------------------
# Import -- header legado 'FECHA INGRESO' -> fecha_firma_contrato (compat)
# ---------------------------------------------------------------------------
def test_import_header_legado_fecha_ingreso_puebla_fecha_firma_contrato(admin_client):
    """Requisito #1: import legacy (solo header 'FECHA INGRESO') puebla
    `fecha_firma_contrato`, NO el `fecha_ingreso` deprecado."""
    archivo = _build_workbook(
        ["Nombre", "Documento", "Cargo", "Salario Base", "Fecha Ingreso", "Fecha Salida"],
        [["Colaborador A3 Legado", "271-A3-0001", "LINIERO_I", 1750905, "2025-02-01", ""]],
    )

    resp = admin_client.post(reverse("cuadrillas:personal_upload"), {"archivo": archivo})
    assert resp.status_code in (200, 302)

    persona = PersonalCuadrilla.objects.get(documento="271-A3-0001")
    assert persona.fecha_firma_contrato == date(2025, 2, 1)
    assert persona.fecha_ingreso_proyecto is None
    # El legacy deprecado ya no se escribe desde este importador (A3, mismo
    # criterio que el form en A2).
    assert persona.fecha_ingreso is None


def test_import_plantilla_legacy_sin_encabezados_usa_posicion_4_para_firma_contrato(admin_client):
    """La plantilla histórica sin encabezados canónicos (Nombre | Documento
    en posiciones 0/1) sigue resolviendo la fecha en la posición legacy 4."""
    archivo = _build_workbook(
        ["Nombre", "Documento", "Cargo"],
        [["Formato Posicional", "271-A3-POS1", "LINIERO_I", 1500000, "2024-06-15", ""]],
    )

    admin_client.post(reverse("cuadrillas:personal_upload"), {"archivo": archivo})

    persona = PersonalCuadrilla.objects.get(documento="271-A3-POS1")
    assert persona.fecha_firma_contrato == date(2024, 6, 15)


# ---------------------------------------------------------------------------
# Import -- ambos headers -> las 2 fechas
# ---------------------------------------------------------------------------
def test_import_con_ambos_headers_puebla_las_2_fechas(admin_client):
    """Requisito #2: import con ambos headers puebla las 2 fechas."""
    archivo = _build_workbook(
        [
            "Nombre", "Documento", "Cargo", "Salario Base",
            "Fecha Ingreso", "Fecha Ingreso Proyecto", "Fecha Salida",
        ],
        [[
            "Colaborador A3 Completo", "271-A3-0002", "LINIERO_I", 1800000,
            "2025-01-10", "2025-01-20", "",
        ]],
    )

    resp = admin_client.post(reverse("cuadrillas:personal_upload"), {"archivo": archivo})
    assert resp.status_code in (200, 302)

    persona = PersonalCuadrilla.objects.get(documento="271-A3-0002")
    assert persona.fecha_firma_contrato == date(2025, 1, 10)
    assert persona.fecha_ingreso_proyecto == date(2025, 1, 20)
    assert persona.activo


def test_edge_columnas_reordenadas_resuelven_fecha_ingreso_proyecto_por_encabezado(admin_client):
    """Edge case: encabezados reordenados/no contiguos -- 'Fecha Ingreso
    Proyecto' se resuelve por nombre, no por posición."""
    archivo = _build_workbook(
        [
            "Cargo", "Documento", "Fecha Ingreso Proyecto", "Nombre",
            "Salario Base", "Fecha Ingreso",
        ],
        [[
            "LINIERO_I", "271-A3-ORDEN1", "2026-03-01", "Columnas Reordenadas",
            1500000, "2026-01-01",
        ]],
    )

    admin_client.post(reverse("cuadrillas:personal_upload"), {"archivo": archivo})

    persona = PersonalCuadrilla.objects.get(documento="271-A3-ORDEN1")
    assert persona.nombre == "Columnas Reordenadas"
    assert persona.fecha_firma_contrato == date(2026, 1, 1)
    assert persona.fecha_ingreso_proyecto == date(2026, 3, 1)


def test_edge_sin_columna_fecha_ingreso_proyecto_queda_none_sin_reventar(admin_client):
    """Edge case: archivo SIN el header nuevo -- opcional, no debe romper el
    import ni caer en una posición legacy inexistente."""
    archivo = _build_workbook(
        ["Nombre", "Documento", "Cargo", "Salario Base", "Fecha Ingreso"],
        [["Sin Columna Nueva", "271-A3-SINCOL", "LINIERO_I", 1000000, "2025-05-05"]],
    )

    resp = admin_client.post(reverse("cuadrillas:personal_upload"), {"archivo": archivo})
    assert resp.status_code in (200, 302)

    persona = PersonalCuadrilla.objects.get(documento="271-A3-SINCOL")
    assert persona.fecha_firma_contrato == date(2025, 5, 5)
    assert persona.fecha_ingreso_proyecto is None


# ---------------------------------------------------------------------------
# Export -- columna única reemplazada por las 2 nuevas
# ---------------------------------------------------------------------------
def test_export_headers_y_valores_en_las_2_columnas_nuevas(admin_client):
    """Requisito #3: export abre con openpyxl, headers y valores correctos
    en las 2 columnas nuevas."""
    PersonalCuadrilla.objects.create(
        nombre="Exportar A3",
        documento="271-A3-EXPORT",
        rol_cuadrilla_id="LINIERO_I",
        fecha_firma_contrato=date(2023, 7, 1),
        fecha_ingreso_proyecto=date(2023, 7, 15),
    )

    resp = admin_client.get(reverse("cuadrillas:colaboradores_export"))
    assert resp.status_code == 200

    wb = openpyxl.load_workbook(BytesIO(resp.content))
    ws = wb.active
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]

    assert "Fecha Firma Contrato" in headers
    assert "Fecha Ingreso Proyecto" in headers
    assert "Fecha Ingreso" not in headers

    idx_firma = headers.index("Fecha Firma Contrato")
    idx_proyecto = headers.index("Fecha Ingreso Proyecto")
    fila = next(
        row for row in ws.iter_rows(min_row=2, values_only=True)
        if row[0] == "271-A3-EXPORT"
    )
    assert fila[idx_firma] == "2023-07-01"
    assert fila[idx_proyecto] == "2023-07-15"


def test_export_registro_sin_fechas_deja_celdas_vacias(admin_client):
    """Edge case: colaborador sin ninguna de las 2 fechas -- export no debe
    reventar ni escribir 'None'."""
    PersonalCuadrilla.objects.create(
        nombre="Sin Fechas A3", documento="271-A3-VACIO", rol_cuadrilla_id="LINIERO_I",
    )

    resp = admin_client.get(reverse("cuadrillas:colaboradores_export"))
    wb = openpyxl.load_workbook(BytesIO(resp.content))
    ws = wb.active
    headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    idx_firma = headers.index("Fecha Firma Contrato")
    idx_proyecto = headers.index("Fecha Ingreso Proyecto")

    fila = next(
        row for row in ws.iter_rows(min_row=2, values_only=True)
        if row[0] == "271-A3-VACIO"
    )
    # openpyxl escribe '' pero al releer una celda vacía devuelve None
    # (no hay diferencia observable en el .xlsx real).
    assert fila[idx_firma] in (None, "")
    assert fila[idx_proyecto] in (None, "")


# ---------------------------------------------------------------------------
# Round-trip contra dato LEGACY (backfillado por A1, no solo fixture propia)
# ---------------------------------------------------------------------------
def test_roundtrip_export_import_preserva_las_2_fechas_de_dato_legacy(admin_client):
    """Colaborador con `fecha_ingreso` legacy YA backfillada por A1 en
    `fecha_firma_contrato` (migración 0041) -- el export/import round-trip
    debe preservar ambas fechas sin perder el dato legacy."""
    legado = PersonalCuadrilla.objects.create(
        nombre="Registro Legacy 271 A3",
        documento="271-A3-LEGACY",
        rol_cuadrilla_id="LINIERO_I",
        fecha_ingreso=date(2019, 3, 1),  # legacy físico, ya deprecado
        fecha_firma_contrato=date(2019, 3, 1),  # backfillado por A1
        fecha_ingreso_proyecto=date(2019, 4, 1),
    )

    exportado = admin_client.get(reverse("cuadrillas:colaboradores_export"))
    assert exportado.status_code == 200

    archivo = BytesIO(exportado.content)
    archivo.name = "colaboradores_271_a3_roundtrip.xlsx"
    admin_client.post(reverse("cuadrillas:personal_upload"), {"archivo": archivo})

    legado.refresh_from_db()
    assert PersonalCuadrilla.objects.filter(documento="271-A3-LEGACY").count() == 1
    assert legado.fecha_firma_contrato == date(2019, 3, 1)
    assert legado.fecha_ingreso_proyecto == date(2019, 4, 1)
